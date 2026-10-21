"""Точечная правка (app/core/graph.py, app/pipeline/patch.py, docs/ARCHITECTURE.md §5) в режиме mock.

Инвариант §5.3: объекты вне changed_ids побайтно совпадают (хеши model_dump_json), суммы часов сходятся,
ID не переиспользуются; неудачная правка комплект не меняет и объясняет причину.
"""

from __future__ import annotations

import asyncio
import io
import re
from pathlib import Path

import docx
import pytest
from fastapi.testclient import TestClient

from app.config import Settings
from app.core.graph import DependencyGraph, affected, kind_changes
from app.llm import BudgetExceeded
from app.main import create_app
from app.models import HoursRow, Kit, PatchOperation, PatchPlanDraft, QuestionDraft, TestsDraft, WorkKind
from app.pipeline.base import RequestInput, StepContext
from app.pipeline.default import DefaultPipeline
from app.pipeline.patch import (
    PatchFailed,
    diff_ids,
    object_hashes,
    parse_command,
    parse_hours,
    plan_from_draft,
)
from app.render.report_md import build_report, render_report_md
from tests.intake_queries import OPEN

FIXTURES = Path(__file__).parent / "fixtures" / "kits"
SERVICE_ID = re.compile(r"\b(?:T|L|S|Q|IQ|LO|TS|M)\d+(?:\.\d+)?\b")


def load(n: int) -> Kit:
    return Kit.model_validate_json((FIXTURES / f"open_q{n}.json").read_text(encoding="utf-8"))


@pytest.fixture
def pipeline(tmp_path: Path) -> DefaultPipeline:
    return DefaultPipeline(
        Settings(llm_mode="mock", sessions_dir=tmp_path / "sessions", cost_log=tmp_path / "cost.jsonl")
    )


def run_patch(pipeline: DefaultPipeline, kit: Kit, target: str, instruction: str, tmp: Path):  # noqa: ANN201
    ctx = StepContext(session_id="s", inputs_dir=tmp / "in", files_dir=tmp / "out", request=RequestInput())
    return asyncio.run(pipeline.patch(kit.model_copy(deep=True), target, instruction, ctx))


def assert_only_changed(before: Kit, after: Kit, changed: list[str], removed: list[str] = ()) -> None:
    """Всё, чего нет в changed_ids, побайтно прежнее; удалённых нет в комплекте."""
    old, new = object_hashes(before), object_hashes(after)
    for obj_id, value in old.items():
        if obj_id in removed:
            assert obj_id not in new
        elif obj_id not in changed:
            assert new[obj_id] == value, f"{obj_id} изменился, но его нет в changed_ids"


# --------------------------------------------------------------------------------------------------
# Граф зависимостей
# --------------------------------------------------------------------------------------------------


def test_graph_reverse_edges_and_closure() -> None:
    kit = load(3)
    graph = DependencyGraph.build(kit)
    assert {"L2.3", "S6", "Q5", "Q6", "Q7"} <= graph.direct("T2.3")
    assert {q.id for q in kit.questions if q.test_id == "TS2"} <= graph.direct("TS2")
    module_topics = {t.id for t in kit.topics_of("M2")}
    closure = graph.closure(["M2"])
    assert module_topics <= closure and "L2.3" in closure and "Q5" in closure


def test_affected_replace_topic_follows_table() -> None:
    kit = load(3)
    aff = affected(kit, PatchOperation.REPLACE_TOPIC, "T2.3")
    assert aff.content == ["T2.3"] and aff.lectures == ["L2.3"] and aff.seminars == ["S6"]
    assert set(aff.questions) == {q.id for q in kit.questions if q.topic_id == "T2.3"}
    assert not aff.retire and not aff.new_lectures and not aff.final_attestation


def test_affected_harden_test_is_only_its_questions() -> None:
    kit = load(1)
    aff = affected(kit, PatchOperation.HARDEN_TEST, "TS1")
    assert set(aff.questions) == set(kit.tests[0].question_ids) and aff.llm_targets() == aff.questions


def test_kind_changes_when_topic_loses_theory_and_gains_lab() -> None:
    before = load(3)
    after = before.model_copy(deep=True)
    topic = next(t for t in after.topics if t.id == "T2.3")
    topic.hours = HoursRow(lk=0, pr=3.5, lb=1.5, srs=0.5)
    aff = kind_changes(before, after)
    assert "L2.3" in aff.retire  # лекция была только по этой теме
    assert aff.content == ["T2.3"] and aff.drop_kinds == {"T2.3": [WorkKind.LK]}
    assert not aff.new_seminars  # задание у темы уже есть


# --------------------------------------------------------------------------------------------------
# Разбор команд кодом
# --------------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("target", "text", "operation"),
    [
        ("T2.3", "часы темы 2.3: ЛК 2, ПР 4", PatchOperation.CHANGE_HOURS),
        ("T2.3", "тема 2.3 на 6 часов", PatchOperation.CHANGE_HOURS),
        ("T2.3", "пересобери тему 2.3", PatchOperation.REGENERATE),
        ("T2.3", "", PatchOperation.REGENERATE),
        ("TS1", "усложни итоговый тест", PatchOperation.HARDEN_TEST),
        ("T2.3", "удали тему 2.3", PatchOperation.REMOVE_TOPIC),
        ("T2.3", "замени тему 2.3 на «Сводные таблицы»", PatchOperation.REPLACE_TOPIC),
        ("M2", "добавь тему «Очистка данных» в модуль 2", PatchOperation.ADD_TOPIC),
        ("LO1.2", "замени результат 1.2 на «Умеет строить сводные таблицы»", PatchOperation.REPLACE_LO),
    ],
)
def test_simple_commands_are_parsed_by_code(target: str, text: str, operation: PatchOperation) -> None:
    plan = parse_command(load(3), target, text)
    assert plan is not None and plan.operation == operation and plan.planned_by == "rules"


@pytest.mark.parametrize(
    "text",
    [
        "убери из темы 2.3 упоминание Excel",  # не удаление темы
        "замени в теме 2.3 пример на «Продажи магазина»",  # не замена темы
        "сделай задания ближе к работе бухгалтера",
        "усложни задание темы 2.3",  # задание, а не тест
    ],
)
def test_free_text_goes_to_the_skill(text: str) -> None:
    assert parse_command(load(3), "T2.3", text) is None


def asks(text: str, bare: bool = False) -> list[tuple]:
    return [(a.kind.value if a.kind else None, a.value, a.sign) for a in parse_hours(text, bare)]


def test_hours_are_read_by_code() -> None:
    assert asks("часы темы 2.3: ЛК 2, ПР 4", bare=True) == [("lk", 2.0, 0), ("pr", 4.0, 0)]
    assert asks("2 ч лекций и 4,5 часа практики") == [("lk", 2.0, 0), ("pr", 4.5, 0)]
    assert asks("сделай тему на 6 часов") == [(None, 6.0, 0)]
    # у каждой пары свой знак — по ближайшему глаголу; «до» — ровно столько
    assert asks("сократи практику на 1 час и добавь 1 час лекций") == [("pr", 1.0, -1), ("lk", 1.0, 1)]
    assert asks("уменьши лекции до 1 часа") == [("lk", 1.0, 0)]
    plan = parse_command(load(3), "T2.3", "добавь 2 часа практики")
    assert plan is not None and plan.operation == PatchOperation.CHANGE_HOURS
    assert [(a.kind, a.value, a.sign) for a in plan.hours] == [(WorkKind.PR, 2.0, 1)]
    replace = parse_command(load(3), "T2.3", "замени тему 2.3 на «Сводные таблицы»")
    assert replace is not None and replace.title == "Сводные таблицы"


@pytest.mark.parametrize(
    "text",
    [
        "в лекции 3 примера устарели, перепиши",
        "лекция 1 слишком длинная",
        "в практике 2 задания слишком простые",
        "в теме 2.3 срс 5 вопросов",
        "сделай задание на 2 часа работы",
        "добавь в тему «Сводные таблицы» пример с продажами",
        "добавь к теме «Excel» кейс",
    ],
)
def test_hours_and_new_topics_only_in_explicit_form(text: str) -> None:
    plan = parse_command(load(3), "T2.3", text)
    assert plan is None or plan.operation not in (PatchOperation.CHANGE_HOURS, PatchOperation.ADD_TOPIC)


def test_model_plan_is_checked_by_code() -> None:
    kit = load(3)
    instruction = "в теме 2.3 больше практики: 4 часа практики, лекций поменьше"
    plan = plan_from_draft(
        kit, "T2.3", instruction, PatchPlanDraft(operation="change_hours", quote="4 часа практики")
    )
    assert [(a.kind, a.value, a.sign) for a in plan.hours] == [(WorkKind.PR, 4.0, 0)]
    assert plan.planned_by == "model"
    with pytest.raises(PatchFailed, match="не нашлось часов"):  # чисел, которых нет в указании, не берём
        plan_from_draft(kit, "T2.3", instruction, PatchPlanDraft(operation="change_hours", quote="ЛК 9"))
    with pytest.raises(PatchFailed, match="Сформулируйте правку"):
        plan_from_draft(kit, "T2.3", instruction, PatchPlanDraft(operation="переделай"))
    moved = plan_from_draft(  # методист выбрал модуль, а заменить нужно тему — тема из плана модели
        kit,
        "M2",
        "замени тему 2.3",
        PatchPlanDraft(operation="replace_topic", target="тема 2.3", title="X Y Z"),
    )
    assert moved.target_id == "T2.3" and moved.title == "X Y Z"
    # знак из указания, если цитата без глагола; «добавь» и «убавь» сразу — просим написать часы явно
    more = plan_from_draft(
        kit,
        "T2.3",
        "дай теме 2.3 ещё 2 часа практики",
        PatchPlanDraft(operation="change_hours", quote="2 часа практики"),
    )
    assert [(a.value, a.sign) for a in more.hours] == [(2.0, 1)]
    with pytest.raises(PatchFailed, match="явно"):
        plan_from_draft(
            kit,
            "T2.3",
            "убавь лекции и добавь 2 часа практики",
            PatchPlanDraft(operation="change_hours", quote="2 часа практики"),
        )
    with pytest.raises(PatchFailed, match="только явной командой"):  # удаление — не по догадке модели
        plan_from_draft(kit, "T2.3", "тема 2.3 не нужна", PatchPlanDraft(operation="remove_topic"))


# --------------------------------------------------------------------------------------------------
# Правки в mock
# --------------------------------------------------------------------------------------------------


def test_replace_topic_changes_only_affected(pipeline: DefaultPipeline, tmp_path: Path) -> None:
    before = load(3)
    result = run_patch(pipeline, before, "T2.3", "замени тему 2.3 на «Сводные таблицы в pandas»", tmp_path)
    after = result.kit
    assert result.operation == "replace_topic"
    assert set(result.changed_ids) <= {
        "T2.3",
        "L2.3",
        "S6",
        *(q.id for q in before.questions if q.topic_id == "T2.3"),
    }
    assert "T2.3" in result.changed_ids and "L2.3" in result.changed_ids
    assert_only_changed(before, after, result.changed_ids)
    topic = next(t for t in after.topics if t.id == "T2.3")
    lecture = next(lec for lec in after.lectures if lec.id == "L2.3")
    assert topic.title == lecture.title == "Сводные таблицы в pandas"
    assert topic.hours == next(t for t in before.topics if t.id == "T2.3").hours
    assert lecture.slides  # у лекции были слайды — собраны заново
    assert after.total_hours() == before.total_hours() == after.program.total_hours
    assert after.revision == before.revision + 1 and len(after.history) == 1
    record = after.history[0]
    assert record.operation == PatchOperation.REPLACE_TOPIC and record.changed_ids == result.changed_ids
    assert record.planned_by == "rules" and not SERVICE_ID.search(record.summary)
    assert after.report is not None and after.report.status.value != "errors"


def test_harden_final_test_keeps_ids_and_outcomes(pipeline: DefaultPipeline, tmp_path: Path) -> None:
    before = load(1)
    result = run_patch(pipeline, before, "TS1", "усложни итоговый тест", tmp_path)
    after = result.kit
    test_ids = before.tests[0].question_ids
    assert result.operation == "harden_test" and set(result.changed_ids) <= set(test_ids)
    assert_only_changed(before, after, result.changed_ids)
    assert after.tests[0].question_ids == test_ids
    old = {q.id: q for q in before.questions}
    for question in after.questions:
        prev = old[question.id]
        assert (question.test_id, question.topic_id, question.lo_ids) == (
            prev.test_id,
            prev.topic_id,
            prev.lo_ids,
        )
        assert question.difficulty == min(3, prev.difficulty + 1) or question.difficulty > prev.difficulty
        assert question.stem != prev.stem


def test_change_hours_keeps_program_total(pipeline: DefaultPipeline, tmp_path: Path) -> None:
    before = load(3)
    result = run_patch(pipeline, before, "T2.3", "часы темы 2.3: ЛК 2, ПР 4", tmp_path)
    after = result.kit
    topic = next(t for t in after.topics if t.id == "T2.3")
    assert (topic.hours.lk, topic.hours.pr) == (2, 4)
    assert after.total_hours() == before.total_hours() == after.program.total_hours == 72
    assert after.module_hours("M2").total == before.module_hours("M2").total  # в пределах модуля
    others = [t.id for t in before.topics if t.module_id != "M2"]
    assert not set(result.changed_ids) & set(others)
    calendar = sum(day.hours for day in after.calendar.days)
    assert calendar == sum(day.hours for day in before.calendar.days) == after.total_hours()
    assert_only_changed(before, after, result.changed_ids)


def test_remove_and_add_topic_never_reuse_ids(pipeline: DefaultPipeline, tmp_path: Path) -> None:
    before = load(3)
    removed = run_patch(pipeline, before, "T2.2", "удали тему 2.2", tmp_path)
    kit = removed.kit
    assert "T2.2" in kit.retired_ids and not any(t.id == "T2.2" for t in kit.topics)
    assert {"L2.2", "S5"} <= set(kit.retired_ids) and kit.total_hours() == 72
    assert_only_changed(before, kit, removed.changed_ids, diff_ids(before, kit)[1])
    added = run_patch(pipeline, kit, "M2", "добавь тему «Очистка данных» в модуль 2", tmp_path)
    new = [t for t in added.kit.topics if t.id not in {x.id for x in kit.topics}]
    assert len(new) == 1 and new[0].id not in kit.retired_ids and new[0].title == "Очистка данных"
    assert any(lec.topic_ids == [new[0].id] for lec in added.kit.lectures)  # лекция к новой теме
    assert any(s.topic_id == new[0].id for s in added.kit.seminars)  # и задание
    assert added.kit.total_hours() == 72 and len(added.kit.history) == 2 and added.kit.revision == 2


def test_failed_patch_explains_and_keeps_kit(pipeline: DefaultPipeline, tmp_path: Path) -> None:
    before = load(3)
    snapshot = before.model_dump_json()
    with pytest.raises(PatchFailed, match="Часы не сходятся.*Комплект не изменён"):
        run_patch(pipeline, before, "T2.3", "часы темы 2.3: ЛК 40", tmp_path)
    with pytest.raises(PatchFailed, match="одна тема"):
        run_patch(pipeline, before, "T6.1", "удали тему 6.1", tmp_path)
    with pytest.raises(PatchFailed, match="нет тестовых вопросов"):
        run_patch(pipeline, before, "T2.2", "усложни вопросы темы 2.2", tmp_path)
    assert before.model_dump_json() == snapshot


def test_gateway_failure_is_explained(
    pipeline: DefaultPipeline, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    async def broke(*args, **kwargs):  # noqa: ANN002, ANN003, ANN202
        raise BudgetExceeded("лимит")

    monkeypatch.setattr(pipeline.content, "_draft_tests", broke)
    with pytest.raises(PatchFailed, match="бюджет.*Комплект не изменился"):
        run_patch(pipeline, load(1), "TS1", "усложни итоговый тест", tmp_path)


def test_report_shows_patch_history(pipeline: DefaultPipeline, tmp_path: Path) -> None:
    kit = run_patch(pipeline, load(3), "T2.3", "часы темы 2.3: ЛК 2, ПР 4", tmp_path).kit
    md = render_report_md(build_report(kit))
    assert "## История правок" in md and "часы темы 2.3: ЛК 2, ПР 4" in md and "версия комплекта — 1" in md
    docx_doc = build_report(kit, for_docx=True)
    history = next(s for s in docx_doc.sections if s.title == "История правок")
    assert "₽" not in history.blocks[1].header  # стоимость — не для программы на утверждение


# --------------------------------------------------------------------------------------------------
# Сквозной сценарий API: done → patch → done
# --------------------------------------------------------------------------------------------------


def test_api_done_patch_done(tmp_path: Path) -> None:
    settings = Settings(sessions_dir=tmp_path / "sessions", llm_mode="mock", cost_log=tmp_path / "cost.jsonl")
    query = OPEN[1]
    with TestClient(create_app(settings=settings, pipeline=DefaultPipeline(settings))) as client:
        done = client.post(
            "/api/v1/run", json={"request": query["request"], "answers": query.get("methodist_answers", {})}
        ).json()
        assert done["status"] == "done", done.get("message")
        sid = done["session_id"]
        docx_before = client.get(done["files"]["program_docx"]).content

        title = "Нейросети в помощь программисту"
        body = client.post(
            f"/api/v1/sessions/{sid}/patch",
            json={"target": "тема 1.2", "instruction": f"замени тему 1.2 на «{title}»"},
        )
        assert body.status_code == 200
        patched = body.json()
        assert patched["status"] == "done" and patched["patch"]["ok"], patched.get("message")
        assert patched["patch"]["revision"] == 1 and patched["patch"]["operation"] == "replace_topic"
        assert "T1.2" in patched["changed_ids"] and patched["kit"]["revision"] == 1
        assert patched["message"].startswith("Правка выполнена (версия 1)")
        docx_after = client.get(patched["files"]["program_docx"]).content
        assert docx_after != docx_before
        text = "\n".join(p.text for p in docx.Document(io.BytesIO(docx_after)).paragraphs)
        tables = " ".join(
            c.text for t in docx.Document(io.BytesIO(docx_after)).tables for r in t.rows for c in r.cells
        )
        assert title in text + tables
        page = client.get(f"/sessions/{sid}/report").text
        assert "История правок" in page and title in page

        failed = client.post(
            f"/api/v1/sessions/{sid}/patch",
            json={"target": "тема 1.2", "instruction": "часы темы 1.2: ЛК 40"},
        ).json()
        assert failed["status"] == "done" and failed["patch"]["ok"] is False
        assert "Часы не сходятся" in failed["patch"]["error"] and failed["kit"]["revision"] == 1
        assert failed["message"].startswith("Часы не сходятся") or failed["message"].startswith("Правка не")

        missing = client.post(
            f"/api/v1/sessions/{sid}/patch", json={"target": "тема 9.9", "instruction": "x"}
        )
        assert missing.status_code == 422 and "9.9" in missing.json()["error"]["message"]
        assert client.get(f"/api/v1/sessions/{sid}").json()["kit"]["revision"] == 1


def test_shared_lecture_is_rebuilt_for_all_its_topics(pipeline: DefaultPipeline, tmp_path: Path) -> None:
    before = load(3)
    shared = next(lec for lec in before.lectures if lec.id == "L2.1")
    shared.topic_ids = ["T2.1", "T2.2"]  # одна лекция на две темы, как бывает у живой модели
    before.lectures = [lec for lec in before.lectures if lec.id != "L2.2"]
    old_plan = len(shared.plan)
    result = run_patch(pipeline, before, "T2.2", "замени тему 2.2 на «Дашборды для руководителя»", tmp_path)
    lecture = next(lec for lec in result.kit.lectures if lec.id == "L2.1")
    assert lecture.topic_ids == ["T2.1", "T2.2"] and "L2.1" in result.changed_ids
    assert lecture.plan and len(lecture.plan) >= min(old_plan, 2)
    assert not any(lec.id not in {x.id for x in before.lectures} for lec in result.kit.lectures)  # лишних нет
    assert_only_changed(before, result.kit, result.changed_ids)


def test_topic_patch_keeps_questions_the_model_repeated(
    pipeline: DefaultPipeline, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Живой прогон 21.10: модель вернула те же вопросы — правка темы проходит, вопросы прежние, есть заметка;
    «усложни тест» в таком случае не применяется."""
    before = load(3)
    own = [q for q in before.questions if q.topic_id == "T2.3"]
    repeated = TestsDraft(
        questions=[
            QuestionDraft(
                topic_id=q.topic_id,
                lo_ids=q.lo_ids,
                stem=q.stem,
                options=q.options,
                correct_answers=[q.options[i] for i in q.correct],
            )
            for q in own
        ]
    )

    async def same(*args, **kwargs):  # noqa: ANN002, ANN003, ANN202
        return repeated

    monkeypatch.setattr(pipeline.content, "_draft_tests", same)
    result = run_patch(pipeline, before, "T2.3", "пересобери тему 2.3", tmp_path)
    assert not {q.id for q in own} & set(result.changed_ids) and "T2.3" in result.changed_ids
    assert any("оставлены прежними" in n for n in result.notes)
    assert "оставлены прежними" in result.kit.history[-1].summary
    with pytest.raises(PatchFailed, match="повторяют прежние"):
        run_patch(pipeline, before, "TS2", "усложни тест 2", tmp_path)


def test_mentioned_object_must_be_the_target(pipeline: DefaultPipeline, tmp_path: Path) -> None:
    kit = load(3)
    for text in ("удали тему 2.2", "часы темы 2.2: ЛК 2, ПР 4", "замени тему 2.2 на «Новая»"):
        with pytest.raises(PatchFailed, match="тема 2.2"):
            run_patch(pipeline, kit, "T2.3", text, tmp_path)
    with pytest.raises(PatchFailed, match="нет объекта"):
        run_patch(pipeline, kit, "T2.3", "удали тему 2.9", tmp_path)
    # в кавычках — новое название, а не ссылка на другую тему
    result = run_patch(pipeline, kit, "T2.3", "замени тему 2.3 на «Тема 3.1 из старой программы»", tmp_path)
    assert result.operation == "replace_topic"


def test_patch_with_new_errors_is_rejected(pipeline: DefaultPipeline, tmp_path: Path) -> None:
    kit = load(3)
    with pytest.raises(PatchFailed, match="автопроверка нашла ошибки.*Комплект прежний"):
        run_patch(pipeline, kit, "T2.3", "тема 2.3 на 0 часов", tmp_path)


def test_new_titles_and_outcomes_get_unverified_marks(pipeline: DefaultPipeline, tmp_path: Path) -> None:
    kit = load(3)
    text = "замени результат 1.1 на «Умеет применять ПС 06.999 и приказ № 9999 в работе»"
    result = run_patch(pipeline, kit, "LO1.1", text, tmp_path)
    outcome = next(o for o in result.kit.outcomes if o.id == "LO1.1")
    assert "[не подтверждено]" in outcome.text and result.operation == "replace_lo"


def test_free_text_goes_through_the_skill(pipeline: DefaultPipeline, tmp_path: Path) -> None:
    result = run_patch(pipeline, load(3), "T2.3", "сделай задания ближе к работе кассира", tmp_path)
    record = result.kit.history[-1]
    assert record.planned_by == "model" and record.operation == PatchOperation.REGENERATE
    assert "T2.3" in result.changed_ids
