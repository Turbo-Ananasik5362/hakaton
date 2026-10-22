"""Программа в .docx (app/render): снимок таблиц, оформление по спецификациям шаблонов, V31 и выдача через API."""

from __future__ import annotations

from pathlib import Path

import docx
import pytest
from docx.oxml.ns import qn

from app.config import Settings, now_msk
from app.core.copycheck import copy_ratio
from app.core.hours import fmt_hours
from app.core.registry import default_registry
from app.models import (
    Kit,
    ProgramType,
    Severity,
    Source,
    SourceKind,
    SourceSection,
    StudyFormKind,
    Verification,
)
from app.pipeline.base import StepContext
from app.pipeline.default import DefaultPipeline
from app.render import docx_dop, docx_pk
from app.render.bundle import FAILED_NOTE
from app.render.common import BLANK, FILL_MARK, load_organization
from app.render.program import DOCX_NAME, render_program_docx
from app.render.tables import build_tables
from app.render.verify import verify_docx

FIXTURES = Path(__file__).parent / "fixtures" / "kits"
ROOT = Path(__file__).resolve().parent.parent
PK_TEXT_TW = 9779
DOP_PORTRAIT_TW = 9355
DOP_LANDSCAPE_TW = 14570


def load(n: int) -> Kit:
    return Kit.model_validate_json((FIXTURES / f"open_q{n}.json").read_text(encoding="utf-8"))


def render(kit: Kit, tmp_path: Path) -> tuple[Path, list]:
    org = load_organization(Settings().organization_file)
    name, issues = render_program_docx(kit, tmp_path, default_registry(), org)
    return tmp_path / name, issues


def grid_width(table) -> int:  # noqa: ANN001
    return sum(int(c.get(qn("w:w"))) for c in table._tbl.tblGrid.findall(qn("w:gridCol")))


def headings(document) -> list[str]:  # noqa: ANN001
    return [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]


# --------------------------------------------------------------------------------------------------
# Снимок таблиц
# --------------------------------------------------------------------------------------------------


@pytest.mark.parametrize("n", [1, 2, 3, 4, 5])
def test_tables_take_numbers_from_kit(n: int) -> None:
    kit = load(n)
    tables = build_tables(kit)
    modules = [r for r in tables.plan if r.kind == "module"]
    assert [r.hours.total for r in modules] == [kit.module_hours(m.id).total for m in kit.modules_sorted()]
    total = tables.plan[-1]
    assert total.kind == "total" and total.hours.total == kit.total_hours() == kit.program.total_hours
    assert tables.plan[-2].hours.total == kit.fa_hours()
    for section in tables.sections:
        for topic in section.topics:
            source = next(t for t in kit.topics if t.id == topic.topic_id)
            assert sum(b.hours for b in topic.blocks) == source.hours.total
            for block in topic.blocks:  # метка — из того же числа, что колонка
                assert block.hours == source.hours.get(block.kind)
                assert (
                    block.label.split("(" if kit.program.type != ProgramType.DOP else "–")[1]
                    .strip()
                    .startswith(f"{block.hours:g}".replace(".", ","))
                )
    assert sum(r.hours for r in tables.calendar) == sum(d.hours for d in kit.calendar.days)


# --------------------------------------------------------------------------------------------------
# Документы по фикстурам открытых запросов
# --------------------------------------------------------------------------------------------------


@pytest.mark.parametrize("n", [1, 2, 3, 4, 5])
def test_rendered_program_passes_v31(n: int, tmp_path: Path) -> None:
    path, issues = render(load(n), tmp_path)
    assert path.is_file() and issues == []


@pytest.mark.parametrize("n", [2, 3, 5])
def test_pk_layout_follows_spec(n: int, tmp_path: Path) -> None:
    kit = load(n)
    path, _ = render(kit, tmp_path)
    document = docx.Document(str(path))
    assert headings(document) == docx_pk.HEADINGS
    first = next(p for p in document.paragraphs if p.text == docx_pk.H1)
    assert first.style.name == "Heading 1"  # в образце раздел 1 — Normal (дефект)
    assert all(grid_width(t) <= PK_TEXT_TW for t in document.tables)
    for table in document.tables:  # шапка повторяется на каждой странице
        assert table._tbl.tr_lst[0].trPr.find(qn("w:tblHeader")) is not None
    texts = [p.text for p in document.paragraphs]
    order = next(t for t in texts if t.startswith("Приказом"))
    assert BLANK in order and "17-ОБР" not in order
    stazh = "ПОВЫШЕНИЯ КВАЛИФИКАЦИИ В ФОРМЕ СТАЖИРОВКИ"
    assert (stazh in texts) == (kit.program.type == ProgramType.PK_INTERNSHIP)
    assert not any("ЛР (" in t for t in texts)
    assert not document.core_properties.last_modified_by


def test_pk_goal_lists_labor_functions_from_registry(tmp_path: Path) -> None:
    kit = load(3)
    registry = default_registry()
    path, _ = render(kit, tmp_path)
    texts = [p.text for p in docx.Document(str(path)).paragraphs]
    goal = texts.index(docx_pk.SUB["1.2"])
    section = texts[goal : texts.index(docx_pk.SUB["1.3"])]
    for comp in kit.competencies:
        for ref in comp.tf_refs:
            function = registry.labor_function(ref.ps_code, ref.tf_code)
            assert any(
                t.startswith(f"{ref.tf_code} ") and function.title[:30].lower() in t.lower() for t in section
            )


def test_pk_with_distance_learning_lists_decree_and_letter(tmp_path: Path) -> None:
    kit = load(3)
    assert kit.program.form.dot and "letter_vk1013" in kit.program.norm_ids
    path, _ = render(kit, tmp_path)
    texts = [p.text for p in docx.Document(str(path)).paragraphs]
    norms = texts[texts.index(docx_pk.SUB["3.2.1"]) : texts.index(docx_pk.SUB["3.2.2"])]
    assert any("№ 1678" in t for t in norms) and any("ВК-1013/06" in t for t in norms)
    assert not any("№ 816" in t or "№ 499" in t for t in norms)


def test_internship_has_no_contact_hours_in_plan(tmp_path: Path) -> None:
    kit = load(5)
    path, _ = render(kit, tmp_path)
    document = docx.Document(str(path))
    plan = next(t for t in document.tables if t.rows[0].cells[1].text == "Наименование разделов программы")
    for row in plan.rows[1:-2]:  # разделы: ЛК = ПР = ЛБ = 0
        assert [row.cells[c].text for c in (3, 4, 5)] == ["0", "0", "0"]


def test_dop_layout_follows_spec(tmp_path: Path) -> None:
    kit = load(1)
    path, _ = render(kit, tmp_path)
    document = docx.Document(str(path))
    assert headings(document) == docx_dop.HEADINGS
    orientations = [s.orientation for s in document.sections]
    assert len(orientations) == 3 and orientations[1] != orientations[0] == orientations[2]
    fonts = document.styles["Normal"].element.rPr.rFonts
    assert all(
        fonts.get(qn(f"w:{slot}")) == "Times New Roman" for slot in ("ascii", "hAnsi", "eastAsia", "cs")
    )
    plan = document.tables[2]  # сводка графика, график по неделям, учебный план
    assert [c.text for c in plan.rows[1].cells][4:6] == ["Лекции / теория", "Практические занятия"]
    assert plan.rows[0].cells[0]._tc is plan.rows[1].cells[0]._tc  # вертикальное объединение шапки
    for idx, table in enumerate(document.tables):  # учебный план и УТП (2, 3) — в альбомной секции
        assert grid_width(table) <= (DOP_LANDSCAPE_TW if idx in (2, 3) else DOP_PORTRAIT_TW)
    texts = [p.text for p in document.paragraphs]
    year = (kit.meta.created_at or now_msk()).year
    assert f"Магас {year}" in texts and f"{BLANK} {BLANK}" not in texts  # город и год, как в образце и в ПК
    assert any("273-ФЗ" in t for t in texts) and any("№ 629" in t for t in texts)
    assert any(t.rows[0].cells[0].text == "Дата начала обучения" for t in document.tables)
    staffing = next(t for t in document.tables if "ФИО преподавателя" in t.rows[0].cells[2].text)
    assert all(r.cells[2].text == BLANK for r in staffing.rows[1:])
    assert not any("профессиональный стандарт" in t.lower() for t in texts)


def test_missing_texts_get_fill_mark_and_unverified_sources_are_marked(tmp_path: Path) -> None:
    kit = load(3)
    kit.texts.annotation = ""
    kit.texts.org_process = ""
    kit.sources = [s for s in kit.sources if s.section != SourceSection.MAIN] + [
        Source(
            id="SRC99",
            kind=SourceKind.BOOK,
            section=SourceSection.MAIN,
            title="Непроверенная книга",
            verification=Verification.UNVERIFIED,
        )
    ]
    path, issues = render(kit, tmp_path)
    texts = [p.text for p in docx.Document(str(path)).paragraphs]
    annotation = texts.index(docx_pk.SUB["1.1"])
    assert texts[annotation + 1] == FILL_MARK
    assert texts[texts.index(docx_pk.SUB["3.3"]) + 1] == FILL_MARK
    assert any("Непроверенная книга" in t and "[не подтверждено]" in t for t in texts)
    assert issues == []


def test_typical_paragraphs_are_not_copied_from_samples(tmp_path: Path) -> None:
    samples = [
        (ROOT / "docs" / "case" / name).read_text(encoding="utf-8")
        for name in ("sample_pk_stazhirovka_bsa_248h.md", "sample_dop_intensiv_ai_24h.md")
    ]
    for n in (1, 3):
        path, _ = render(load(n), tmp_path / str(n))
        text = "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
        assert copy_ratio(text, samples) < 0.05


# --------------------------------------------------------------------------------------------------
# V31 ловит испорченный документ
# --------------------------------------------------------------------------------------------------


def _tamper(path: Path, edit) -> None:  # noqa: ANN001
    document = docx.Document(str(path))
    edit(document)
    document.save(str(path))


def _set(cell, text: str) -> None:  # noqa: ANN001
    cell.paragraphs[0].runs[0].text = text


def test_v31_catches_wrong_number_in_plan(tmp_path: Path) -> None:
    kit = load(3)
    path, _ = render(kit, tmp_path)
    _tamper(path, lambda d: _set(d.tables[1].rows[1].cells[3], "99"))
    issues = verify_docx(path, kit, docx_pk.HEADINGS)
    assert any(i.severity == Severity.ERROR and "Учебный план" in i.message for i in issues)


def test_v31_catches_label_differing_from_column(tmp_path: Path) -> None:
    kit = load(3)
    path, _ = render(kit, tmp_path)
    wp_index = 4  # компетенции, УП, УТП, график, рабочая программа

    def edit(document) -> None:  # noqa: ANN001
        row = document.tables[wp_index].rows[2]
        row.cells[1].paragraphs[0].runs[0].text = "ПР (99 часов):"

    _tamper(path, edit)
    issues = verify_docx(path, kit, docx_pk.HEADINGS)
    assert any("метка" in i.message for i in issues)


def test_v31_catches_filled_requisites_and_missing_section(tmp_path: Path) -> None:
    kit = load(3)
    path, _ = render(kit, tmp_path)

    def edit(document) -> None:  # noqa: ANN001
        order = next(p for p in document.paragraphs if p.text.startswith("Приказом"))
        order.runs[0].text = "Приказом № 17-ОБР от «31» июля 2026 г."
        heading = next(p for p in document.paragraphs if p.text == docx_pk.SUB["2.3"])
        heading._p.getparent().remove(heading._p)

    _tamper(path, edit)
    messages = [i.message for i in verify_docx(path, kit, docx_pk.HEADINGS)]
    assert any("Реквизиты приказа" in m for m in messages)
    assert any("нет разделов шаблона" in m and "2.3." in m for m in messages)


def test_v31_catches_dop_topic_hours(tmp_path: Path) -> None:
    kit = load(1)
    path, _ = render(kit, tmp_path)

    def edit(document) -> None:  # noqa: ANN001
        utp = next(t for t in document.tables if t.rows[0].cells[1].text == "Виды учебных занятий")
        _set(utp.rows[2].cells[1], "Теоретическое занятие – 5 часов")

    _tamper(path, edit)
    assert any("Учебно-тематический план" in i.message for i in verify_docx(path, kit, docx_dop.HEADINGS))


# --------------------------------------------------------------------------------------------------
# Пайплайн и API
# --------------------------------------------------------------------------------------------------


def _ctx(tmp_path: Path) -> StepContext:
    return StepContext(session_id="s_000000000001", inputs_dir=tmp_path, files_dir=tmp_path, request=None)  # type: ignore[arg-type]


def test_render_failure_keeps_json_and_explains(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    def broken(*_args, **_kwargs):  # noqa: ANN002, ANN003, ANN202
        raise RuntimeError("шаблон повреждён")

    monkeypatch.setattr("app.render.bundle.render_program_docx", broken)
    kit = load(1)
    files = DefaultPipeline(Settings(llm_mode="mock")).render(kit, _ctx(tmp_path))
    assert "program_docx" not in files and files["kit_json"] == "kit.json"
    assert not (tmp_path / DOCX_NAME).exists()
    note = FAILED_NOTE.format(title="Программа .docx", error="").split("(")[0]
    assert any(n.startswith(note) and "RuntimeError" in n for n in kit.meta.generation_notes)
    assert not any(
        "шаблон повреждён" in n for n in kit.meta.generation_notes
    )  # текст исключения — не методисту
    assert any("шаблон повреждён" in i.detail for i in kit.report.issues if i.check == "V31")
    assert any(i.check == "V31" and i.severity == Severity.ERROR for i in kit.report.issues)
    saved = Kit.model_validate_json((tmp_path / "kit.json").read_text(encoding="utf-8"))
    assert any(i.check == "V31" for i in saved.report.issues)


def test_render_replaces_previous_v31_issues(tmp_path: Path) -> None:
    kit = load(4)
    pipeline = DefaultPipeline(Settings(llm_mode="mock"))
    pipeline.render(kit, _ctx(tmp_path))
    pipeline.render(kit, _ctx(tmp_path))
    assert [i for i in kit.report.issues if i.check == "V31"] == []


def test_api_returns_program_docx(tmp_path: Path) -> None:
    from fastapi.testclient import TestClient

    from app.main import create_app
    from tests.intake_queries import OPEN
    from tests.test_content import run

    body = run(tmp_path, OPEN[3])
    assert body["files"]["program_docx"].endswith("/files/program.docx")
    settings = Settings(sessions_dir=tmp_path / "sessions", llm_mode="mock")
    with TestClient(create_app(settings=settings)) as client:
        response = client.get(body["files"]["program_docx"])
    assert response.status_code == 200
    document = docx.Document(__import__("io").BytesIO(response.content))
    assert headings(document) == docx_pk.HEADINGS
    assert not [i for i in body["report"]["issues"] if i["check"] == "V31"]


def test_form_kind_text_is_rendered(tmp_path: Path) -> None:
    kit = load(4)
    kit.program.form.kind = StudyFormKind.ZAOCHNAYA
    path, _ = render(kit, tmp_path)
    texts = [p.text for p in docx.Document(str(path)).paragraphs]
    assert any(t.startswith("Форма обучения — заочная") for t in texts)


def test_dop_assessment_lists_every_seminar_criterion(tmp_path: Path) -> None:
    """2.3 ДОП: критерии всех практических работ, без обрезки списка (обрезанный выглядел незаполненным)."""
    kit = load(1)
    for seminar in kit.seminars:  # много критериев: прежний вывод обрезал на восьми
        seminar.criteria = [*seminar.criteria, *seminar.criteria[:1]]
    lines = docx_dop.seminar_criteria(kit)
    assert len(lines) == len([s for s in kit.seminars if s.criteria])
    every = {c.text.strip().rstrip(".") for s in kit.seminars for c in s.criteria}
    assert all(any(text in line for line in lines) for text in every)
    path, _ = render(kit, tmp_path)
    texts = [p.text for p in docx.Document(str(path)).paragraphs]
    assert all(line in texts for line in lines)


def test_pk_with_distance_learning_splits_contact_hours(tmp_path: Path) -> None:
    """Очно-заочная программа с ДОТ: под учебным планом и УТП — очные часы и часы с ДОТ по разделам и темам
    (оценка 21.10, п. 2.2: разбивка была одной цифрой только в 1.8)."""
    from app.render.docx_pk import distance_split_lines

    kit = load(3)
    tables = build_tables(kit)
    plan_lines = distance_split_lines(tables.plan)
    assert plan_lines and plan_lines[-1].startswith("Всего контактной работы")
    volume = tables.volume
    assert f"с применением ДОТ — {fmt_hours(volume.distance)} ч" in plan_lines[-1]
    path, issues = render(kit, tmp_path)
    texts = [p.text for p in docx.Document(str(path)).paragraphs]
    assert all(line in texts for line in plan_lines)
    assert not [i for i in issues if i.severity.value == "error"]
