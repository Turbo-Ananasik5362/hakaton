from __future__ import annotations

import copy
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from pypdf import PdfWriter

from app.core.reference import (
    extract_text,
    find_order_refs,
    find_ps_codes,
    find_tf_codes,
    match_label,
    number_key,
    parse_hours,
    parse_reference,
    reference_outline,
    requisites_filled,
)
from app.models import ParsedReference, ProgramType, RefHoursTable, WorkKind

TEMPLATES = Path(__file__).resolve().parents[1] / "data" / "templates"
PK_SAMPLE = TEMPLATES / "pk_stazhirovka_bsa_248h.docx"
DOP_SAMPLE = TEMPLATES / "dop_intensiv_ai_24h.docx"


def table_of(ref: ParsedReference, kind: str) -> RefHoursTable:
    return next(table for table in ref.hour_tables if table.kind == kind)


def row_values(table: RefHoursTable, row_type: str) -> list[float | None]:
    row = next(r for r in table.rows if r.row_type == row_type)
    return [row.values[key] for key in table.columns]


# --------------------------------------------------------------------------------------------------
# Образец ПК (стажировка, 248 ч)
# --------------------------------------------------------------------------------------------------


@pytest.fixture(scope="module")
def pk() -> ParsedReference:
    return parse_reference(PK_SAMPLE)


def test_pk_type_and_source(pk: ParsedReference) -> None:
    assert pk.source_name == PK_SAMPLE.name
    assert pk.program_type_guess is ProgramType.PK_INTERNSHIP


def test_pk_headings(pk: ParsedReference) -> None:
    titles = [heading.title for heading in pk.headings]
    assert "Учебный план" in titles
    assert "Учебно-тематический план" in titles
    assert any(title.startswith("Рабочая программа") for title in titles)
    first = pk.headings[0]  # раздел 1 набран жирным Normal, а не стилем заголовка
    assert (first.number, first.level) == ("1", 1)
    assert first.title.startswith("ОБЩАЯ ХАРАКТЕРИСТИКА")
    by_number = {heading.number: heading for heading in pk.headings if heading.number}
    assert by_number["2.5.1"].level == 3
    assert "Приложение 1" in titles
    assert not any(
        "Письмо" in title or "Опишите" in title for title in titles
    )  # пункты списков — не заголовки
    outline = reference_outline(pk)
    assert "2.1. Учебный план" in outline and "    2.5.1. " in outline


def test_pk_plan_numbers_as_in_document(pk: ParsedReference) -> None:
    plan = table_of(pk, "plan")
    assert plan.columns == ["total", "lk", "pr", "lb", "srs"]
    assert "Форма аттестации" in plan.raw_header
    section = next(row for row in plan.rows if row.row_type == "section")
    assert section.number == "1."
    assert [section.values[key] for key in plan.columns] == [246, 0, 0, 0, 246]
    assert row_values(plan, "total") == [248, None, 66, 6, 174]  # ошибка образца сохранена
    fa = next(row for row in plan.rows if row.row_type == "fa")
    assert fa.values["total"] == 2 and fa.form == "зачет"


def test_pk_utp_keeps_sample_errors(pk: ParsedReference) -> None:
    utp = table_of(pk, "utp")
    topics = [row for row in utp.rows if row.row_type == "topic"]
    assert len(topics) == 10
    assert sum(row.values["srs"] or 0 for row in topics) == 168
    assert row_values(utp, "total") == [248, 0, 66, 6, 174]
    control = next(row for row in utp.rows if row.row_type == "control")
    assert number_key(control.number) == "1.11" and control.title == "Текущий контроль"
    assert [control.values[key] for key in utp.columns] == [6, 0, 0, 0, 0]
    topic_13 = next(row for row in topics if number_key(row.number) == "1.3")
    assert [topic_13.values[key] for key in utp.columns] == [30, 0, 9, 6, 15]


def test_pk_workprogram_labels_vs_column(pk: ParsedReference) -> None:
    items = pk.workprogram

    def pr_of(topic: str) -> tuple[float | None, float | None]:
        item = next(i for i in items if i.topic_number == topic and i.kind is WorkKind.PR)
        return item.label_hours, item.column_hours

    assert pr_of("1.6") == (9, 12)
    assert pr_of("1.7") == (6, 9)
    assert pr_of("1.2") == (9, 9)
    lab = next(i for i in items if i.topic_number == "1.3" and i.kind is WorkKind.LB)
    assert lab.kind_label == "ЛР" and lab.label_hours == 6 and lab.column_hours == 6
    srs_15 = [i for i in items if i.topic_number == "1.5" and i.kind is WorkKind.SRS]
    assert [(i.label_hours, i.column_hours) for i in srs_15] == [(12, 12), (3, 3)]  # «СРС (12часов)»
    controls = [i for i in items if i.kind_label == "Текущий контроль"]
    assert [(i.kind, i.label_hours, i.column_hours) for i in controls] == [
        (None, 3, 3),
        (None, 3, 3),
        (None, None, 6),
    ]
    fa = next(i for i in items if i.kind_label == "Итоговая аттестация")
    assert fa.column_hours == 2 and fa.topic_number is None
    assert next(i for i in items if i.topic_number == "1.1").topic_title == "Введение в профессию аналитика"


def test_pk_labels_requisites_and_codes(pk: ParsedReference) -> None:
    assert {"ЛР", "ЛБ", "ПР", "СРС"} <= set(pk.work_kind_labels)
    assert pk.order_requisites is not None and "17-ОБР" in pk.order_requisites
    assert pk.requisites_filled is True
    assert "06.015" in pk.ps_codes
    assert "C/14.6" in pk.tf_codes
    refs = " | ".join(pk.order_refs)
    assert "ВК-1013/06" in refs and "№ 273-ФЗ" in refs and "№ 816" in refs and "№ 586н" in refs
    assert "17-ОБР" not in refs  # реквизиты утверждения — не нормативная ссылка


# --------------------------------------------------------------------------------------------------
# Образец ДОП (интенсив, 24 ч)
# --------------------------------------------------------------------------------------------------


@pytest.fixture(scope="module")
def dop() -> ParsedReference:
    return parse_reference(DOP_SAMPLE.read_bytes(), filename=DOP_SAMPLE.name)


def test_dop_type_and_headings(dop: ParsedReference) -> None:
    assert dop.source_name == DOP_SAMPLE.name
    assert dop.program_type_guess is ProgramType.DOP
    first = dop.headings[0]  # «1.» — автонумерация списка
    assert (first.number, first.title) == ("1", "Пояснительная записка")
    titles = [heading.title for heading in dop.headings]
    assert {"Учебный план", "Учебно-тематический план", "Приложение № 1"} <= set(titles)
    assert "Задачи Программы" not in titles


def test_dop_plan_two_level_header(dop: ParsedReference) -> None:
    plan = table_of(dop, "dop_plan")
    assert plan.columns == ["total", "contact", "lk", "pr", "srs"]
    assert "Лекции/теория" in plan.raw_header and "Практические занятия" in plan.raw_header
    assert row_values(plan, "total") == [24, 8, 4, 4, 10]
    assert row_values(plan, "fa") == [6, None, None, None, None]
    intro = plan.rows[0]
    assert intro.title == "Вводное занятие" and intro.number is None
    assert [intro.values[key] for key in plan.columns] == [2, 1, 0.5, 0.5, 1]
    assert [row.number for row in plan.rows if row.row_type == "section"] == [None, "1", "2", "3"]
    assert intro.form is None  # прочерк в колонке формы


def test_dop_requisites_codes_and_utp(dop: ParsedReference) -> None:
    assert dop.order_requisites is not None and "____" in dop.order_requisites
    assert dop.requisites_filled is False
    assert dop.ps_codes == [] and dop.tf_codes == []
    intro = [(i.kind, i.label_hours) for i in dop.workprogram if i.topic_number == "0.1"]
    assert intro == [(WorkKind.LK, 0.5), (WorkKind.PR, 0.5)]
    fa = next(i for i in dop.workprogram if i.kind_label == "Итоговая аттестация")
    assert fa.label_hours == 6


def test_dop_sdt_text_is_extracted() -> None:
    text = extract_text(DOP_SAMPLE)
    assert "проводится очно в формате хакатона" in text  # абзац целиком внутри w:sdt
    assert "Репетиция выступления (питча)" in text  # ячейка таблицы с w:sdt


# --------------------------------------------------------------------------------------------------
# Синтетические документы: универсальность разбора
# --------------------------------------------------------------------------------------------------


def _add_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.cell(r, c).text = value


def test_generic_pk_program(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("ДОПОЛНИТЕЛЬНАЯ ПРОФЕССИОНАЛЬНАЯ ПРОГРАММА")
    document.add_paragraph("ПОВЫШЕНИЯ КВАЛИФИКАЦИИ")
    document.add_paragraph("Приказом № ____ от «__» ________ 20__ г.")
    document.add_heading("2.1. Учебный план", level=2)
    _add_table(
        document,
        [
            [
                "№",
                "Наименование модулей",
                "Всего, час.",
                "Лекции",
                "Практические занятия",
                "СРС",
                "Форма контроля",
            ],
            ["1", "2", "3", "4", "5", "6", "7"],
            ["1.", "Основы", "10", "4", "4,5", "1,5", "Тест"],
            ["", "Итоговая аттестация", "2", "—", "-", "", "Зачёт"],
            ["", "Итого", "12", "4", "4,5", "1,5", ""],
        ],
    )
    path = tmp_path / "program.docx"
    document.save(path)

    ref = parse_reference(path)

    assert ref.program_type_guess is ProgramType.PK
    assert ref.headings[0].number == "2.1" and ref.headings[0].level == 2
    plan = table_of(ref, "plan")
    assert plan.columns == ["total", "lk", "pr", "srs"]
    assert [row.row_type for row in plan.rows] == ["section", "fa", "total"]  # строка «1 2 3 …» пропущена
    assert row_values(plan, "section") == [10, 4, 4.5, 1.5]
    assert row_values(plan, "fa") == [2, None, None, None]
    assert ref.requisites_filled is False


@pytest.mark.parametrize(
    ("title", "expected_kind"),
    [
        ("ДОПОЛНИТЕЛЬНАЯ ПРОФЕССИОНАЛЬНАЯ ПРОГРАММА ПОВЫШЕНИЯ КВАЛИФИКАЦИИ", "plan"),
        ("ДОПОЛНИТЕЛЬНАЯ ОБЩЕРАЗВИВАЮЩАЯ ПРОГРАММА", "dop_plan"),
        ("Программа курса", "dop_plan"),  # тип неизвестен — по раскладке колонок
    ],
)
def test_contact_column_kind_depends_on_program_type(tmp_path: Path, title: str, expected_kind: str) -> None:
    document = Document()
    document.add_paragraph(title)
    document.add_heading("2. Учебный план", level=1)
    _add_table(
        document,
        [
            ["№", "Наименование", "Общая трудоемкость", "Всего аудиторных", "Лекции", "Практика", "СР"],
            ["1", "Модуль", "8", "6", "2", "4", "2"],
            ["", "Итого", "8", "6", "2", "4", "2"],
        ],
    )
    path = tmp_path / "plan.docx"
    document.save(path)

    plan = parse_reference(path).hour_tables[0]

    assert plan.kind == expected_kind
    assert plan.columns == ["total", "contact", "lk", "pr", "srs"]
    assert row_values(plan, "total") == [8, 6, 2, 4, 2]


def test_sdt_wrapped_text_is_parsed(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    sdt, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
    run = paragraph.add_run("Профессиональный стандарт 06.001, ТФ D/03.6")._r
    content.append(copy.deepcopy(run))
    sdt.append(content)
    paragraph._p.remove(run)
    paragraph._p.append(sdt)
    path = tmp_path / "sdt.docx"
    document.save(path)

    assert document.paragraphs[0].text == ""  # python-docx текст в w:sdt не видит
    ref = parse_reference(path)
    assert ref.ps_codes == ["06.001"] and ref.tf_codes == ["D/03.6"]


def test_txt_reference_cp1251() -> None:
    text = "\n".join(
        [
            "ДОПОЛНИТЕЛЬНАЯ ОБЩЕОБРАЗОВАТЕЛЬНАЯ ОБЩЕРАЗВИВАЮЩАЯ ПРОГРАММА",
            "Приказом № ____ от «___» ______ 2026 г.",
            "1. ПОЯСНИТЕЛЬНАЯ ЗАПИСКА",
            "1.1. Актуальность программы",
            "1. Федеральный закон от 29.12.2012 № 273-ФЗ «Об образовании в Российской Федерации».",
            "ПР (4 часа): практикум по прототипированию",
        ]
    )
    ref = parse_reference(text.encode("cp1251"), filename="запрос.txt")
    assert ref.source_name == "запрос.txt"
    assert ref.program_type_guess is ProgramType.DOP
    assert [(h.number, h.title) for h in ref.headings] == [
        ("1", "ПОЯСНИТЕЛЬНАЯ ЗАПИСКА"),
        ("1.1", "Актуальность программы"),
    ]
    assert ref.requisites_filled is False
    assert ref.work_kind_labels == ["ПР"]
    assert ref.order_refs == ["Федеральный закон от 29.12.2012 № 273-ФЗ"]
    assert ref.hour_tables == [] and ref.workprogram == []


def test_empty_and_broken_pdf_do_not_fail(tmp_path: Path) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    empty_pdf = tmp_path / "empty.pdf"
    with empty_pdf.open("wb") as handle:
        writer.write(handle)

    empty = parse_reference(empty_pdf)
    broken = parse_reference(b"%PDF-1.4\n%%%% oops", filename="broken.pdf")

    assert empty.source_name == "empty.pdf" and empty.headings == [] and empty.program_type_guess is None
    assert broken.source_name == "broken.pdf" and broken.hour_tables == []


def test_unreadable_inputs_do_not_fail(tmp_path: Path) -> None:
    assert parse_reference(b"PK\x03\x04 not a zip", filename="fake.docx").source_name == "fake.docx"
    assert parse_reference(tmp_path / "missing.docx").source_name == "missing.docx"
    assert parse_reference(b"", filename="empty.txt").headings == []
    assert extract_text(b"\x00\x01", filename="x.docx") == ""


# --------------------------------------------------------------------------------------------------
# Помощники
# --------------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("246", 246.0),
        ("0,5", 0.5),
        (" 1.5 ", 1.5),
        ("12 ч", 12.0),
        ("", None),
        ("—", None),
        ("-", None),
        ("зачет", None),
    ],
)
def test_parse_hours(text: str, expected: float | None) -> None:
    assert parse_hours(text) == expected


@pytest.mark.parametrize(
    ("paragraph", "label", "kind", "hours"),
    [
        ("СРС (15 часов):", "СРС", WorkKind.SRS, 15),
        ("СРС (12часов)", "СРС", WorkKind.SRS, 12),
        ("СРС (3 часов): Практикум", "СРС", WorkKind.SRS, 3),
        ("ПР (9 часов): Практикум «Интервью»:", "ПР", WorkKind.PR, 9),
        ("ЛР (6 часов):", "ЛР", WorkKind.LB, 6),
        ("ЛК (1,5 ч)", "ЛК", WorkKind.LK, 1.5),
        ("Практическое занятие – 1 час", "Практическое занятие", WorkKind.PR, 1),
        ("Теоретическое занятие – 0,5 часа", "Теоретическое занятие", WorkKind.LK, 0.5),
        ("СРС: чтение литературы", "СРС", WorkKind.SRS, None),
    ],
)
def test_match_label(paragraph: str, label: str, kind: WorkKind, hours: float | None) -> None:
    found = match_label(paragraph)
    assert found is not None
    assert (found.text, found.kind, found.hours) == (label, kind, hours)


@pytest.mark.parametrize(
    "paragraph", ["Лекция о главном", "Практика: разбор", "ПРоект (3 часа)", "Сравнение (2 часа)"]
)
def test_match_label_rejects_plain_text(paragraph: str) -> None:
    assert match_label(paragraph) is None


def test_requisites_filled() -> None:
    assert requisites_filled("Приказом № 17-ОБР от «31» июля 2026 г.")
    assert requisites_filled("Приказом № ____ от 31.07.2026")
    assert not requisites_filled("Приказом № ____-ОБР от «___» __________ 2026 г.")
    assert not requisites_filled("Приказом № ____ от «____» ____________ ____ г.")
    assert not requisites_filled(None)


def test_codes_and_refs() -> None:
    text = "ПС 06.015 (приказ от 13.07.2023 № 586н), ТФ С/11.6 и C/14.6; версия 3.10; 12.345,6"
    assert find_ps_codes(text) == ["06.015"]
    assert find_tf_codes(text) == ["C/11.6", "C/14.6"]  # кириллическая «С» → латинская
    refs = find_order_refs(["Письмом Минобрнауки России от 21.04.2015 № ВК-1013/06 рекомендовано."])
    assert refs == ["Письмом Минобрнауки России от 21.04.2015 № ВК-1013/06"]
    assert number_key("1.6.") == "1.6" and number_key(None) == ""
