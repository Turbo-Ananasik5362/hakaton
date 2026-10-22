"""Выгрузки слайдов и тестов (app/render/pptx.py, gift.py, moodle_xml.py, csv_tests.py, bundle.py) и сквозная выдача
всех файлов через API."""

from __future__ import annotations

import csv
import io
import re
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.config import Settings
from app.core.registry import default_registry
from app.models import Kit, QuestionKind, TestScope
from app.pipeline.base import StepContext
from app.pipeline.content import slides_lectures
from app.pipeline.default import DefaultPipeline
from app.render.bundle import NO_SLIDES_NOTE, render_all
from app.render.common import load_organization
from app.render.csv_tests import render_csv
from app.render.gift import escape, parse_gift, render_gift
from app.render.kit_json import kit_export
from app.render.moodle_xml import render_moodle_xml
from app.render.pptx import render_pptx
from app.render.questions import question_bank

FIXTURES = Path(__file__).parent / "fixtures" / "kits"


def load(n: int) -> Kit:
    return Kit.model_validate_json((FIXTURES / f"open_q{n}.json").read_text(encoding="utf-8"))


# --------------------------------------------------------------------------------------------------
# Свой мини-парсер GIFT — независимо от app/render/gift.parse_gift
# --------------------------------------------------------------------------------------------------

QUESTION_RE = re.compile(r"::((?:\\.|[^:\\]|:(?!:))*)::((?:\\.|[^{\\])*)\{((?:\\.|[^}\\])*)\}", re.S)
ANSWER_RE = re.compile(r"([=~])(%-?[\d.]+%)?((?:\\.|[^=~#\\]|#(?!###))*)")


def unescape(text: str) -> str:
    return re.sub(r"\\(.)", lambda m: "\n" if m.group(1) == "n" else m.group(1), text).strip()


def mini_gift(source: str) -> list[dict]:
    """Вопросы: категория, название, текст, варианты (текст, вес), общая обратная связь «####»."""
    out: list[dict] = []
    category = ""
    events = [(m.start(), "cat", m) for m in re.finditer(r"^\$CATEGORY:\s*(.+)$", source, re.M)]
    events += [(m.start(), "q", m) for m in QUESTION_RE.finditer(source)]
    for _, kind, match in sorted(events, key=lambda e: e[0]):
        if kind == "cat":
            category = match.group(1).strip()
            continue
        body = match.group(3)
        feedback = ""
        if "####" in body:
            body, feedback = body.split("####", 1)
        answers = []
        for a in ANSWER_RE.finditer(body):
            weight = float(a.group(2).strip("%")) if a.group(2) else (100.0 if a.group(1) == "=" else 0.0)
            answers.append((unescape(a.group(3)), weight))
        out.append(
            {
                "category": category,
                "name": unescape(match.group(1)),
                "text": unescape(match.group(2)),
                "answers": answers,
                "feedback": unescape(feedback),
            }
        )
    return out


# --------------------------------------------------------------------------------------------------
# GIFT, Moodle XML, CSV
# --------------------------------------------------------------------------------------------------


@pytest.mark.parametrize("n", [1, 2, 3, 4])
def test_gift_is_parsed_back_to_the_bank(n: int, tmp_path: Path) -> None:
    kit = load(n)
    bank = question_bank(kit)
    path = tmp_path / "tests.gift"
    assert render_gift(kit, path) == len(bank) > 0
    parsed = mini_gift(path.read_text(encoding="utf-8"))
    assert len(parsed) == len(bank)
    for item, got in zip(bank, parsed, strict=True):
        q = item.question
        assert got["name"].startswith(q.id) and got["text"] == q.stem.strip()
        assert [a for a, _ in got["answers"]] == [o.strip() for o in q.options]
        right = [i for i, (_, w) in enumerate(got["answers"]) if w > 0]
        assert right == sorted(q.correct)
        assert abs(sum(w for _, w in got["answers"] if w > 0) - 100) < 0.01
        assert got["feedback"] == q.explanation.strip()
        assert got["category"].endswith(item.category)
    module_tests = [t for t in kit.tests if t.scope == TestScope.MODULE]
    categories = {p["category"] for p in parsed}
    assert len(categories) == len({t.module_id for t in module_tests}) + any(
        t.scope == TestScope.FINAL for t in kit.tests
    )
    assert len(parse_gift(path.read_text(encoding="utf-8"))) == len(bank)  # и разбор для V31


def test_gift_escapes_special_characters(tmp_path: Path) -> None:
    kit = load(1)
    tricky = "Что выведет print({'a': 1}) ~ x = y # z \\ путь: C:\\temp?\nВторая строка"
    question = kit.questions[0]
    question.stem = tricky
    question.options[0] = "=равно ~тильда {скобки}"
    question.explanation = "Пояснение: #### не обратная связь"
    path = tmp_path / "tests.gift"
    render_gift(kit, path)
    text = path.read_text(encoding="utf-8")
    assert escape("a{b}") == "a\\{b\\}" and "\\~" in text and "\\=" in text and "\\#" in text
    parsed = {p["name"].split()[0]: p for p in mini_gift(text)}
    got = parsed[question.id]
    assert got["text"] == tricky and got["answers"][0][0] == "=равно ~тильда {скобки}"
    assert got["feedback"] == "Пояснение: #### не обратная связь"


def test_multiple_answers_get_weights(tmp_path: Path) -> None:
    kit = load(3)
    question = next(q for q in kit.questions if len(q.options) >= 4)
    question.correct = [0, 2]
    question.kind = QuestionKind.MULTIPLE
    gift, xml_path = tmp_path / "t.gift", tmp_path / "t.xml"
    render_gift(kit, gift)
    render_moodle_xml(kit, xml_path)
    got = next(
        p for p in mini_gift(gift.read_text(encoding="utf-8")) if p["name"].startswith(question.id + " ")
    )
    weights = [w for _, w in got["answers"]]
    assert weights[0] == weights[2] == 50 and all(w < 0 for i, w in enumerate(weights) if i not in (0, 2))
    node = next(
        q
        for q in ET.parse(str(xml_path)).getroot().findall("question[@type='multichoice']")
        if q.find("name/text").text.startswith(question.id + " ")
    )
    assert node.find("single").text == "false"
    assert [a.get("fraction") for a in node.findall("answer")][:3] == ["50", weights_str(weights[1]), "50"]


def weights_str(value: float) -> str:
    return f"{value:g}" if value != int(value) else str(int(value))


@pytest.mark.parametrize("n", [1, 3])
def test_moodle_xml_is_valid(n: int, tmp_path: Path) -> None:
    kit = load(n)
    kit.questions[0].stem = "Что вернёт df[df.x < 5] & df.y?"  # символы XML и HTML
    path = tmp_path / "tests.xml"
    count = render_moodle_xml(kit, path)
    root = ET.parse(str(path)).getroot()
    assert root.tag == "quiz"
    questions = root.findall("question[@type='multichoice']")
    assert len(questions) == count == len(question_bank(kit))
    categories = [c.find("category/text").text for c in root.findall("question[@type='category']")]
    assert all(c.startswith("$course$/top/") for c in categories)
    assert [c.rsplit("/", 1)[-1] for c in categories] == list(
        dict.fromkeys(i.category for i in question_bank(kit))
    )
    by_id = {q.id: q for q in kit.questions}
    for node in questions:
        q = by_id[node.find("name/text").text.split()[0]]
        answers = node.findall("answer")
        assert len(answers) == len(q.options)
        assert (
            abs(sum(float(a.get("fraction")) for a in answers if float(a.get("fraction")) > 0) - 100) < 0.01
        )
        assert node.find("questiontext").get("format") == "html"
    first = by_id[questions[0].find("name/text").text.split()[0]]
    if first.id == kit.questions[0].id:
        assert "&lt; 5" in questions[0].find("questiontext/text").text


def test_csv_has_one_row_per_question(tmp_path: Path) -> None:
    kit = load(3)
    path = tmp_path / "tests.csv"
    count = render_csv(kit, path)
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM: Excel открывает кириллицу без мастера импорта
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter=";"))
    assert len(rows) - 1 == count == len(question_bank(kit))
    header = rows[0]
    by_id = {q.id: q for q in kit.questions}
    for row in rows[1:]:
        record = dict(zip(header, row, strict=True))
        q = by_id[record["ID вопроса"]]
        assert record["Вопрос"] == q.stem and record["Вариант а"] == q.options[0]


# --------------------------------------------------------------------------------------------------
# Слайды
# --------------------------------------------------------------------------------------------------


@pytest.mark.parametrize("n", [1, 2, 3, 4])
def test_slides_follow_lecture_plans(n: int, tmp_path: Path) -> None:
    kit = load(n)
    lectures = slides_lectures(kit)
    path = tmp_path / "slides.pptx"
    total = render_pptx(kit, path)
    deck = Presentation(str(path))
    slides = list(deck.slides)
    assert len(slides) == total == sum(len(lec.plan) + 2 for lec in lectures)
    topics = {t.id: t.title for t in kit.topics}
    titles = [s.shapes.title.text_frame.text for s in slides]
    pos = 0
    for lecture in sorted(lectures, key=lambda lec: titles.index(lec.title)):
        assert titles[pos] == topics[lecture.topic_ids[0]]  # титул = тема учебного плана
        assert titles[pos + 1 : pos + 1 + len(lecture.plan)] == [p.title for p in lecture.plan]
        assert titles[pos + len(lecture.plan) + 1] == "Итоги лекции"
        pos += len(lecture.plan) + 2
    for slide in slides:
        assert slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()
        assert not any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
    body_counts = [
        len([p for p in s.shapes if p.name == "Тезисы"][0].text_frame.paragraphs)
        for s in slides
        if any(p.name == "Тезисы" for p in s.shapes)
    ]
    assert all(1 <= c <= 5 for c in body_counts)


def test_only_lectures_from_slides_scope(tmp_path: Path) -> None:
    kit = load(1)  # «слайды для первых двух занятий»
    lectures = slides_lectures(kit)
    assert {lec.session_no for lec in lectures} == {1, 2} and len(lectures) < len(kit.lectures)
    render_pptx(kit, tmp_path / "slides.pptx")
    titles = {s.shapes.title.text_frame.text for s in Presentation(str(tmp_path / "slides.pptx")).slides}
    others = [
        lec.title
        for lec in kit.lectures
        if lec not in lectures and lec.title not in {x.title for x in lectures}
    ]
    assert not titles & set(others)


# --------------------------------------------------------------------------------------------------
# Сборка выгрузок, JSON и API
# --------------------------------------------------------------------------------------------------


def _ctx(path: Path) -> StepContext:
    return StepContext(session_id="s_000000000001", inputs_dir=path, files_dir=path, request=None)  # type: ignore[arg-type]


def test_render_produces_all_files_and_json_roundtrips(tmp_path: Path) -> None:
    kit = load(3)
    files = DefaultPipeline(Settings(llm_mode="mock")).render(kit, _ctx(tmp_path))
    assert set(files) == {
        "program_docx",
        "slides_pptx",
        "tests_gift",
        "tests_moodle_xml",
        "tests_csv",
        "kit_json",
        "report_json",
        "report_md",
    }
    assert all((tmp_path / name).is_file() for name in files.values())
    loaded = Kit.model_validate_json((tmp_path / "kit.json").read_text(encoding="utf-8"))
    assert kit_export(loaded) == kit_export(kit)  # обратимо
    assert not [i for i in kit.report.issues if i.check == "V31"]


def test_internship_without_tests_and_lectures_explains(tmp_path: Path) -> None:
    kit = load(5)
    outcome = render_all(kit, tmp_path, default_registry(), load_organization(Settings().organization_file))
    assert set(outcome.files) == {"program_docx"} and outcome.issues == []
    assert NO_SLIDES_NOTE in outcome.notes
    assert any("Тестовых вопросов в комплекте нет" in n and "собеседование" in n for n in outcome.notes)


def test_one_failed_export_keeps_the_others(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    def broken(*_args, **_kwargs):  # noqa: ANN002, ANN003, ANN202
        raise RuntimeError("сломан шаблон слайдов")

    monkeypatch.setattr("app.render.bundle.render_pptx", broken)
    kit = load(3)
    files = DefaultPipeline(Settings(llm_mode="mock")).render(kit, _ctx(tmp_path))
    assert "slides_pptx" not in files and not (tmp_path / "slides.pptx").exists()
    assert {"program_docx", "tests_gift", "kit_json"} <= set(files)
    assert any(n.startswith("Слайды .pptx: файл не сформирован") for n in kit.meta.generation_notes)
    issues = [i for i in kit.report.issues if i.check == "V31"]
    assert len(issues) == 1 and issues[0].id == "V31-1" and "сломан шаблон слайдов" in issues[0].detail


def test_all_files_download_through_api(tmp_path: Path) -> None:
    """Сквозной прогон запроса 1 в mock: каждый файл из files скачивается и разбирается."""
    from docx import Document

    from tests.intake_queries import OPEN
    from tests.test_content import run

    body = run(tmp_path, OPEN[1])
    files = body["files"]
    assert set(files) == {
        "program_docx",
        "slides_pptx",
        "tests_gift",
        "tests_moodle_xml",
        "tests_csv",
        "kit_json",
        "report_json",
        "report_md",
    }
    from fastapi.testclient import TestClient

    from app.main import create_app

    settings = Settings(sessions_dir=tmp_path / "sessions", llm_mode="mock")
    with TestClient(create_app(settings=settings)) as client:
        got = {key: client.get(url) for key, url in files.items()}
    assert all(r.status_code == 200 and r.content for r in got.values())
    Document(io.BytesIO(got["program_docx"].content))
    deck = Presentation(io.BytesIO(got["slides_pptx"].content))
    kit = Kit.model_validate_json(got["kit_json"].content)
    assert len(deck.slides) == sum(len(lec.plan) + 2 for lec in slides_lectures(kit))
    assert len(mini_gift(got["tests_gift"].content.decode("utf-8"))) == len(kit.questions) == 10
    assert len(ET.fromstring(got["tests_moodle_xml"].content).findall("question[@type='multichoice']")) == 10
    assert (
        len(list(csv.reader(io.StringIO(got["tests_csv"].content.decode("utf-8-sig")), delimiter=";"))) == 11
    )
    assert kit_export(kit) == body["kit"]
    assert not [i for i in body["report"]["issues"] if i["check"] == "V31"]


@pytest.mark.parametrize(
    ("hours", "expected"),
    [
        (4, "не более 4 часов"),
        (1, "не более 1 часа"),
        (21, "не более 21 часа"),
        (11, "не более 11 часов"),
        (0.5, "не более 0,5 часа"),
        (6, "не более 6 часов"),
        (2, "не более 2 часов"),
    ],
)
def test_at_most_hours_uses_genitive(hours: float, expected: str) -> None:
    """«не более 4 часа в день» в программе (оценка 21.10): после «более» — родительный падеж."""
    from app.render.common import at_most_hours

    assert at_most_hours(hours) == expected


def test_code_lines_in_slides_are_line_breaks_in_monospace(tmp_path: Path) -> None:
    """Строки кода в тезисе — переносы строк (<a:br/>) моноширинным шрифтом с отступами, а не «\\n» внутри текста."""
    from pptx import Presentation

    from app.render.pptx import MONO_FONT, _thesis_runs

    prs = Presentation()
    box = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, 100, 100)
    paragraph = box.text_frame.paragraphs[0]
    _thesis_runs(paragraph, "Цикл:\nfor i in range(3):\n    print(i)", 20)
    xml = paragraph._p.xml
    assert xml.count("<a:br") == 2 and "\n" not in "".join(r.text for r in paragraph.runs)
    runs = paragraph.runs
    assert [r.font.name for r in runs[1:]] == [MONO_FONT, MONO_FONT]
    assert runs[2].text == " " * 4 + "print(i)"


@pytest.mark.parametrize(
    ("points", "expected"),
    [(3, "до 3 баллов"), (1, "до 1 балла"), (21, "до 21 балла"), (11, "до 11 баллов"), (0.5, "до 0,5 балла")],
)
def test_up_to_points_uses_genitive(points: float, expected: str) -> None:
    """«до 3 балл.» 36 раз в программе для школьников (прогон 22.10): после «до» — родительный падеж."""
    from app.render.common import up_to_points

    assert up_to_points(points) == expected
