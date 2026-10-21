"""Отчёт агента для методиста (P6.1, docs/ARCHITECTURE.md §1.3): предупреждения по общим правилам — с цифрами
и вариантами решения, без жаргона; сводки report.json; блок предупреждений в приложении .docx; раздел отчёта
в веб-интерфейсе."""

from __future__ import annotations

import json
import re
from pathlib import Path

import docx
import pytest
from fastapi.testclient import TestClient

from app.config import Settings
from app.core import methodist_warnings as mw
from app.core.calendar import build_calendar
from app.core.hours import HoursConstraints, ModuleSpec, TopicSpec, constraints_from_request, solve_hours
from app.core.registry import default_registry
from app.core.validate import ValidationContext, validate
from app.main import create_app
from app.models import (
    Constraints,
    HoursRow,
    Kit,
    MethodistWarning,
    ProgramType,
    RequestSpec,
    ValidationReport,
    Verification,
    WarningType,
    WorkMix,
)
from app.pipeline.default import DefaultPipeline
from app.pipeline.editing import conflict_details
from app.render.common import load_organization
from app.render.program import render_program_docx
from app.render.report_md import api_message, build_report
from app.web.report_page import render_report_page
from tests.intake_queries import OPEN
from tests.test_calendar import MIX_DOP, make_kit
from tests.test_content import attachments_payload
from tests.validate_kits import make_internship_kit, make_pk_kit, make_registry
from tests.validate_mutations import load_reference

FIXTURES = Path(__file__).parent / "fixtures" / "kits"


def load(n: int) -> Kit:
    return Kit.model_validate_json((FIXTURES / f"open_q{n}.json").read_text(encoding="utf-8"))


def warnings_of(kit: Kit, ctx: ValidationContext | None = None) -> dict[WarningType, MethodistWarning]:
    report = validate(kit, ctx or ValidationContext(registry=make_registry()))
    return {w.type: w for w in report.warnings_for_methodist}


def assert_complete(warning: MethodistWarning) -> None:
    """Общее требование: заголовок, цифры, не меньше двух вариантов решения, без служебных обозначений."""
    assert warning.title and warning.text
    assert warning.figures and any(re.search(r"\d", f.value) for f in warning.figures)
    assert len(warning.options) >= 2, warning.options
    assert mw.missing_parts(warning) == [] and mw.jargon_in(warning) is None


# --------------------------------------------------------------------------------------------------
# Предупреждения по общим правилам
# --------------------------------------------------------------------------------------------------


def test_reference_errors_counted_by_kind_with_concrete_numbers() -> None:
    kit = make_internship_kit()
    kit.request.reference = load_reference("bsa_248_sample.json")
    warning = warnings_of(kit)[WarningType.REFERENCE_ERRORS]
    assert_complete(warning)
    figures = {f.label: f.value for f in warning.figures}
    total = int(figures["Всего расхождений"])
    assert total >= 10 and sum(int(v) for k, v in figures.items() if k != "Всего расхождений") == total
    # конкретика из чек-листа Q5.3b — прямо в тексте предупреждения
    assert "66" in warning.text and "174" in warning.text and "168" in warning.text
    # в программу (.docx) — без цитат и чисел образца
    assert warning.program_text and "«" not in warning.program_text and "174" not in warning.program_text


def test_expectation_gap_from_request_wording() -> None:
    kit = make_pk_kit()
    kit.request.raw_text += " Хотим, чтобы после курса слушатели могли пойти работать в эту сферу."
    warning = warnings_of(kit)[WarningType.EXPECTATION_GAP]
    assert_complete(warning)
    figures = {f.label: f.value for f in warning.figures}
    assert figures["Ожидание из запроса"] == "«пойти работать в эту сферу»"
    assert figures["Объём программы"] == f"{kit.program.total_hours:g} ч"
    assert any("переподготовк" in option for option in warning.options)
    assert WarningType.EXPECTATION_GAP not in warnings_of(make_pk_kit())  # нет ожидания — нет предупреждения


def test_schedule_days_options_are_computed() -> None:
    request = RequestSpec(
        program_type=ProgramType.PK_INTERNSHIP,
        total_hours=256,
        fa_hours=2,
        hours_per_day_max=4,
        weeks=9,
        duration_text="около двух месяцев",
        constraints=Constraints(all_srs=True),
    )
    constraints = constraints_from_request(request, ProgramType.PK_INTERNSHIP)
    kit = make_kit(constraints, [2, 3, 3, 2], duration_text=request.duration_text)
    _, warnings = build_calendar(kit, constraints)
    warning = warnings[0]
    assert_complete(warning)
    figures = {f.label: f.value for f in warning.figures}
    assert figures["Нужно учебных дней"] == "64" and figures["Расчётный срок"] == "11 недель"
    accept, load_more, cut = warning.options
    assert "11 недель" in accept
    assert "5 ч в день при 6 учебных днях" in load_more  # 256 ч / (9 недель × 6 дней) → 5 ч
    assert "216 ч" in cut  # 9 недель × 6 дней × 4 ч


def test_schedule_sessions_options_are_computed() -> None:
    constraints = HoursConstraints(
        total_hours=32, program_type=ProgramType.DOP, session_hours=2, sessions_per_week=2, weeks_target=4
    )
    kit = make_kit(constraints, [3, 3], mix=MIX_DOP)
    _, warnings = build_calendar(kit, constraints)
    warning = warnings[0]
    assert_complete(warning)
    assert "Проводить 4 занятия в неделю" in warning.options[1]
    assert "16 ч" in warning.options[2]  # 4 недели × 2 занятия × 2 ч
    assert any("Удлинить занятия до 4 ч" in o for o in warning.options)


def test_norm_replaced_keeps_old_number_out_of_program_text() -> None:
    kit = make_pk_kit()
    kit.request.raw_text += " Программа должна соответствовать приказу № 499."
    warning = warnings_of(kit, ValidationContext(registry=default_registry()))[WarningType.NORM_REPLACED]
    assert_complete(warning)
    assert "№ 266" in warning.text and "№ 499" in warning.text and "упоминается в запросе" in warning.text
    assert "№ 266" in warning.program_text and "499" not in warning.program_text
    assert not any("499" in option for option in warning.options)
    assert any(f.label == "Источник" and f.value.startswith("http") for f in warning.figures)


def test_unverified_sources_say_what_remains_after_removal() -> None:
    kit = make_pk_kit()
    book = next(s for s in kit.sources if s.section.value == "main")
    book.verification = Verification.UNVERIFIED
    kit.meta.literature_checks = {"rejected": 2}
    warning = warnings_of(kit)[WarningType.UNVERIFIED_SOURCES]
    assert_complete(warning)
    assert warning.target_ids == [book.id]
    figures = {f.label: f.value for f in warning.figures}
    assert (
        figures["Без подтверждения"] == f"1 из {len(kit.sources)}"
        and figures["Отброшено при проверке"] == "2"
    )
    assert "«Основная литература»" in warning.options[2]


def test_solver_conflict_has_options_and_figures() -> None:
    constraints = HoursConstraints(
        total_hours=72, program_type=ProgramType.PK, all_srs=True, min_practice_share=0.5
    )
    modules = [ModuleSpec(key="M1", weight=1)]
    topics = [
        TopicSpec(key="T1.1", module_key="M1", weight=1, mix=WorkMix(lk=20, pr=60, srs=20), title="Сети"),
        TopicSpec(key="T1.2", module_key="M1", weight=1, mix=WorkMix(lk=40, pr=40, srs=20), title="Linux"),
    ]
    warning = solve_hours(modules, topics, constraints).warnings()[0]
    assert warning.type == WarningType.CONSTRAINT_CONFLICT and warning.code == "all_srs_vs_practice"
    assert_complete(warning)
    assert {f.label: f.value for f in warning.figures}["Доля практики в запросе"] == "50 %"


def test_solver_messages_name_topics_not_keys() -> None:
    constraints = HoursConstraints(total_hours=20, program_type=ProgramType.PK, dot=True)
    modules = [ModuleSpec(key="M1", weight=1)]
    topics = [
        TopicSpec(
            key="T1.1",
            module_key="M1",
            weight=1,
            mix=WorkMix(lk=40, pr=40, srs=20),
            locked=HoursRow(lk=2, srs=2, distance=5),
            title="Права доступа",
        ),
        TopicSpec(key="T1.2", module_key="M1", weight=1, mix=WorkMix(lk=40, pr=40, srs=20), title="Сети"),
    ]
    result = solve_hours(modules, topics, constraints)
    message = result.conflicts[0].message
    assert "«Права доступа»" in message and "T1.1" not in message
    details = conflict_details(result)[0]  # то же — в ответе правки часов (редактор каркаса)
    assert details["title"] and details["options"]


def test_methodist_warnings_go_actionable_first() -> None:
    kit = make_internship_kit()
    kit.request.reference = load_reference("bsa_248_sample.json")
    ctx = ValidationContext(
        registry=default_registry(),
        extra_warnings=[
            mw.schedule_days(
                target_weeks=9,
                target_text="9 недель",
                total=256,
                cap=4,
                days=64,
                per_week=6,
                weeks=11,
                need_per_day=5,
            )
        ],
    )
    types = [w.type for w in validate(kit, ctx).warnings_for_methodist]
    assert types.index(WarningType.SCHEDULE_TIGHT) < types.index(WarningType.REFERENCE_ERRORS)
    assert types[-1] == WarningType.NORM_REPLACED


# --------------------------------------------------------------------------------------------------
# V46: полнота и язык предупреждений (заведомо ошибочные примеры)
# --------------------------------------------------------------------------------------------------


def v46(kit: Kit, extra: MethodistWarning) -> list[str]:
    ctx = ValidationContext(registry=make_registry(), extra_warnings=[extra])
    return [i.message for i in validate(kit, ctx, only=["V46"]).issues]


def test_v46_flags_warning_without_numbers_and_options() -> None:
    messages = v46(
        make_pk_kit(), MethodistWarning(type=WarningType.CONSTRAINT_CONFLICT, text="Требования спорят.")
    )
    assert any("без заголовка, цифр и вариантов решения" in m for m in messages)


def test_v46_accepts_carried_over_warning_from_old_report() -> None:
    """kit.json прежней версии: «график плотный» без заголовка и вариантов — не замечание, а умолчания рендера."""
    kit = make_pk_kit()
    old = MethodistWarning(type=WarningType.SCHEDULE_TIGHT, text="Нужно 6 ч в день при лимите 4 ч.")
    kit.report = ValidationReport(warnings_for_methodist=[old])
    report = validate(kit, ValidationContext(registry=make_registry()), only=["V46"])
    assert report.issues == []
    carried = next(w for w in report.warnings_for_methodist if w.type == WarningType.SCHEDULE_TIGHT)
    assert carried.title == mw.title_of(old) and carried.options == mw.options_of(old)


def test_v46_flags_object_ids_in_warning_text() -> None:
    warning = mw.schedule_days(
        target_weeks=4, target_text="4 недели", total=72, cap=6, days=12, per_week=5, weeks=5, need_per_day=4
    )
    warning.text += " Сократите тему T2.3."
    messages = v46(make_pk_kit(), warning)
    assert any("служебное обозначение «T2.3»" in m for m in messages)
    assert v46(make_pk_kit(), warning.model_copy(update={"text": "Сократите тему 2.3."})) == []
    quoted = warning.model_copy(update={"text": "Не подтверждён источник «Хранилище Amazon S3 и Q1»."})
    assert v46(make_pk_kit(), quoted) == []  # в кавычках — название, а не служебный ID


# --------------------------------------------------------------------------------------------------
# report.json
# --------------------------------------------------------------------------------------------------


def test_legacy_report_from_old_kit_json_still_loads() -> None:
    legacy = {
        "hours": {"M1": {"lk": 2, "pr": 2}, "TOTAL": {"lk": 2, "pr": 2}},
        "sources": {"registry": 1, "catalog": 2, "online": 0, "unverified": 0},
        "warnings_for_methodist": [{"type": "schedule_tight", "text": "Срок не сходится."}],
    }
    report = ValidationReport.model_validate(legacy)
    assert report.hours is None and report.sources is None
    warning = report.warnings_for_methodist[0]
    assert mw.title_of(warning) and mw.options_of(
        warning
    )  # показывается с заголовком и вариантами по умолчанию


@pytest.fixture(scope="module")
def q5_session(tmp_path_factory: pytest.TempPathFactory) -> tuple[dict, dict[str, bytes], dict[str, str]]:
    """Запрос 5 (стажировка с образцом) через API в mock: ответ done, файлы и страница отчёта."""
    tmp = tmp_path_factory.mktemp("q5")
    settings = Settings(sessions_dir=tmp / "sessions", llm_mode="mock", cost_log=tmp / "cost.jsonl")
    query = OPEN[5]
    with TestClient(create_app(settings=settings, pipeline=DefaultPipeline(settings))) as client:
        body = client.post(
            "/api/v1/sessions",
            json={
                "request": query["request"],
                "attachments": attachments_payload(query),
                "answers": query.get("methodist_answers", {}),
                "auto_approve": True,
            },
        ).json()
        files = {key: client.get(url).content for key, url in body.get("files", {}).items()}
        sid = body["session_id"]
        pages = {
            "page": client.get(f"/sessions/{sid}/report").text,
            "partial": client.get(f"/sessions/{sid}/report?partial=1").text,
            "missing": client.get("/sessions/s_000000000000/report").text,
            "bad": client.get("/sessions/<script>alert(1)</script>/report?partial=1").text,
        }
    assert body["status"] == "done", body.get("message")
    return body, files, pages


def test_report_json_has_every_part_of_the_report(q5_session) -> None:  # noqa: ANN001
    body, files, _ = q5_session
    report = body["report"]
    assert json.loads(files["report_json"])["warnings_for_methodist"] == report["warnings_for_methodist"]
    hours = report["hours"]
    assert hours["matches"] and hours["declared"] == hours["total"] == 256 and hours["contact"] == 0
    assert hours["lines"][-1]["level"] == "total" and hours["lines"][-2]["level"] == "final_attestation"
    assert report["coverage"] and all(
        c["status"] and c["text"] and "by_kind" in c for c in report["coverage"]
    )
    sources = report["sources"]
    assert sources["confirmed"] + sources["unverified"] == sources["total"] > 0
    assert report["sections"] and all(isinstance(v, bool) for v in report["sections"].values())
    assert report["copying"]["references"] == 1 and report["copying"]["passed"]
    assert report["judge"] and all(j["title"] for j in report["judge"])
    assert report["llm_calls"] > 0 and report["cost_rub"] == 0  # mock — бесплатно
    assert report["generation_seconds"] is not None and report["timing"]
    assert report["summary"].startswith("Комплект собран")
    types = {w["type"] for w in report["warnings_for_methodist"]}
    assert {"reference_errors", "schedule_tight", "norm_replaced"} <= types
    for warning in report["warnings_for_methodist"]:
        assert warning["title"] and warning["figures"] and warning["options"]
    assert not [i for i in report["issues"] if i["check"] == "V46"]
    assert body["report_page"] == f"/sessions/{body['session_id']}/report"
    assert "Предупреждения методисту (с цифрами и вариантами решения): Желаемый срок" in body["message"]


# --------------------------------------------------------------------------------------------------
# Приложение .docx и report.md
# --------------------------------------------------------------------------------------------------


def docx_text(path: Path) -> str:
    document = docx.Document(str(path))
    cells = [c.text for t in document.tables for r in t.rows for c in r.cells]
    return "\n".join([p.text for p in document.paragraphs] + cells)


def test_docx_appendix_has_methodist_warnings_without_forbidden_texts(tmp_path: Path) -> None:
    kit = load(5)
    types = {w.type for w in kit.report.warnings_for_methodist}
    assert {WarningType.REFERENCE_ERRORS, WarningType.NORM_REPLACED, WarningType.SCHEDULE_TIGHT} <= types
    name, issues = render_program_docx(
        kit, tmp_path, default_registry(), load_organization(Settings().organization_file)
    )
    assert issues == []
    text = docx_text(tmp_path / name)
    for expected in (
        "Предупреждения методисту",
        "Желаемый срок не сходится с нагрузкой",
        "Цифры: Желаемый срок",
        "Ошибки в приложенном образце",
        "Варианты решения:",
        "Приказ Минобрнауки России от 24.03.2025 № 266",
    ):
        assert expected in text, expected
    for forbidden in ("№ 499", "№ 816", "«Итого» 66", "в «Итого» 174", "pk_stazhirovka"):
        assert forbidden not in text, forbidden


def test_report_md_keeps_full_texts(q5_session) -> None:  # noqa: ANN001
    _, files, _ = q5_session
    md = files["report_md"].decode("utf-8")
    assert "## Предупреждения методисту" in md and "**Варианты решения:**" in md
    assert (
        "№ 499" in md and "в «Итого» 174" in md
    )  # полная версия — с номером отменённого акта и цифрами образца
    assert "## Ошибки приложенного образца" in md


def test_api_message_names_warnings() -> None:
    message = api_message(load(4))
    assert "Предупреждения методисту (с цифрами и вариантами решения): Ожидания заказчика выше" in message


# --------------------------------------------------------------------------------------------------
# Раздел в интерфейсе
# --------------------------------------------------------------------------------------------------


def test_report_page_shows_status_in_color_and_words(q5_session) -> None:  # noqa: ANN001
    body, _, pages = q5_session
    page = pages["page"]
    assert page.startswith("<!doctype html>") and 'lang="ru"' in page
    assert "status-warnings" in page and "Есть предупреждения — проверьте" in page
    assert page.count('class="advice') == len(body["report"]["warnings_for_methodist"])
    assert "Нужно решение" in page and "Для сведения" in page and "Варианты решения" in page
    assert 'class="tone tone-ok"' in page and "в порядке" in page  # цвет и текст статуса строки
    assert f"/api/v1/sessions/{body['session_id']}/files/program.docx" in page
    partial = pages["partial"]
    assert partial.startswith('<section class="report"') and "<html" not in partial


def test_report_page_for_missing_session_and_hostile_id(q5_session) -> None:  # noqa: ANN001
    _, _, pages = q5_session
    assert "не найдена" in pages["missing"]
    assert "<script>" not in pages["bad"]


def test_report_page_escapes_kit_texts() -> None:
    kit = load(1)
    kit.program.title = "<script>alert(1)</script>"
    page = render_report_page(build_report(kit), "s_000000000001")
    assert "<script>alert(1)</script>" not in page and "&lt;script&gt;" in page


def test_report_page_waits_for_kit(tmp_path: Path) -> None:
    settings = Settings(sessions_dir=tmp_path / "sessions", llm_mode="mock", cost_log=tmp_path / "cost.jsonl")
    query = OPEN[4]
    with TestClient(create_app(settings=settings, pipeline=DefaultPipeline(settings))) as client:
        body = client.post(
            "/api/v1/sessions",
            json={"request": query["request"], "answers": query.get("methodist_answers", {})},
        ).json()
        assert body["status"] == "skeleton_ready", body.get("message")
        response = client.get(f"/sessions/{body['session_id']}/report")
    assert response.status_code == 200
    assert "Каркас готов и ждёт утверждения" in response.text and 'http-equiv="refresh"' in response.text
    assert "report_page" not in body
