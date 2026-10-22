"""Исправления по замечаниям ревью генерации: шлюз лёг — шаги не ждут повторов, ссылки лекций на РО своего модуля,
лимиты собеседования и размера теста, метки кодов во всех полях, отмена частей теста, честный учёт повтора."""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.config import Settings
from app.core.registry import default_registry
from app.core.validate import mark_unknown_codes
from app.llm import GatewayError
from app.main import create_app
from app.models import (
    ExampleTask,
    InterviewDraft,
    Kit,
    LectureDraft,
    PlanItemDraft,
    SkeletonDraft,
    Source,
    SourceKind,
    SourceSection,
    TestsDraft,
    Verification,
)
from app.pipeline import content
from app.pipeline import content_apply as apply
from app.pipeline.base import AgentError, Step, StepContext
from app.pipeline.default import DefaultPipeline
from app.pipeline.request_rules import MAX_TEST_QUESTIONS, extract_facts
from app.render.bundle import NO_SLIDES_NOTE, NO_TESTS_NOTE
from app.render.report_md import api_message
from tests.fake_pipeline import FakePipeline
from tests.validate_kits import make_pk_kit

FIXTURES = Path(__file__).parent / "fixtures" / "kits"


def load(n: int) -> Kit:
    return Kit.model_validate_json((FIXTURES / f"open_q{n}.json").read_text(encoding="utf-8"))


def ctx(tmp_path: Path) -> StepContext:
    return StepContext(session_id="s_000000000001", inputs_dir=tmp_path, files_dir=tmp_path, request=None)  # type: ignore[arg-type]


# --------------------------------------------------------------------------------------------------
# 1. Шлюз не отвечает: оставшиеся вызовы модели не запускаются
# --------------------------------------------------------------------------------------------------


class OutagePipeline(FakePipeline):
    def generation_steps(self, kit: Kit, finished: set[str]) -> list[Step]:
        waves = [
            [("content:M1", True), ("content:M2", True)],
            [("lectures:M1", True), ("lectures:M2", True)],
            [("finalize", False)],
            [("validate", False)],
        ]
        for wave in waves:
            pending = [(n, llm) for n, llm in wave if n not in finished]
            if pending:
                return [
                    Step(
                        n,
                        f"Шаг {n}",
                        self._validate_step if n == "validate" else self._runner(n),
                        critical=n == "validate",
                        uses_llm=llm,
                    )
                    for n, llm in pending
                ]
        return []


def _run_outage(tmp_path: Path, failures: dict) -> tuple[OutagePipeline, dict]:  # noqa: ANN001
    pipeline = OutagePipeline(fail_once=failures)
    settings = Settings(sessions_dir=tmp_path / "sessions")
    with TestClient(create_app(settings=settings, pipeline=pipeline)) as c:
        body = c.post(
            "/api/v1/sessions",
            json={"request": "Программа на 36 часов", "answers": {"hours": "36"}, "auto_approve": True},
        ).json()
    return pipeline, body


def test_gateway_outage_skips_remaining_model_steps(tmp_path: Path) -> None:
    """Вся волна упала с «шлюз недоступен» — дальше вызовы модели не запускаются, код выполняется."""
    down = GatewayError("Шлюз недоступен: ConnectError")
    pipeline, body = _run_outage(
        tmp_path, {"content:M1": down, "content:M2": GatewayError("down", status=502)}
    )
    assert body["status"] == "done"
    assert pipeline.calls["lectures:M1"] == pipeline.calls["lectures:M2"] == 0  # не запускались
    assert pipeline.calls["finalize"] == 1 and pipeline.calls["validate"] == 1  # код выполнен
    steps = {s["name"]: s["status"] for s in body["progress"]["steps"]}
    assert steps["lectures:M1"] == steps["lectures:M2"] == "skipped"
    note = next(n for n in body["notes"] if "не отвечал" in n)
    assert "«Шаг lectures:M1»" in note and "«Шаг lectures:M2»" in note  # одна заметка на причину


def test_single_failure_with_live_gateway_does_not_stop_generation(tmp_path: Path) -> None:
    """Один шаг упал, соседний прошёл — шлюз жив: следующие шаги с моделью запускаются."""
    pipeline, body = _run_outage(tmp_path, {"content:M1": GatewayError("Шлюз недоступен: ReadTimeout")})
    assert body["status"] == "done"
    assert pipeline.calls["lectures:M1"] == pipeline.calls["lectures:M2"] == 1
    assert not any("не отвечал" in n for n in body["notes"])


def test_rate_limit_429_does_not_stop_other_steps(tmp_path: Path) -> None:
    pipeline = OutagePipeline(fail_once={"content:M1": GatewayError("busy", status=429)})
    settings = Settings(sessions_dir=tmp_path / "sessions")
    with TestClient(create_app(settings=settings, pipeline=pipeline)) as c:
        body = c.post(
            "/api/v1/sessions",
            json={"request": "Программа на 36 часов", "answers": {"hours": "36"}, "auto_approve": True},
        ).json()
    assert body["status"] == "done" and pipeline.calls["lectures:M1"] == 1


# --------------------------------------------------------------------------------------------------
# 2–5. Применение черновиков и метки
# --------------------------------------------------------------------------------------------------


def test_lecture_plan_links_only_to_outcomes_of_its_module() -> None:
    kit = make_pk_kit()
    first, second = kit.modules_sorted()[:2]
    topic = next(t for t in apply.lecture_topics(kit, first))
    foreign = next(
        lo
        for t in kit.topics_of(second.id)
        for lo in t.lo_ids
        if lo not in {x for t2 in kit.topics_of(first.id) for x in t2.lo_ids}
    )
    kit.lectures = [lec for lec in kit.lectures if lec.module_id != first.id]
    draft = LectureDraft(
        topic_ids=[topic.id],
        title=topic.title,
        plan=[PlanItemDraft(title="Раздел", theses=["Тезис"], lo_ids=[foreign])],
    )
    apply.apply_lectures(kit, first, [draft])
    lecture = next(lec for lec in kit.lectures if lec.module_id == first.id)
    assert foreign not in lecture.plan[0].lo_ids and lecture.plan[0].lo_ids == list(topic.lo_ids)


def test_interview_questions_are_capped_by_quota() -> None:
    kit = load(5)  # стажировка: текущий контроль — собеседование
    module = kit.modules_sorted()[0]
    kit.interview_questions = [iq for iq in kit.interview_questions if iq.module_id != module.id]
    quota = content.interview_quota(kit, module)
    los = [lo for t in kit.topics_of(module.id) for lo in t.lo_ids]
    draft = TestsDraft(
        interview=[
            InterviewDraft(module_id=module.id, lo_ids=los[:1], text=f"Вопрос {i} о практике")
            for i in range(quota + 7)
        ]
    )
    generator = DefaultPipeline(Settings(llm_mode="mock")).content
    generator._apply_tests(kit, module, draft, {}, quota)
    assert len([iq for iq in kit.interview_questions if iq.module_id == module.id]) == quota


def test_huge_test_is_capped_with_explanation() -> None:
    facts = extract_facts("Программа на 72 часа. Нужен итоговый тест на 1000 вопросов.")
    assert facts.test_questions == MAX_TEST_QUESTIONS
    assert any("1000 вопросов" in n and str(MAX_TEST_QUESTIONS) in n for n in facts.notes)
    assert extract_facts("Нужен итоговый тест на 30 вопросов.").test_questions == 30


def test_codes_are_marked_in_all_model_fields() -> None:
    kit = load(3)
    seminar = kit.seminars[0]
    seminar.title = "Разбор по приказу № 4321"
    kit.topics[0].control_form = "опрос по профстандарту 06.099"
    fa = kit.final_attestation
    fa.form_text = "Защита по письму № 77-ЗЗ"
    fa.example_task = ExampleTask(title="Задание по указу № 9999", description="Описание")
    registry = default_registry()
    content.mark_model_texts(kit, lambda text: mark_unknown_codes(text, registry))
    for text in (seminar.title, kit.topics[0].control_form, fa.form_text, fa.example_task.title):
        assert "[не подтверждено]" in text, text


# --------------------------------------------------------------------------------------------------
# 7–8. Части теста и повтор шага
# --------------------------------------------------------------------------------------------------


def test_failed_test_part_cancels_the_others() -> None:
    cancelled = []

    async def slow() -> int:
        try:
            await asyncio.sleep(5)
        except asyncio.CancelledError:
            cancelled.append(True)
            raise
        return 1

    async def broken() -> int:
        await asyncio.sleep(0.01)
        raise GatewayError("busy", status=429)

    async def main() -> None:
        with pytest.raises(GatewayError):
            await content.gather_or_cancel([slow(), broken(), slow()])

    started = asyncio.get_event_loop_policy().new_event_loop()
    try:
        started.run_until_complete(main())
    finally:
        started.close()
    assert cancelled == [True, True]


def test_retry_without_result_is_not_counted_as_success(tmp_path: Path) -> None:
    kit = make_pk_kit()
    kit.sources = [s for s in kit.sources if s.section == SourceSection.NORMATIVE]
    generator = DefaultPipeline(Settings(llm_mode="mock")).content

    async def nothing(kit: Kit, ctx: StepContext) -> None:
        return None

    async def adds_book(kit: Kit, ctx: StepContext) -> None:
        kit.sources.append(
            Source(
                id="SRC90",
                kind=SourceKind.BOOK,
                section=SourceSection.MAIN,
                title="Книга",
                verification=Verification.CATALOG,
            )
        )

    with pytest.raises(AgentError, match="повтор тоже не дал результата"):
        asyncio.run(generator._checked_retry("literature", nothing)(kit, ctx(tmp_path)))
    asyncio.run(
        generator._checked_retry("literature", adds_book)(kit, ctx(tmp_path))
    )  # результат есть — без ошибки


# --------------------------------------------------------------------------------------------------
# 10. Мелочи
# --------------------------------------------------------------------------------------------------


def test_api_message_does_not_promise_missing_report_md() -> None:
    kit = load(3)
    assert "report.md" in api_message(kit)
    without = api_message(kit, has_report_md=False)
    assert "report.md" not in without and "report.json" in without


def test_material_absence_notes_are_not_limitations() -> None:
    from scripts.run_open_query import NO_MATERIAL_NOTES

    assert NO_SLIDES_NOTE.startswith(NO_MATERIAL_NOTES)
    assert NO_TESTS_NOTE.format(forms="собеседование").startswith(NO_MATERIAL_NOTES)


# --------------------------------------------------------------------------------------------------
# Живой прогон 20.10: каркас без результатов обучения, тема с теорией без лекции
# --------------------------------------------------------------------------------------------------


def _skeleton_draft() -> SkeletonDraft:
    import json

    data = json.loads(
        (Path(__file__).parent / "fixtures" / "llm" / "skeleton" / "abc33f5e259a5fbe.json").read_text()
    )
    return SkeletonDraft.model_validate({k: v for k, v in data.items() if not k.startswith("_")})


def test_skeleton_without_outcomes_is_asked_again_with_correction() -> None:
    from types import SimpleNamespace

    from app.pipeline import skeleton
    from tests.intake_queries import OPEN

    good = _skeleton_draft()
    bad = good.model_copy(deep=True)
    for comp in bad.competencies:
        comp.outcomes = []
    assert "outcomes" in skeleton.draft_problem(bad) and skeleton.draft_problem(good) == ""

    calls: list[str] = []

    class FakeLLM:
        settings = SimpleNamespace(llm_mode="mock")

        async def complete_json(self, step, model, messages, schema, **kwargs):  # noqa: ANN001, ANN003, ANN201
            calls.append(messages[-1]["content"])
            return bad if len(calls) == 1 else good

    from app.models import ProgramType, RequestSpec
    from app.pipeline.default import DefaultPipeline

    spec = RequestSpec(raw_text=OPEN[3]["request"], program_type=ProgramType.PK, total_hours=72)
    pipeline = DefaultPipeline(Settings(llm_mode="mock"))
    kit = asyncio.run(
        skeleton.run_skeleton(
            spec, None, FakeLLM(), pipeline.registry, lambda k, w: pipeline.validate(k, w, None, "skeleton")
        )
    )
    assert len(calls) == 2 and "ВНИМАНИЕ, исправь предыдущий ответ" in calls[1] and "outcomes" in calls[1]
    assert kit.outcomes


def test_topics_without_lecture_are_refilled_only_where_missing() -> None:
    kit = load(3)
    module = next(m for m in kit.modules_sorted() if len(apply.lecture_topics(kit, m)) >= 2)
    topics = apply.lecture_topics(kit, module)
    dropped = topics[-1].id
    kit.lectures = [lec for lec in kit.lectures if dropped not in lec.topic_ids]
    gaps = content.ContentGenerator.lecture_gaps(kit)
    assert gaps == {module.id: [dropped]}
    before = len(kit.lectures)
    drafts = [
        LectureDraft(topic_ids=[t.id], title=t.title, plan=[PlanItemDraft(title="Раздел", theses=["Тезис"])])
        for t in topics
    ]
    created = apply.apply_lectures(kit, module, drafts, only=gaps[module.id])
    assert len(created) == 1 and len(kit.lectures) == before + 1
    assert content.ContentGenerator.lecture_gaps(kit) == {}


# --------------------------------------------------------------------------------------------------
# Живой прогон 20.10, запрос 1: литература, вводный модуль, слайды, отчёт в программе
# --------------------------------------------------------------------------------------------------


def test_catalog_key_in_url_and_misspelled_author_do_not_make_duplicates() -> None:
    from app.core import bibliography
    from app.models import BookCandidate

    registry = default_registry()
    record = registry.literature("python_lubanovic_2021")
    assert registry.find_literature(record.title, ["Лубанович Б."]).id == record.id  # опечатка в фамилии
    kit = load(1)
    items = [bibliography.ShortlistItem(f"lit:{record.id}", record)]
    result = bibliography.build_sources(
        kit,
        registry,
        items,
        picks=[],
        candidates=[
            BookCandidate(
                title="Простой Python",
                authors=["Лубанович Б."],
                url="lit:python_lubanovic_2021",
                section=SourceSection.MAIN,
            ),
            BookCandidate(title=record.title, authors=["Лубанович"], section=SourceSection.ADDITIONAL),
            BookCandidate(
                title="Книга со странным адресом",
                publisher="Питер",
                year=2024,
                url="ftp-ключ",
                section=SourceSection.ADDITIONAL,
            ),
        ],
        current_year=2026,
    )
    titles = [s.title for s in result.sources]
    assert titles.count(record.title) == 1  # ключ каталога и дубль кандидатом — одна запись из каталога
    assert all(not (s.url or "").startswith(("lit:", "doc:", "ftp")) for s in result.sources)


def test_multi_topic_first_module_is_not_an_unnumbered_intro() -> None:
    from app.models import ModuleKind
    from app.pipeline import skeleton

    draft = _skeleton_draft()
    draft.modules[0].kind = ModuleKind.INTRO
    assert len(draft.modules[0].topics) > 1
    from app.models import ProgramType, RequestSpec

    spec = RequestSpec(raw_text="ДОП", program_type=ProgramType.DOP, total_hours=16)
    final = skeleton.finalize_spec(spec)
    kit = skeleton.assemble(final, draft, [], default_registry(), lambda k, w: None)  # type: ignore[arg-type]
    assert kit.modules_sorted()[0].kind == ModuleKind.REGULAR


def test_term_is_not_repeated_on_summary_slide() -> None:
    assert (
        content._term_line("типы данных", "Тип данных — это вид значения.")
        == "Тип данных — это вид значения."
    )
    assert content._term_line("цикл", "повторение действий") == "цикл — повторение действий"


def test_title_slide_shows_lecture_plan(tmp_path: Path) -> None:
    from pptx import Presentation

    from app.render.pptx import render_pptx

    kit = load(1)
    render_pptx(kit, tmp_path / "s.pptx")
    first = list(Presentation(str(tmp_path / "s.pptx")).slides)[0]
    text = "\n".join(sh.text_frame.text for sh in first.shapes if sh.has_text_frame)
    lecture = content.slides_lectures(kit)[0]
    assert "План лекции:" in text and all(item.title in text for item in lecture.plan[:5])


def test_program_appendix_has_no_cost_section() -> None:
    from app.render.report_md import build_report

    kit = load(3)
    assert "Стоимость и время" in [s.title for s in build_report(kit).sections]
    assert "Стоимость и время" not in [s.title for s in build_report(kit, for_docx=True).sections]


# --------------------------------------------------------------------------------------------------
# Живой прогон 20.10, запрос 5: тексты стажировки, РО во многих модулях, 1.6, заметки в .docx, повторы
# --------------------------------------------------------------------------------------------------


def test_internship_texts_promising_lectures_are_flagged() -> None:
    from app.core.validate import ValidationContext, validate

    kit = load(5)
    kit.texts.form_description = (
        "Обучение проходит на базе учебных аудиторий: занятия включают лекции и практику."
    )
    issues = validate(kit, ValidationContext(registry=default_registry()), only=["V10"]).issues
    assert any("у стажировки нет аудиторных занятий" in i.message for i in issues)
    kit.texts.form_description = (
        "Стажёр выполняет задания на рабочем месте под руководством руководителя стажировки."
    )
    kit.texts.org_process = ""
    kit.texts.mto = ["Рабочее место с доступом к стенду"]
    issues = validate(kit, ValidationContext(registry=default_registry()), only=["V10"]).issues
    assert not any("аудиторных занятий" in i.message for i in issues)


def test_outcome_spread_over_many_modules_is_flagged() -> None:
    from app.core.validate import ValidationContext, validate

    kit = load(5)
    lo = kit.outcomes[0].id
    for topic in kit.topics:  # один РО — в темах всех модулей
        if lo not in topic.lo_ids:
            topic.lo_ids.append(lo)
    issues = validate(kit, ValidationContext(registry=default_registry()), only=["V43"]).issues
    assert any("раскрывается в темах" in i.message and lo in i.target_ids for i in issues)


def test_pk_requirements_always_name_those_still_studying(tmp_path: Path) -> None:
    import docx

    from app.render.common import load_organization
    from app.render.docx_pk import SUB
    from app.render.program import render_program_docx

    kit = load(5)
    kit.texts.audience_requirements = "Системные администраторы с опытом работы от года."
    name, _ = render_program_docx(
        kit, tmp_path, default_registry(), load_organization(Settings().organization_file)
    )
    texts = [p.text for p in docx.Document(str(tmp_path / name)).paragraphs]
    section = texts[texts.index(SUB["1.6"]) : texts.index(SUB["1.7"])]
    assert any("получающие среднее профессиональное" in t for t in section)
    assert any("опытом работы от года" in t for t in section)


def test_docx_appendix_knows_about_missing_slides_and_tests(tmp_path: Path) -> None:
    import docx

    from app.render.bundle import NO_SLIDES_NOTE, render_all
    from app.render.common import load_organization

    kit = load(5)  # стажировка: нет лекций и тестов
    render_all(kit, tmp_path, default_registry(), load_organization(Settings().organization_file))
    text = "\n".join(p.text for p in docx.Document(str(tmp_path / "program.docx")).paragraphs)
    assert NO_SLIDES_NOTE in text and "Тестовых вопросов в комплекте нет" in text
    assert "Ничего не пропущено" not in text


def test_repeated_words_from_generation_glitch_are_collapsed() -> None:
    assert apply.clean("strace, ltrace, ltrace, ltrace, ltrace и gdb") == "strace, ltrace и gdb"
    assert apply.clean("Python, Java, Python") == "Python, Java, Python"  # обычный перечень не трогаем


def test_prompt_inputs_name_audience_and_internship_without_classes() -> None:
    from app.pipeline import content_prompts as prompts

    school, internship, pk = load(1), load(5), load(3)
    assert prompts.audience_word(school) == "ученики" and prompts.audience_word(pk) == "слушатели"
    assert "«ученики»" in prompts.describe_program(school)
    assert "Аудиторных занятий нет" in prompts.describe_program(internship)
    assert "Аудиторных занятий нет" not in prompts.describe_program(pk)
    second = pk.modules_sorted()[1]
    earlier = prompts.earlier_topics(pk, second)
    first_topic = pk.topics_of(pk.modules_sorted()[0].id)[0]
    assert earlier.startswith("Уже изучено") and first_topic.title in earlier
    assert prompts.earlier_topics(pk, pk.modules_sorted()[0]).startswith("Это первый модуль")


def test_final_test_topic_is_not_part_of_the_plan() -> None:
    from app.models import ProgramType, RequestSpec
    from app.pipeline import skeleton

    draft = _skeleton_draft()
    module = draft.modules[-1]
    module.topics.append(module.topics[0].model_copy(update={"key": "t_fin", "title": "Итоговый тест"}))
    final = skeleton.finalize_spec(RequestSpec(raw_text="ПК", program_type=ProgramType.PK, total_hours=72))
    kit = skeleton.assemble(final, draft, [], default_registry(), lambda k, w: None)  # type: ignore[arg-type]
    assert not any(t.title == "Итоговый тест" for t in kit.topics)
    assert any("«Итоговый тест» не включена" in a for a in kit.request.assumptions)


def test_code_options_differing_only_in_brackets_are_not_duplicates() -> None:
    from app.core.validate import ValidationContext, validate

    kit = load(1)
    question = kit.questions[0]
    question.options = ["print(5 + 3)", "x = 5 + 3", "print(5) + 3", "5 + 3"]
    question.correct = [0]
    issues = validate(kit, ValidationContext(registry=default_registry()), only=["V18"]).issues
    assert not any("повторяются" in i.message and question.id in i.target_ids for i in issues)
    question.options = ["print(5 + 3)", "Print(5 + 3)", "x", "y"]
    issues = validate(kit, ValidationContext(registry=default_registry()), only=["V18"]).issues
    assert any("повторяются" in i.message for i in issues)


def test_one_lecture_for_several_topics_covers_them_all() -> None:
    kit = load(3)
    module = next(m for m in kit.modules_sorted() if len(apply.lecture_topics(kit, m)) >= 3)
    topics = apply.lecture_topics(kit, module)
    kit.lectures = [lec for lec in kit.lectures if lec.module_id != module.id]
    merged = LectureDraft(
        topic_ids=[t.id for t in topics],
        title="Всё сразу",
        plan=[PlanItemDraft(title="Раздел", theses=["Тезис"])],
    )
    apply.apply_lectures(kit, module, [merged])
    assert content.ContentGenerator.lecture_gaps(kit) == {}
    assert next(lec for lec in kit.lectures if lec.module_id == module.id).topic_ids == [t.id for t in topics]


def test_interview_only_module_request_can_be_split(tmp_path: Path) -> None:
    from app.llm import TruncatedResponse

    kit = load(5)
    module = kit.modules_sorted()[0]
    generator = DefaultPipeline(Settings(llm_mode="mock")).content
    seen: list[str] = []

    async def fake_ask(client, call):  # noqa: ANN001, ANN202
        seen.append(call.fixture_key)
        if len(seen) == 1:
            raise TruncatedResponse("tests", {"completion_tokens": 4000})
        return TestsDraft(interview=[])

    generator._ask = fake_ask  # type: ignore[method-assign]
    asyncio.run(generator._draft_tests(kit, module, None, {}, 4))  # type: ignore[arg-type]
    assert (
        len(seen) == 3 and seen[1].endswith("-a") and seen[2].endswith("-b")
    )  # обрезан → две половины по РО


def test_pk_skeleton_without_function_or_practical_experience_is_corrected() -> None:
    from app.models import OutcomeKind, ProgramType
    from app.pipeline import skeleton

    draft = _skeleton_draft()
    for comp in draft.competencies:
        comp.tf_ids = comp.tf_ids or ["06.015:C/11.6"]
        if not any(o.kind == OutcomeKind.MASTER for o in comp.outcomes):
            comp.outcomes[0].kind = OutcomeKind.MASTER
    # в заготовке один РО в 3 модулях из 5 — это отдельная причина повтора (V43); здесь проверяются ТФ и опыт
    good = skeleton.draft_problem(draft, ProgramType.PK, has_functions=True)
    assert "трудовых функций" not in good and "практический опыт" not in good
    draft.competencies[0].tf_ids = []
    for outcome in draft.competencies[-1].outcomes:
        outcome.kind = OutcomeKind.KNOW
    problem = skeleton.draft_problem(draft, ProgramType.PK, has_functions=True)
    assert "нет трудовых функций" in problem and "практический опыт" in problem
    dop = skeleton.draft_problem(draft, ProgramType.DOP)  # у общеразвивающей нет ТФ и опыта
    assert "трудовых функций" not in dop and "практический опыт" not in dop


def test_negated_classroom_mentions_are_not_flagged() -> None:
    from app.core.validate import ValidationContext, validate

    kit = load(5)
    kit.texts.form_description = (
        "Стажировка без аудиторных занятий: лекции не проводятся, работа — на рабочем месте."
    )
    kit.texts.org_process = ""
    kit.texts.mto = ["Рабочее место"]
    issues = validate(kit, ValidationContext(registry=default_registry()), only=["V10"]).issues
    assert not any("аудиторных занятий" in i.message for i in issues)


def test_interview_module_is_topped_up_to_norm() -> None:
    kit = load(5)
    module = kit.modules_sorted()[0]
    generator = DefaultPipeline(Settings(llm_mode="mock")).content
    quota = content.interview_quota(kit, module, generator.limits)
    kit.interview_questions = [iq for iq in kit.interview_questions if iq.module_id != module.id][:] + [
        iq for iq in kit.interview_questions if iq.module_id == module.id
    ][: quota - 1]
    needs = generator.refill_needs(kit)
    assert module.id in needs and generator._interview_gap(kit, module) == 1


def test_empty_book_sections_are_filled_from_catalog_by_title() -> None:
    from app.core import bibliography
    from app.models import CatalogPick

    kit = load(1)
    items = bibliography.shortlist(kit, default_registry())
    picks = content.with_book_sections(
        kit, items, [CatalogPick(catalog_id=items[0].key, section=SourceSection.MAIN)]
    )
    sections = [p.section for p in picks]
    assert (
        SourceSection.ADDITIONAL in sections and sections.count(SourceSection.MAIN) == 1
    )  # основную не трогаем
    added = [p.catalog_id for p in picks[1:]]
    assert all(bibliography.matches_title(kit, next(i for i in items if i.key == key)) for key in added)


def test_failed_skeleton_retry_keeps_usable_draft() -> None:
    """Повтор каркаса — улучшение, а не условие: сбой повторного вызова не роняет сессию, если исходный каркас
    собирается (без практического опыта — только предупреждение V17)."""
    from types import SimpleNamespace

    from app.llm import GatewayError
    from app.models import OutcomeKind, ProgramType, RequestSpec
    from app.pipeline import skeleton
    from tests.intake_queries import OPEN

    draft = _skeleton_draft()
    for comp in draft.competencies:
        for outcome in comp.outcomes:
            if outcome.kind == OutcomeKind.MASTER:
                outcome.kind = OutcomeKind.CAN  # изъян, который просит повтора, но не мешает сборке
    calls: list[int] = []

    class FlakyLLM:
        settings = SimpleNamespace(llm_mode="mock")

        async def complete_json(self, step, model, messages, schema, **kwargs):  # noqa: ANN001, ANN003, ANN201
            calls.append(1)
            if len(calls) > 1:
                raise GatewayError("Шлюз недоступен: ReadTimeout")
            return draft

    spec = RequestSpec(raw_text=OPEN[3]["request"], program_type=ProgramType.PK, total_hours=72)
    pipeline = DefaultPipeline(Settings(llm_mode="mock"))
    kit = asyncio.run(
        skeleton.run_skeleton(
            spec, None, FlakyLLM(), pipeline.registry, lambda k, w: pipeline.validate(k, w, None, "skeleton")
        )
    )
    assert len(calls) == 2 and kit.outcomes  # повтор был, упал, каркас собран из исходного


def test_repeat_cleanup_does_not_touch_code() -> None:
    assert apply.clean("[0, 0, 0]") == "[0, 0, 0]" and apply.clean("print(1, 1, 1)") == "print(1, 1, 1)"
