"""Отчёт автопроверки (app/render/report_md.py, report_docx.py): итог одной фразой, разделы, отдельный .docx и
сообщение API — на трёх комплектах: без замечаний, с предупреждениями, с ошибками."""

from __future__ import annotations

import re
from pathlib import Path

import docx
import pytest

from app.config import Settings
from app.core.registry import default_registry
from app.core.validate import ValidationContext, validate
from app.models import JudgeResult, Kit, StepCost
from app.render.bundle import render_all
from app.render.common import load_organization
from app.render.program import render_program_docx
from app.render.report_docx import render_report_docx
from app.render.report_md import Advice, Table, api_message, build_report, render_report_md, summary_line

FIXTURES = Path(__file__).parent / "fixtures" / "kits"
ROOT = Path(__file__).resolve().parent.parent
SAMPLE_PK = (ROOT / "docs" / "case" / "sample_pk_stazhirovka_bsa_248h.md").read_text(encoding="utf-8")

BASE_SECTIONS = [
    "Коротко",
    "Предупреждения методисту",
    "Часы по уровням",
    "Покрытие результатов обучения",
    "Источники",
    "Полнота разделов шаблона",
    "Совпадение с образцом",
    "Смысловая проверка",
    "Что пропущено и почему",
    "Стоимость и время",
]


def load(n: int) -> Kit:
    return Kit.model_validate_json((FIXTURES / f"open_q{n}.json").read_text(encoding="utf-8"))


def ok_kit() -> Kit:
    kit = load(1)
    kit.meta.generation_notes = []
    assert kit.report.status.value == "ok"
    return kit


def warnings_kit() -> Kit:
    kit = load(3)
    assert kit.report.status.value == "warnings"
    return kit


def errors_kit() -> Kit:
    """Настоящая ошибка: часы темы изменены вручную — объём не сходится, проверка это находит."""
    kit = load(4)
    kit.topics[0].hours.lk += 2
    kit.report = validate(kit, ValidationContext(registry=default_registry()))
    assert kit.report.status.value == "errors"
    return kit


def titles(kit: Kit, **kwargs) -> list[str]:  # noqa: ANN003
    return [s.title for s in build_report(kit, **kwargs).sections]


# --------------------------------------------------------------------------------------------------
# Три статуса
# --------------------------------------------------------------------------------------------------


def test_ok_kit_says_so_in_one_phrase() -> None:
    kit = ok_kit()
    assert summary_line(kit) == "Комплект собран, ошибок и предупреждений нет."
    assert titles(kit) == BASE_SECTIONS
    md = render_report_md(build_report(kit))
    assert (
        md.startswith("# Отчёт автопроверки: «")
        and "**Итог.** Комплект собран, ошибок и предупреждений нет." in md
    )
    assert "Ничего не пропущено" in md
    warnings = {sec.title: sec for sec in build_report(kit).sections}["Предупреждения методисту"]
    assert [b.block for b in warnings.blocks] == ["para"] and "Предупреждений нет" in warnings.blocks[0].text


def test_warnings_kit_lists_warnings_with_options_and_separates_sample_errors() -> None:
    kit = warnings_kit()
    summary = summary_line(kit)
    assert re.match(
        r"Комплект собран, ошибок нет; \d+ предупреждени\w+ методисту( и \d+ замечани\w+ автопроверки)? — см\. ниже\.",
        summary,
    )
    overview = {s.title: s for s in build_report(kit).sections}["Коротко"].blocks[0].items
    methodist = int(re.search(r"(\d+) предупреждени\w+ методисту", summary).group(1))
    assert f"предупреждения методисту (с вариантами решения): {methodist}" in overview
    assert "В приложенном образце найдено ошибок: 15; в программу они не перенесены." in summary
    report = build_report(kit)
    sections = {s.title: s for s in report.sections}
    cards = [b for b in sections["Предупреждения методисту"].blocks if isinstance(b, Advice)]
    assert {c.kind for c in cards} == {"reference_errors", "norm_replaced"}
    assert all(c.title and c.text and c.figures and len(c.options) >= 2 for c in cards)
    remarks = sections["Замечания автопроверки"].blocks[0].items
    assert remarks and all("Как исправить:" in line for line in remarks)
    sample = sections["Ошибки приложенного образца"]
    assert "делать ничего не нужно" in sample.blocks[0].text and len(sample.blocks[1].items) == 15
    assert "Что нужно исправить" not in sections
    md = render_report_md(report)
    assert "### 1. Ошибки в приложенном образце" in md and "**Варианты решения:**" in md


def test_errors_kit_says_what_to_fix_and_how() -> None:
    kit = errors_kit()
    summary = summary_line(kit)
    assert (
        summary.startswith("Комплект собран, но автопроверка нашла") and "до утверждения программы" in summary
    )
    sections = {s.title: s for s in build_report(kit).sections}
    fixes = sections["Что нужно исправить"].blocks[0].items
    assert fixes and all("Как исправить:" in line for line in fixes)
    hours = sections["Часы по уровням"]
    assert "не сходится" in hours.blocks[0].text
    message = api_message(kit)
    assert message.startswith("Комплект собран, но автопроверка нашла") and "Главное:" in message
    assert "см. отчёт" in message and "files.report_md" in message


@pytest.mark.parametrize("make", [ok_kit, warnings_kit, errors_kit])
def test_report_has_no_jargon(make) -> None:  # noqa: ANN001
    md = render_report_md(build_report(make(), references=[SAMPLE_PK]))
    for jargon in ("fix_by", "detail", "Issue", "LO1", "None", "True", "False"):
        assert jargon not in md, jargon
    assert not re.search(r"\bV\d{2}-\d", md)  # номера замечаний — для разработчика, не для методиста


# --------------------------------------------------------------------------------------------------
# Разделы
# --------------------------------------------------------------------------------------------------


def test_hours_table_takes_numbers_from_kit() -> None:
    kit = warnings_kit()
    table = next(
        b
        for s in build_report(kit).sections
        if s.title == "Часы по уровням"
        for b in s.blocks
        if isinstance(b, Table)
    )
    assert len(table.rows) == len(kit.modules) + len(kit.topics) + 2
    first_module = kit.modules_sorted()[0]
    assert table.rows[0][1] == f"{kit.module_hours(first_module.id).total:g}".replace(".", ",")
    assert table.rows[-1][:2] == ["Итого", f"{kit.total_hours():g}"]
    assert table.rows[-2][:2] == ["Итоговая аттестация", f"{kit.fa_hours():g}"]


def test_coverage_lists_every_outcome() -> None:
    kit = warnings_kit()
    section = next(s for s in build_report(kit).sections if s.title == "Покрытие результатов обучения")
    table = section.blocks[1]
    assert len(table.rows) == len(kit.outcomes)
    kit.questions = [q for q in kit.questions if kit.outcomes[0].id not in q.lo_ids]
    kit.seminars = [s for s in kit.seminars if kit.outcomes[0].id not in s.lo_ids]
    kit.final_attestation.lo_ids = [x for x in kit.final_attestation.lo_ids if x != kit.outcomes[0].id]
    kit.report.coverage = []  # пересчитать по комплекту
    table = next(s for s in build_report(kit).sections if s.title == "Покрытие результатов обучения").blocks[
        1
    ]
    assert table.rows[0][-1] == "не проверяется"


def test_sources_count_confirmed_unverified_and_dropped() -> None:
    kit = warnings_kit()
    kit.meta.literature_checks = {"rejected": 2, "dropped": 1, "unverified": 1}
    kit.sources[-1].verification = kit.sources[-1].verification.__class__.UNVERIFIED
    kit.report = validate(kit, ValidationContext(registry=default_registry()))
    assert kit.report.sources.unverified == 1 and kit.report.sources.rejected == 2
    section = next(s for s in build_report(kit).sections if s.title == "Источники")
    lines = section.blocks[0].items
    assert lines[1].endswith(": 1") and "отброшено при проверке: 3" in lines[2]
    by_section = section.blocks[1]
    assert isinstance(by_section, Table) and "warn" in by_section.tones
    assert "[не подтверждено]" in section.blocks[3].items[0]


def copying_text(kit: Kit, references: list[str]) -> str:
    return (
        next(
            s for s in build_report(kit, references=references).sections if s.title == "Совпадение с образцом"
        )
        .blocks[0]
        .text
    )


def test_copying_against_attached_sample() -> None:
    kit = warnings_kit()
    kit.report.copying = None  # пересчитать по переданному образцу
    assert "в норме" in copying_text(kit, [SAMPLE_PK])
    kit.texts.annotation = SAMPLE_PK[2000:6000]  # дословный кусок образца
    assert "выше порога" in copying_text(kit, [SAMPLE_PK])
    assert "не приложен" in copying_text(kit, [])
    # сводка из report.json важнее пересчёта: её считает validate по тем же вложениям
    kit.report = validate(kit, ValidationContext(registry=default_registry(), reference_texts=[SAMPLE_PK]))
    assert kit.report.copying.references == 1 and not kit.report.copying.passed
    text = copying_text(kit, [])
    assert "выше порога" in text and "Почти дословно повторяют образец 1 фрагмент" in text


def test_judge_results_and_absence() -> None:
    kit = warnings_kit()
    kit.report.judge = [
        JudgeResult(criterion="J1", passed=True, reason="Термины поясняются."),
        JudgeResult(
            criterion="J3", passed=False, reason="Результат «освоить машинное обучение» сверх объёма."
        ),
    ]
    section = next(s for s in build_report(kit).sections if s.title == "Смысловая проверка")
    assert "пройдено 1 из 2" in section.blocks[0].text
    table = section.blocks[1]
    assert table.header[-1] == "Итог" and table.rows[1][-1] == "есть замечания"
    assert table.tones == ["ok", "warn"]
    kit.report.judge, kit.meta.judge = [], []
    empty = next(s for s in build_report(kit).sections if s.title == "Смысловая проверка")
    assert "не выполнялась" in empty.blocks[0].text


def test_skipped_steps_and_cost() -> None:
    kit = warnings_kit()
    kit.report.cost = [
        StepCost(
            step="lectures:M1", calls=2, prompt_tokens=3000, completion_tokens=1500, cost_rub=0.61, seconds=42
        )
    ]
    kit.report.timing = {"lectures:M1": 42.0}
    report = build_report(kit, notes=["Лекции модуля 2 пропущены: дневной бюджет исчерпан."])
    sections = {s.title: s for s in report.sections}
    assert (
        "Лекции модуля 2 пропущены: дневной бюджет исчерпан."
        in sections["Что пропущено и почему"].blocks[0].items
    )
    cost = sections["Стоимость и время"]
    assert "0,61 ₽" in cost.blocks[0].text and "«Лекции», 42 с" in cost.blocks[0].text
    assert cost.blocks[1].rows[0][0] == "Лекции"
    assert "Пояснения к сборке" in report.summary


# --------------------------------------------------------------------------------------------------
# Отдельный report.docx, чистая программа, report.md и API
# --------------------------------------------------------------------------------------------------


@pytest.mark.parametrize("n", [1, 2, 3, 4, 5])
def test_separate_docx_report(n: int, tmp_path: Path) -> None:
    kit = load(n)
    report = build_report(kit)
    outcome = render_all(kit, tmp_path, default_registry(), load_organization(Settings().organization_file))
    assert outcome.issues == []
    program = docx.Document(str(tmp_path / outcome.files["program_docx"]))
    paragraphs = [p.text for p in program.paragraphs]
    cells = [c.text for table in program.tables for row in table.rows for c in row.cells]
    assert "Отчёт автопроверки" not in "\n".join(paragraphs + cells)
    assert outcome.files["report_docx"] == "report.docx"
    document = docx.Document(str(tmp_path / "report.docx"))
    texts = [p.text for p in document.paragraphs]
    assert texts[0] == report.title and f"Итог. {report.summary}" in texts
    assert "Коротко" in texts and "Часы по уровням" in texts and "Покрытие результатов обучения" in texts
    assert not any(
        p.text.startswith("Приложение") for p in document.paragraphs if p.style.name.startswith("Heading")
    )
    assert not document.core_properties.author and not document.core_properties.last_modified_by


def test_done_message_and_report_md_through_api(tmp_path: Path) -> None:
    from fastapi.testclient import TestClient

    from app.main import create_app
    from tests.intake_queries import OPEN
    from tests.test_content import run

    body = run(tmp_path, OPEN[3])
    assert body["files"]["report_md"].endswith("/files/report.md")
    assert body["files"]["report_docx"].endswith("/files/report.docx")
    kit = Kit.model_validate(body["kit"])
    assert body["message"].startswith(summary_line(kit, see="см. отчёт").split(" Пояснения")[0])
    assert "Подробный отчёт — файл report.md" in body["message"]
    settings = Settings(sessions_dir=tmp_path / "sessions", llm_mode="mock")
    with TestClient(create_app(settings=settings)) as client:
        md = client.get(body["files"]["report_md"]).content.decode("utf-8")
    assert md.startswith("# Отчёт автопроверки") and "## Проверка выгрузок" in md
    assert "перечитаны после записи и сверены с комплектом: расхождений нет" in md
    for name in ("program.docx", "slides.pptx", "tests.gift"):
        assert name in md


def test_program_keeps_out_report_details(tmp_path: Path) -> None:
    """В программу не идут: ISBN отброшенной моделью книги, часы и цитаты образца, свободный текст судьи, номера
    отменённых актов. В report.md всё это остаётся."""
    kit = load(3)
    reason = "На итоговую аттестацию отведено 7 часов, а в запросе — 2 часа."
    kit.report.judge = [JudgeResult(criterion="J3", passed=False, reason=reason)]
    assert any(n.startswith("Модель предложила источник") for n in kit.meta.generation_notes)
    name, issues = render_program_docx(
        kit, tmp_path, default_registry(), load_organization(Settings().organization_file)
    )
    assert issues == []
    text = "\n".join(p.text for p in docx.Document(str(tmp_path / name)).paragraphs)
    text += "\n".join(
        c.text for t in docx.Document(str(tmp_path / name)).tables for r in t.rows for c in r.cells
    )
    for forbidden in (
        "978-5-4461-0000-1",
        "Полное руководство по всему",
        "ПР (9 ч)",
        "в «Итого» 174",
        reason,
        "№ 816",
        "№ 499",
    ):
        assert forbidden not in text, forbidden
    assert "Отчёт автопроверки" not in text
    # Сокращённый ReportDoc остаётся совместимым; прежние проверки очистки не теряются.
    render_report_docx(build_report(kit, for_docx=True), tmp_path / "filtered.docx")
    filtered = "\n".join(p.text for p in docx.Document(str(tmp_path / "filtered.docx")).paragraphs)
    assert "не прошедших проверку: 1" in filtered and "Полная версия отчёта — файл report.md" in filtered
    render_report_docx(build_report(kit), tmp_path / "report.docx")
    full = docx.Document(str(tmp_path / "report.docx"))
    full_text = "\n".join(
        [p.text for p in full.paragraphs] + [c.text for t in full.tables for r in t.rows for c in r.cells]
    )
    assert "978-5-4461-0000-1" in full_text and reason in full_text and "ПР (9 ч)" in full_text
    md = render_report_md(build_report(kit))
    assert "978-5-4461-0000-1" in md and reason in md and "ПР (9 ч)" in md


def test_broken_organization_file_gives_blank_header(tmp_path: Path) -> None:
    for content in ('{"name_lines": "не список"}', "[1, 2]", "не json"):
        path = tmp_path / "org.json"
        path.write_text(content, encoding="utf-8")
        org = load_organization(path)
        assert org.name_lines == ["____"] and org.requisites == "____"
