"""Общее для рендеров .docx: организация из конфигурации, числа и слова для часов, помощники python-docx.

python-docx не умеет напрямую: шрифт eastAsia/cs, повтор шапки таблицы, заливку и границы ячеек, фиксированную
раскладку, определения маркированных списков. Всё это — здесь, через OxmlElement, чтобы рендеры ПК и ДОП
оформляли таблицы и списки одинаково.
"""

from __future__ import annotations

import json
from collections.abc import Iterable, Sequence
from pathlib import Path

from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from pydantic import BaseModel, Field

from app.core.hours import fmt_hours, plural

FONT = "Times New Roman"
FILL_MARK = "[заполнить методисту]"
BLANK = "____"
DASH = "—"
TWIPS_PER_PT = 20

ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


# --------------------------------------------------------------------------------------------------
# Организация (шапка титульного листа) — из конфигурации
# --------------------------------------------------------------------------------------------------


class Organization(BaseModel):
    name_lines: list[str] = Field(default_factory=list)
    requisites: str = ""
    address: str = ""
    city: str = ""
    order_suffix: str = ""  # «-ОБР» в номере приказа ПК


def load_organization(path: Path) -> Organization:
    """Шапка организации. Нет файла — пустые поля «____», а не выдуманные реквизиты."""
    empty = Organization(name_lines=[BLANK], requisites=BLANK, address=BLANK, city=BLANK)
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return Organization.model_validate({k: v for k, v in data.items() if not k.startswith("_")})
    except (OSError, ValueError, AttributeError):  # нет файла, битый JSON, не те поля — пустая шапка
        return empty


# --------------------------------------------------------------------------------------------------
# Числа и слова
# --------------------------------------------------------------------------------------------------


def num(value: float | None) -> str:
    """Число часов в ячейке: «4», «0,5»; нет значения — пусто."""
    return "" if value is None else fmt_hours(value)


def hours_word(value: float) -> str:
    """«час / часа / часов» для числа часов; дробное число — «часа» (0,5 часа, 1,5 часа)."""
    if abs(value - round(value)) > 1e-9:
        return "часа"
    return plural(int(round(value)), "час", "часа", "часов")


def hours_phrase(value: float) -> str:
    """«2 часа», «0,5 часа», «6 часов»."""
    return f"{fmt_hours(value)} {hours_word(value)}"


def parse_num(text: str) -> float | None:
    """Обратное к num(): «4,5» → 4.5; пусто, прочерк или не число — None."""
    value = text.strip().replace(",", ".").replace(" ", "")
    if not value or value in (DASH, "-", "–"):
        return None
    try:
        return float(value)
    except ValueError:
        return None


def upper_first(text: str) -> str:
    return text[:1].upper() + text[1:] if text else text


def or_fill(text: str | None) -> str:
    value = (text or "").strip()
    return value or FILL_MARK


def items_or_fill(items: Iterable[str]) -> list[str]:
    values = [x.strip() for x in items if x and x.strip()]
    return values or [FILL_MARK]


# --------------------------------------------------------------------------------------------------
# Шрифты и стили
# --------------------------------------------------------------------------------------------------


def _set_rfonts(rpr, name: str = FONT) -> None:  # noqa: ANN001 — CT_RPr
    """ascii/hAnsi/eastAsia/cs = name; ссылки на шрифты темы снимаются (иначе Word берёт тему, а не name)."""
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def set_doc_defaults(doc: DocumentObject, size_pt: float, name: str = FONT) -> None:
    """docDefaults: шрифт во всех четырёх слотах, размер, русский язык."""
    styles = doc.styles.element
    defaults = styles.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        styles.insert(0, defaults)
    rpr_default = defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        rpr_default = OxmlElement("w:rPrDefault")
        defaults.insert(0, rpr_default)
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)
    _set_rfonts(rpr, name)
    for tag in ("w:sz", "w:szCs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), str(int(size_pt * 2)))
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "ru-RU")
    lang.set(qn("w:eastAsia"), "ru-RU")


def style_font(style, size_pt: float | None = None, bold: bool | None = None, name: str = FONT) -> None:  # noqa: ANN001
    """Шрифт стиля: явно TNR во всех слотах, чёрный цвет, размер и жирность."""
    rpr = style.element.get_or_add_rPr()
    _set_rfonts(rpr, name)
    font = style.font
    if size_pt is not None:
        font.size = Pt(size_pt)
    if bold is not None:
        font.bold = bold
    color = rpr.find(qn("w:color"))
    if color is not None:
        rpr.remove(color)
    font.italic = False


def font_run(run, size_pt: float | None = None, bold: bool | None = None, italic: bool | None = None) -> None:  # noqa: ANN001
    _set_rfonts(run._r.get_or_add_rPr())
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


# --------------------------------------------------------------------------------------------------
# Абзацы и списки
# --------------------------------------------------------------------------------------------------


def format_par(
    par: Paragraph,
    align: str | None = None,
    before: float | None = None,
    after: float | None = None,
    first_line_cm: float | None = None,
    line_spacing: float | None = None,
    keep_next: bool = False,
    page_break_before: bool = False,
) -> Paragraph:
    fmt = par.paragraph_format
    if align:
        par.alignment = ALIGN[align]
    if before is not None:
        fmt.space_before = Pt(before)
    if after is not None:
        fmt.space_after = Pt(after)
    if first_line_cm is not None:
        from docx.shared import Cm

        fmt.first_line_indent = Cm(first_line_cm)
    if line_spacing is not None:
        fmt.line_spacing = line_spacing
    if keep_next:
        fmt.keep_with_next = True
    if page_break_before:
        fmt.page_break_before = True
    return par


def add_runs(par: Paragraph, parts: Sequence[tuple[str, bool]], size_pt: float | None = None) -> Paragraph:
    """Абзац из частей (текст, жирный): «Направленность:» жирным, значение — обычным."""
    for text, bold in parts:
        font_run(par.add_run(text), size_pt=size_pt, bold=bold)
    return par


def ensure_bullet_numbering(doc: DocumentObject, left_tw: int, hanging_tw: int, marker: str = "–") -> int:
    """Своё определение маркированного списка «–» в numbering.xml (не зависит от нумерации образца).
    Возвращает numId для numPr абзацев."""
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
        '<w:multiLevelType w:val="singleLevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
        f'<w:lvlText w:val="{marker}"/><w:lvlJc w:val="left"/>'
        f'<w:pPr><w:ind w:left="{left_tw}" w:hanging="{hanging_tw}"/></w:pPr>'
        f'<w:rPr><w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:cs="{FONT}"/></w:rPr>'
        "</w:lvl></w:abstractNum>"
    )
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    return int(numbering.add_num(abstract_id).numId)


def set_bullet(par: Paragraph, num_id: int) -> Paragraph:
    num_pr = par._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num_id
    return par


# --------------------------------------------------------------------------------------------------
# Таблицы
# --------------------------------------------------------------------------------------------------


def _tbl_pr(table: Table):  # noqa: ANN202
    return table._tbl.tblPr


def _replace_child(parent, tag: str, element) -> None:  # noqa: ANN001
    old = parent.find(qn(tag))
    if old is not None:
        parent.remove(old)
    parent.append(element)


def setup_table(
    table: Table,
    widths_tw: Sequence[int],
    border_color: str = "000000",
    border_sz: int = 4,
    margins_tw: tuple[int, int, int, int] = (60, 90, 60, 90),
) -> Table:
    """Ширины колонок (gridCol + tcW каждой ячейки), фиксированная раскладка, границы и поля ячеек на таблицу."""
    table.autofit = False
    tbl_pr = _tbl_pr(table)
    style = tbl_pr.find(qn("w:tblStyle"))
    if style is not None:
        tbl_pr.remove(style)  # без стиля из шаблона: оформление задаётся ниже явно
    total = sum(widths_tw)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    _replace_child(tbl_pr, "w:tblW", tbl_w)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    _replace_child(tbl_pr, "w:tblLayout", layout)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(border_sz))
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), border_color)
        borders.append(edge)
    _replace_child(tbl_pr, "w:tblBorders", borders)
    mar = OxmlElement("w:tblCellMar")
    for side, value in zip(("top", "left", "bottom", "right"), margins_tw, strict=True):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:w"), str(value))
        edge.set(qn("w:type"), "dxa")
        mar.append(edge)
    _replace_child(tbl_pr, "w:tblCellMar", mar)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.findall(qn("w:gridCol")), widths_tw, strict=True):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(widths_tw):
                cell.width = widths_tw[idx] * 635  # twips → EMU (1 tw = 635 EMU)
    return table


def repeat_header(table: Table, rows: int = 1) -> None:
    """Строки шапки повторяются на каждой странице (w:tblHeader)."""
    for row in table.rows[:rows]:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:tblHeader")) is None:
            tr_pr.append(OxmlElement("w:tblHeader"))


def shade(cell: _Cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    _replace_child(tc_pr, "w:shd", shd)


def cell_text(
    cell: _Cell,
    text: str | Sequence[str],
    size_pt: float,
    bold: bool = False,
    align: str = "left",
    fill: str | None = None,
    v_center: bool = False,
    bullets: int | None = None,
    line_spacing: float = 1.0,
) -> _Cell:
    """Текст ячейки: строка или список строк (каждая — отдельный абзац; bullets — numId маркера).
    Пустая ячейка тоже получает размер шрифта, чтобы строка не была выше остальных."""
    lines = [text] if isinstance(text, str) else list(text) or [""]
    first = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    for idx, line in enumerate(lines):
        par = first if idx == 0 else cell.add_paragraph()
        for run in list(par.runs):
            run._r.getparent().remove(run._r)
        format_par(par, align=align, before=0, after=0, first_line_cm=0, line_spacing=line_spacing)
        if bullets is not None and line:
            set_bullet(par, bullets)
        font_run(par.add_run(line), size_pt=size_pt, bold=bold)
    if fill:
        shade(cell, fill)
    if v_center:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return cell


def add_cell_block(
    cell: _Cell, label: str, items: Sequence[str], size_pt: float, bullets: int | None, first: bool
) -> None:
    """Блок вида работы в ячейке рабочей программы: метка жирным и пункты списком."""
    if first:
        par = cell.paragraphs[0]
        for run in list(par.runs):
            run._r.getparent().remove(run._r)
    else:
        par = cell.add_paragraph()
    format_par(par, align="left", before=0 if first else 4, after=3, first_line_cm=0, line_spacing=1.0)
    font_run(par.add_run(label), size_pt=size_pt, bold=True)
    for item in items:
        line = format_par(
            cell.add_paragraph(), align="left", before=0, after=1, first_line_cm=0, line_spacing=1.0
        )
        if bullets is not None:
            set_bullet(line, bullets)
        font_run(line.add_run(item), size_pt=size_pt)


def merge_row(table: Table, row_idx: int, start: int, end: int) -> _Cell:
    """Объединение ячеек строки по горизонтали (gridSpan); лишние пустые абзацы убираются."""
    row = table.rows[row_idx]
    merged = row.cells[start].merge(row.cells[end])
    for extra in merged.paragraphs[1:]:
        if not extra.text.strip():
            extra._p.getparent().remove(extra._p)
    return merged
