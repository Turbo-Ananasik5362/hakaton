"""Слайды лекций в .pptx (python-pptx): строгий стиль без внешних картинок.

Только лекции из slides_scope (content.slides_lectures). На лекцию: титул (заголовок = тема учебного плана),
слайд на каждый пункт плана с тезисами, итоговый слайд — len(plan) + 2 слайда; у каждого слайда заметки
докладчика. Оформление: формат 16:9, белый фон, акцентная полоса слева, шрифт Arial, тёмно-синие заголовки.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from app.models import Kit, Lecture, Slide
from app.pipeline.content import build_slides, slides_lectures
from app.render.tables import module_numbers

PPTX_NAME = "slides.pptx"
FONT = "Arial"
TITLE_COLOR = RGBColor(0x1F, 0x2A, 0x44)
TEXT_COLOR = RGBColor(0x22, 0x22, 0x22)
MUTED_COLOR = RGBColor(0x6B, 0x72, 0x80)
ACCENT = RGBColor(0x2E, 0x5A, 0x88)
WIDTH = Emu(12192000)  # 13,333 дюйма — 16:9
HEIGHT = Emu(6858000)  # 7,5 дюйма
LAYOUT_TITLE = 0
LAYOUT_TITLE_ONLY = 5
LEFT = Inches(0.8)
TEXT_W = Inches(11.7)


def lecture_slides(kit: Kit, lecture: Lecture) -> list[Slide]:
    """Слайды лекции: сохранённые в комплекте, если их ровно len(plan) + 2, иначе собранные кодом из плана."""
    slides = lecture.slides or []
    if len(slides) == len(lecture.plan) + 2:
        return slides
    return build_slides(kit, lecture)


def deck_lectures(kit: Kit) -> list[Lecture]:
    """Лекции для слайдов — в порядке занятий, затем модулей и тем."""
    order = {m.id: i for i, m in enumerate(kit.modules_sorted())}
    topics = {t.id: t.order for t in kit.topics}
    return sorted(
        slides_lectures(kit),
        key=lambda lec: (
            lec.session_no or 10**6,
            order.get(lec.module_id, 99),
            min((topics.get(t, 99) for t in lec.topic_ids), default=99),
            lec.id,
        ),
    )


def render_pptx(kit: Kit, path: Path) -> int | None:
    """Пишет слайды; возвращает число слайдов. None — в комплекте нет лекций со слайдами (файл не создаётся)."""
    lectures = deck_lectures(kit)
    if not lectures:
        return None
    prs = Presentation()
    prs.slide_width, prs.slide_height = WIDTH, HEIGHT
    numbers = module_numbers(kit)
    modules = {m.id: m for m in kit.modules}
    decks = [(lecture, lecture_slides(kit, lecture)) for lecture in lectures]
    total = sum(len(slides) for _, slides in decks)
    counter = 0
    for lecture, slides in decks:
        module = modules.get(lecture.module_id)
        number = numbers.get(lecture.module_id, "")
        if module is None:
            label = ""
        else:  # вводное занятие ДОП — без номера («Модуль 0» в программе нет)
            label = module.title if number in ("", "0") else f"Модуль {number}. {module.title}"
        session = f" · занятие {lecture.session_no}" if lecture.session_no else ""
        for idx, slide in enumerate(slides):
            counter += 1
            footer = f"{kit.program.title} · {lecture.title}"
            if idx == 0:
                plan = [p.title for p in lecture.plan] or [
                    b for b in slide.bullets if not b.startswith("Цель:")
                ]
                _title_slide(prs, slide, kit.program.title, f"{label}{session}", lecture.goal, plan)
            else:
                _content_slide(prs, slide)
            _decorate(prs.slides[-1], footer, f"{counter} / {total}")
            _notes(prs.slides[-1], slide, idx == 0)
    path.parent.mkdir(parents=True, exist_ok=True)
    _core_properties(prs, kit)
    prs.save(str(path))
    return total


# --------------------------------------------------------------------------------------------------
# Слайды
# --------------------------------------------------------------------------------------------------


def _core_properties(prs, kit: Kit) -> None:  # noqa: ANN001
    """Свои метаданные вместо оставшихся от шаблона библиотеки (автор, описание, даты 2013 года)."""
    props = prs.core_properties
    props.author = props.last_modified_by = ""
    props.comments = props.subject = props.keywords = props.category = ""
    props.title = kit.program.title
    stamp = datetime.now()
    props.created = props.modified = stamp
    props.revision = 1


def _style_run(run, size: int, bold: bool = False, color: RGBColor = TEXT_COLOR) -> None:  # noqa: ANN001
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _set_text(shape, text: str, size: int, bold: bool, color: RGBColor) -> None:  # noqa: ANN001
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    _style_run(paragraph.add_run(), size, bold, color)
    paragraph.runs[0].text = text


def _place(shape, left, top, width, height) -> None:  # noqa: ANN001
    shape.left, shape.top, shape.width, shape.height = left, top, width, height


def _title_slide(prs, slide: Slide, program: str, module: str, goal: str, plan: list[str]) -> None:  # noqa: ANN001
    page = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE])
    title, subtitle = page.shapes.title, page.placeholders[1]
    _place(title, LEFT, Inches(1.2), TEXT_W, Inches(1.8))
    _set_text(title, slide.title, 36, True, TITLE_COLOR)
    title.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    _place(subtitle, LEFT, Inches(3.2), TEXT_W, Inches(3.5))
    frame = subtitle.text_frame
    frame.clear()
    frame.word_wrap = True
    lines = [(program, 20, MUTED_COLOR), (module, 18, MUTED_COLOR)]
    if goal:
        lines.append((f"Цель: {goal}", 16, TEXT_COLOR))
    if plan:  # план лекции — на самом слайде, а не только в заметках
        lines.append(("План лекции:", 16, TEXT_COLOR))
        lines += [(f"{n}. {item}", 15, TEXT_COLOR) for n, item in enumerate(plan[:6], start=1)]
    first = True
    for text, size, color in lines:
        if not text:
            continue
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        paragraph.alignment = PP_ALIGN.LEFT
        run = paragraph.add_run()
        run.text = text
        _style_run(run, size, color=color)


def _content_slide(prs, slide: Slide) -> None:  # noqa: ANN001
    page = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_ONLY])
    title = page.shapes.title
    _place(title, LEFT, Inches(0.45), TEXT_W, Inches(1.0))
    _set_text(title, slide.title, 28, True, TITLE_COLOR)
    title.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    line = page.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT, Inches(1.5), Inches(1.2), Emu(38100))
    _fill(line, ACCENT)
    body = page.shapes.add_textbox(LEFT, Inches(1.8), TEXT_W, Inches(4.9))
    body.name = "Тезисы"
    frame = body.text_frame
    frame.word_wrap = True
    bullets = [b for b in slide.bullets if b.strip()] or [slide.title]
    size = 20 if sum(len(b) for b in bullets) <= 600 else 16
    for idx, text in enumerate(bullets):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(10)
        _bullet(paragraph)
        _thesis_runs(paragraph, text, size)


MONO_FONT = "Courier New"


def _thesis_runs(paragraph, text: str, size: int) -> None:  # noqa: ANN001
    """Тезис: первая строка — текст; строки кода после неё — через перенос строки, моноширинным шрифтом, отступы
    неразрывными пробелами (python-pptx оставляет «\\n» внутри <a:t>, и PowerPoint склеивает строки)."""
    head, *code = text.split("\n")
    run = paragraph.add_run()
    run.text = head
    _style_run(run, size)
    for line in code:
        paragraph.add_line_break()
        stripped = line.lstrip(" ")
        mono = paragraph.add_run()
        mono.text = "\u00a0" * (len(line) - len(stripped)) + stripped
        _style_run(mono, max(size - 2, 12))
        mono.font.name = MONO_FONT


def _bullet(paragraph) -> None:  # noqa: ANN001
    """Маркер «•» с висячим отступом (в текстовом поле python-pptx маркеров по умолчанию нет)."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("marL", "342900")
    p_pr.set("indent", "-342900")
    for tag in ("a:buNone", "a:buChar", "a:buFont"):
        old = p_pr.find(qn(tag))
        if old is not None:
            p_pr.remove(old)
    font = p_pr.makeelement(qn("a:buFont"), {"typeface": FONT})
    char = p_pr.makeelement(qn("a:buChar"), {"char": "•"})
    p_pr.append(font)
    p_pr.append(char)


def _fill(shape, color: RGBColor) -> None:  # noqa: ANN001
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False


def _decorate(page, footer: str, number: str) -> None:  # noqa: ANN001
    bar = page.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.25), HEIGHT)
    _fill(bar, ACCENT)
    left = page.shapes.add_textbox(LEFT, Inches(6.95), Inches(10.0), Inches(0.4))
    _set_text(left, footer, 11, False, MUTED_COLOR)
    right = page.shapes.add_textbox(Inches(11.3), Inches(6.95), Inches(1.6), Inches(0.4))
    _set_text(right, number, 11, False, MUTED_COLOR)
    right.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT


def _notes(page, slide: Slide, is_title: bool) -> None:  # noqa: ANN001
    """Заметки докладчика: из слайда; у титула — ещё план лекции, у пустых — тезисы слайда."""
    parts = [slide.notes.strip()] if slide.notes.strip() else []
    if is_title and slide.bullets:
        parts.append("План: " + "; ".join(b for b in slide.bullets if not b.startswith("Цель:")))
    if not parts:
        parts = ["; ".join(slide.bullets)]
    page.notes_slide.notes_text_frame.text = "\n".join(p for p in parts if p)
