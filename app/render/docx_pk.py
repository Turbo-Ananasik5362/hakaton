"""Программа повышения квалификации (в т.ч. в форме стажировки) в .docx — docs/templates_spec_pk.md.

Основа — копия образца data/templates/pk_stazhirovka_bsa_248h.docx (стили, поля страницы, шрифт по умолчанию):
тело очищается с сохранением sectPr, все разделы собираются заново из Kit. Исходник только читается.
Ошибки образца (findings.md, §6 спецификации) не переносятся: заголовок раздела 1 — Heading 1, ширины таблиц —
в области текста, «ЛБ» везде, метки «ПР (N часов)» — из того же числа, что колонка, реквизиты — «____».
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from app.config import ROOT, now_msk
from app.core.hours import fmt_hours
from app.core.registry import Registry
from app.models import Kit, ProgramType, SourceSection
from app.render.common import (
    BLANK,
    FILL_MARK,
    Organization,
    add_cell_block,
    at_most_hours,
    cell_text,
    font_run,
    format_par,
    hours_phrase,
    merge_row,
    num,
    or_fill,
    upper_first,
)
from app.render.report_md import ReportDoc
from app.render.tables import (
    CALENDAR_HEADER,
    FA_ROW_TITLE,
    PK_PLAN_HEADER,
    PK_UTP_HEADER,
    PK_WP_HEADER,
    PlanRow,
    ProgramTables,
    form_title,
)
from app.render.texts import (
    competency_lines,
    labor_function_lines,
    norm_lines,
    profstandard_lines,
    software_lines,
    source_lines,
)
from app.render.writer import DocWriter, Look, write_current_control, write_final_attestation

TEMPLATE = ROOT / "data" / "templates" / "pk_stazhirovka_bsa_248h.docx"
TEXT_WIDTH_TW = 9779  # A4 книжная, поля образца 2,25 / 1,5 см

LOOK = Look(
    body_size=12,
    cell_size=10,
    header_size=10,
    header_fill="F2F2F2",
    total_fill="F2F2F2",
    border_color="AAAAAA",
    border_sz=4,
    cell_margins=(80, 120, 80, 120),
    first_line_cm=0,
    line_spacing=1.0,
    para_after=6,
    list_left_tw=538,
    list_hanging_tw=358,
    cell_list_left_tw=360,
    cell_list_hanging_tw=240,
    text_width_tw=TEXT_WIDTH_TW,
)

H1 = "1. ОБЩАЯ ХАРАКТЕРИСТИКА ДОПОЛНИТЕЛЬНОЙ ПРОФЕССИОНАЛЬНОЙ ПРОГРАММЫ"
H2 = "2. СОДЕРЖАНИЕ ПРОГРАММЫ"
H3 = "3. ОРГАНИЗАЦИОННО-ПЕДАГОГИЧЕСКИЕ УСЛОВИЯ РЕАЛИЗАЦИИ ПРОГРАММЫ"
APPENDIX_1 = "Приложение 1"
APPENDIX_2 = "Приложение 2"
SUB = {
    "1.1": "1.1. Аннотация программы",
    "1.2": "1.2. Цель реализации программы",
    "1.3": "1.3. Планируемые результаты обучения",
    "1.4": "1.4. Программа разработана на основе",
    "1.5": "1.5. Перечень профессиональных компетенций в рамках имеющейся квалификации, качественное изменение "
    "которых осуществляется в результате обучения",
    "1.6": "1.6. Требования к слушателям",
    "1.7": "1.7. Форма обучения",
    "1.8": "1.8. Объём программы",
    "1.9": "1.9. Выдаваемый документ",
    "2.1": "2.1. Учебный план",
    "2.2": "2.2. Учебно-тематический план",
    "2.3": "2.3. Календарный учебный график",
    "2.4": "2.4. Рабочая программа",
    "2.5": "2.5. Оценка качества освоения программы",
    "2.5.1": "2.5.1. Формы текущего контроля и итоговой аттестации",
    "2.5.2": "2.5.2. Оценочные материалы",
    "2.5.3": "2.5.3. Анализ удовлетворённости слушателей",
    "3.1": "3.1. Материально-техническое обеспечение",
    "3.2": "3.2. Учебно-методическое и информационное обеспечение",
    "3.2.1": "3.2.1. Нормативно-правовые документы",
    "3.2.2": "3.2.2. Основная литература",
    "3.2.3": "3.2.3. Дополнительная литература",
    "3.2.4": "3.2.4. Интернет-ресурсы",
    "3.2.5": "3.2.5. Программное обеспечение и информационные системы",
    "3.3": "3.3. Организация образовательного процесса",
    "3.4": "3.4. Кадровое обеспечение образовательного процесса",
}
# Порядок разделов документа — его сверяет V31
HEADINGS: list[str] = [
    H1,
    *[SUB[k] for k in ("1.1", "1.2", "1.3", "1.4", "1.5", "1.6", "1.7", "1.8", "1.9")],
    H2,
    *[SUB[k] for k in ("2.1", "2.2", "2.3", "2.4", "2.5", "2.5.1", "2.5.2", "2.5.3")],
    H3,
    *[SUB[k] for k in ("3.1", "3.2", "3.2.1", "3.2.2", "3.2.3", "3.2.4", "3.2.5", "3.3", "3.4")],
    APPENDIX_1,
    APPENDIX_2,
]
ORDER_LINE = "Приказом № ____{suffix} от «___» __________ 20__ г."


def render_pk(
    kit: Kit, tables: ProgramTables, registry: Registry, org: Organization, path: Path, report: ReportDoc
) -> list[str]:
    """Пишет программу ПК в path; report принимается для совместимости и в программу не включается."""
    doc = Document(str(TEMPLATE))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    _core_properties(doc, kit)
    w = DocWriter(doc, LOOK)
    _title_page(w, kit, org)
    _section_1(w, kit, tables, registry)
    _section_2(w, kit, tables)
    _section_3(w, kit, registry)
    _appendices(w, kit, tables)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return w.headings


def _core_properties(doc, kit: Kit) -> None:  # noqa: ANN001
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = kit.program.title
    props.comments = ""
    stamp = datetime.now()
    props.created = stamp
    props.modified = stamp
    props.revision = 1
    props.last_printed = stamp  # в образце осталась дата печати чужого документа


# --------------------------------------------------------------------------------------------------
# Титульный лист
# --------------------------------------------------------------------------------------------------


def _line(
    w: DocWriter,
    text: str,
    align: str,
    size: float | None = None,
    bold: bool = False,
    after: float = 8,
    before: float = 0,
) -> None:
    par = w.doc.add_paragraph()
    format_par(par, align=align, before=before, after=after, first_line_cm=0, line_spacing=1.0)
    if text:
        font_run(par.add_run(text), size_pt=size, bold=bold)


def _title_page(w: DocWriter, kit: Kit, org: Organization) -> None:
    lines = org.name_lines or [BLANK]
    for idx, name in enumerate(lines):
        _line(w, name, "center", size=14 if idx == len(lines) - 1 else None, bold=True)
    _line(w, "_" * 43, "center", after=4)
    _line(w, org.requisites or BLANK, "center", size=10, after=3)
    _line(w, org.address or BLANK, "center", size=10, after=20)
    _line(w, "УТВЕРЖДЕНО", "right", bold=True, after=3)
    _line(w, ORDER_LINE.format(suffix=org.order_suffix), "right", size=11, after=30)
    _line(w, "ДОПОЛНИТЕЛЬНАЯ ПРОФЕССИОНАЛЬНАЯ ПРОГРАММА", "center", size=14, bold=True)
    kind = "ПОВЫШЕНИЯ КВАЛИФИКАЦИИ"
    if kit.program.type == ProgramType.PK_INTERNSHIP:
        kind += " В ФОРМЕ СТАЖИРОВКИ"
    _line(w, kind, "center", size=14, bold=True, after=24)
    _line(w, f"«{kit.program.title.strip() or FILL_MARK}»", "center", size=13, bold=True)
    year = (kit.meta.created_at or now_msk()).year
    _line(w, org.city or BLANK, "center", after=4, before=230)
    _line(w, str(year), "center", after=0)


# --------------------------------------------------------------------------------------------------
# Раздел 1
# --------------------------------------------------------------------------------------------------


def _section_1(w: DocWriter, kit: Kit, tables: ProgramTables, registry: Registry) -> None:
    program = kit.program
    w.heading(H1, level=1, page_break=True)

    w.heading(SUB["1.1"])
    w.para(or_fill(kit.texts.annotation))

    w.heading(SUB["1.2"])
    w.para(or_fill(program.goal))
    functions = labor_function_lines(kit, registry)
    w.para(
        "Программа направлена на совершенствование профессиональных компетенций, необходимых для выполнения "
        "трудовых функций:",
        keep_next=True,
    )
    w.bullets(functions)

    w.heading(SUB["1.3"])
    _competency_table(w, tables)

    w.heading(SUB["1.4"])
    w.bullets(profstandard_lines(kit, registry) + norm_lines(kit, registry))

    w.heading(SUB["1.5"])
    w.bullets(competency_lines(kit, registry))

    w.heading(SUB["1.6"])
    audience = program.audience
    # круг лиц для ДПП задан законом (ч. 3 ст. 76 273-ФЗ) — формулировка всегда; уточнения модели — после неё
    w.para(
        "К освоению программы допускаются лица, имеющие среднее профессиональное и (или) высшее образование, "
        "а также лица, получающие среднее профессиональное и (или) высшее образование."
    )
    extra = kit.texts.audience_requirements.strip() or (audience.requirements or "").strip()
    if extra and "получающ" not in extra.lower():
        w.para(extra)
    elif extra:  # модель сама пересказала формулировку закона — без повтора, только уточнения после неё
        rest = [
            s
            for s in re.split(r"(?<=[.!?])\s+", extra)
            if "получающ" not in s.lower() and "имеющ" not in s.lower()
        ]
        if rest:
            w.para(" ".join(rest))
    if audience.category:
        w.label_value("Категория слушателей", audience.category)
    if audience.entry_level:
        w.label_value("Входной уровень", audience.entry_level)

    w.heading(SUB["1.7"])
    text = f"Форма обучения — {form_title(kit)}."
    if program.type == ProgramType.PK_INTERNSHIP:
        text += " Программа реализуется в форме стажировки."
    w.para(text)
    if kit.texts.form_description:
        w.para(kit.texts.form_description)

    w.heading(SUB["1.8"])
    _volume(w, tables)

    w.heading(SUB["1.9"])
    w.para(
        f"{program.document}. Документ выдаётся лицам, успешно освоившим программу и прошедшим итоговую аттестацию."
    )


def _competency_table(w: DocWriter, tables: ProgramTables) -> None:
    rows = tables.competencies
    if not rows:
        w.para(FILL_MARK)
        return
    widths = [1750, 1960, 2020, 2020, 2020]  # 9770 ≤ 9779
    table = w.table(len(rows) + 1, widths)
    w.header_cells(
        table,
        0,
        ["Виды деятельности", "Профессиональные компетенции", "Практический опыт", "Умения", "Знания"],
    )
    for idx, row in enumerate(rows, start=1):
        w.text_cell(table, idx, 0, row.activity)
        w.text_cell(table, idx, 1, row.title)
        cell_text(table.rows[idx].cells[2], row.master, LOOK.cell_size, bullets=w.cell_bullets)
        cell_text(table.rows[idx].cells[3], row.can, LOOK.cell_size, bullets=w.cell_bullets)
        cell_text(table.rows[idx].cells[4], row.know, LOOK.cell_size, bullets=w.cell_bullets)
    w.after_table()


def _volume(w: DocWriter, tables: ProgramTables) -> None:
    v = tables.volume
    w.para(f"Объём программы — {hours_phrase(v.total)}, в том числе:", keep_next=True)
    contact = f"контактная работа (лекции, практические и лабораторные занятия) — {hours_phrase(v.contact)}"
    if v.distance > 0:
        contact += (
            f", из них с применением дистанционных образовательных технологий — {hours_phrase(v.distance)}"
        )
    term = FILL_MARK
    if v.weeks and v.days:
        term = f"{v.weeks} нед. ({v.days} учебных дн.)"
        if v.hours_per_day:
            term += f", {at_most_hours(v.hours_per_day)} в день"
    w.bullets(
        [
            contact,
            f"самостоятельная работа слушателей — {hours_phrase(v.srs)}",
            f"итоговая аттестация — {hours_phrase(v.fa)}",
            f"срок обучения — {term}",
        ]
    )


# --------------------------------------------------------------------------------------------------
# Раздел 2
# --------------------------------------------------------------------------------------------------


def _section_2(w: DocWriter, kit: Kit, tables: ProgramTables) -> None:
    w.heading(H2, level=1, page_break=True)
    w.heading(SUB["2.1"])
    _plan_table(w, tables.plan, PK_PLAN_HEADER, [544, 1947, 1939, 841, 842, 841, 842, 1842])
    _distance_split(w, tables.plan)
    w.heading(SUB["2.2"])
    _plan_table(w, tables.utp, PK_UTP_HEADER, [640, 2008, 1748, 850, 850, 850, 850, 1842])
    _distance_split(w, tables.utp)
    w.heading(SUB["2.3"])
    _calendar(w, tables)
    w.heading(SUB["2.4"])
    _work_program(w, kit, tables)
    _assessment(w, kit, tables)


def distance_split_lines(rows: list[PlanRow]) -> list[str]:
    """Очные и дистанционные часы контактной работы по разделам (и темам — для УТП) и всего: в таблице шаблона
    колонки ДОТ нет, а очно-заочная программа с ДОТ должна разделять их явно (оценка 21.10, п. 2.2)."""
    topics = [r for r in rows if r.kind == "topic"]
    modules = [r for r in rows if r.kind == "module"]
    parts = topics or modules  # УТП — по темам, учебный план — по разделам
    if not any(r.hours.distance > 0 for r in parts):
        return []
    lines = []
    for row in parts:
        contact, distance = row.hours.contact, row.hours.distance
        if contact <= 0:
            continue
        label = f"{row.number} {row.title}".strip()
        lines.append(
            f"{label}: очно — {fmt_hours(max(contact - distance, 0))} ч, с применением ДОТ — {fmt_hours(distance)} ч"
        )
    total_contact = sum(r.hours.contact for r in parts)
    total_distance = sum(r.hours.distance for r in parts)
    lines.append(
        f"Всего контактной работы: очно — {fmt_hours(total_contact - total_distance)} ч, с применением ДОТ — "
        f"{fmt_hours(total_distance)} ч"
    )
    return lines


def _distance_split(w: DocWriter, rows: list[PlanRow]) -> None:
    lines = distance_split_lines(rows)
    if lines:
        w.para(
            "Контактная работа в очной форме и с применением дистанционных образовательных технологий:",
            first_line=False,
        )
        w.bullets(lines)


def _plan_table(w: DocWriter, rows: list[PlanRow], header: list[str], widths: list[int]) -> None:
    table = w.table(len(rows) + 1, widths)
    w.header_cells(table, 0, header)
    for idx, row in enumerate(rows, start=1):
        if row.kind in ("fa", "total"):
            merged = merge_row(table, idx, 0, 1)
            cell_text(merged, row.title, LOOK.cell_size, bold=True, fill=LOOK.total_fill)
            w.number_cell(table, idx, 2, row.hours.total, bold=True)
            if row.kind == "total":
                for col, value in zip(
                    (3, 4, 5, 6), (row.hours.lk, row.hours.pr, row.hours.lb, row.hours.srs), strict=True
                ):
                    w.number_cell(table, idx, col, value, bold=True)
                w.text_cell(table, idx, 7, "")
            else:
                for col in (3, 4, 5, 6):
                    w.text_cell(table, idx, col, "")
                w.text_cell(table, idx, 7, row.form)
            continue
        bold = row.kind == "module" and header is PK_UTP_HEADER
        cell_text(table.rows[idx].cells[0], row.number, LOOK.cell_size, bold=bold, align="center")
        w.text_cell(table, idx, 1, row.title, bold=bold)
        hours = row.hours
        for col, value in zip(
            (2, 3, 4, 5, 6), (hours.total, hours.lk, hours.pr, hours.lb, hours.srs), strict=True
        ):
            w.number_cell(table, idx, col, value, bold=bold and col in (2, 6))
        w.text_cell(table, idx, 7, row.form)
    w.after_table()


def _calendar(w: DocWriter, tables: ProgramTables) -> None:
    v = tables.volume
    if not tables.calendar:
        w.para(FILL_MARK)
        return
    regime = f", {at_most_hours(v.hours_per_day)} в день" if v.hours_per_day else ""
    w.para(
        f"Срок обучения — {v.weeks} нед., {v.days} учебных дн.{regime}. Даты начала и окончания обучения "
        "устанавливаются приказом о зачислении слушателей."
    )
    rows = tables.calendar
    table = w.table(len(rows) + 2, [1100, 1300, 5638, 1600])
    w.header_cells(table, 0, CALENDAR_HEADER)
    for idx, row in enumerate(rows, start=1):
        cell_text(table.rows[idx].cells[0], str(row.week), LOOK.cell_size, align="center")
        cell_text(table.rows[idx].cells[1], str(row.days), LOOK.cell_size, align="center")
        w.text_cell(table, idx, 2, row.content)
        w.number_cell(table, idx, 3, row.hours)
    last = len(rows) + 1
    merged = merge_row(table, last, 0, 2)
    cell_text(merged, "Итого (час.)", LOOK.cell_size, bold=True, fill=LOOK.total_fill)
    w.number_cell(table, last, 3, tables.volume.calendar_total, bold=True)
    w.after_table()


def _work_program(w: DocWriter, kit: Kit, tables: ProgramTables) -> None:
    n_rows = 1 + sum(1 + sum(max(len(t.blocks), 1) for t in s.topics) for s in tables.sections) + 2
    table = w.table(n_rows, [3400, 5200, 1040])
    w.header_cells(table, 0, PK_WP_HEADER)
    row = 1
    for section in tables.sections:
        merged = merge_row(table, row, 0, 1)
        cell_text(merged, f"Раздел {section.number}. {section.title}", LOOK.cell_size, bold=True)
        w.number_cell(table, row, 2, section.hours.total, bold=True)
        row += 1
        for topic in section.topics:
            blocks = topic.blocks or []
            if not blocks:
                w.text_cell(table, row, 0, f"Тема {topic.number}. {topic.title}", bold=True)
                w.text_cell(table, row, 1, FILL_MARK)
                w.text_cell(table, row, 2, "")
                row += 1
                continue
            for idx, block in enumerate(blocks):
                w.text_cell(
                    table, row, 0, f"Тема {topic.number}. {topic.title}" if idx == 0 else "", bold=True
                )
                cell = table.rows[row].cells[1]
                add_cell_block(cell, block.label, block.items, LOOK.cell_size, w.cell_bullets, first=True)
                w.number_cell(table, row, 2, block.hours)
                row += 1
    fa = kit.final_attestation
    cell_text(table.rows[row].cells[0], FA_ROW_TITLE, LOOK.cell_size, bold=True, fill=LOOK.total_fill)
    w.text_cell(
        table, row, 1, upper_first(fa.form_text) if fa is not None and fa.form_text else tables.fa_title
    )
    w.number_cell(table, row, 2, tables.fa_hours, bold=True)
    row += 1
    cell_text(table.rows[row].cells[0], "Итого", LOOK.cell_size, bold=True, fill=LOOK.total_fill)
    w.text_cell(table, row, 1, "")
    w.number_cell(table, row, 2, tables.total, bold=True)
    w.after_table()


def _assessment(w: DocWriter, kit: Kit, tables: ProgramTables) -> None:
    w.heading(SUB["2.5"])
    w.para(
        "Качество освоения программы оценивается в ходе текущего контроля по разделам и на итоговой аттестации. "
        "К итоговой аттестации допускаются слушатели, выполнившие все задания текущего контроля."
    )
    w.heading(SUB["2.5.1"])
    w.para("Текущий контроль по разделам программы:", keep_next=True)
    w.bullets(
        [f"раздел {row.number} «{row.title}» — {row.form}" for row in tables.plan if row.kind == "module"]
    )
    fa = kit.final_attestation
    text = f"Итоговая аттестация — {tables.fa_title.lower()} ({hours_phrase(tables.fa_hours)})."
    if fa is not None and fa.pass_rule:
        text += f" {upper_first(fa.pass_rule.strip())}"
    w.para(text)
    w.para("Результат итоговой аттестации — «зачтено» или «не зачтено».")
    w.heading(SUB["2.5.2"])
    w.para(
        "Оценочные материалы текущего контроля приведены в Приложении 1, оценочные материалы итоговой аттестации — "
        "в Приложении 2."
    )
    w.heading(SUB["2.5.3"])
    w.para(
        kit.texts.satisfaction.strip()
        or "В начале и в конце обучения слушатели заполняют анкету: ожидания от программы, полезность тем, "
        "качество материалов и работы преподавателей. Результаты анкетирования учитываются при обновлении программы."
    )


# --------------------------------------------------------------------------------------------------
# Раздел 3 и приложения
# --------------------------------------------------------------------------------------------------


def _section_3(w: DocWriter, kit: Kit, registry: Registry) -> None:
    form = kit.program.form
    w.heading(H3, level=1, page_break=True)
    w.heading(SUB["3.1"])
    if form.eo or form.dot:
        w.para(
            "Для занятий с применением электронного обучения и дистанционных образовательных технологий слушателям "
            "предоставляется доступ к электронной информационно-образовательной среде организации."
        )
    w.bullets(kit.texts.mto)
    w.heading(SUB["3.2"])
    w.para("Учебно-методическое и информационное обеспечение программы включает:", keep_next=True)
    w.heading(SUB["3.2.1"])
    w.numbered(norm_lines(kit, registry) + profstandard_lines(kit, registry))
    w.heading(SUB["3.2.2"])
    w.numbered(source_lines(kit, SourceSection.MAIN))
    w.heading(SUB["3.2.3"])
    w.numbered(source_lines(kit, SourceSection.ADDITIONAL))
    w.heading(SUB["3.2.4"])
    w.numbered(source_lines(kit, SourceSection.INTERNET))
    w.heading(SUB["3.2.5"])
    w.numbered(software_lines(kit))
    w.heading(SUB["3.3"])
    w.para(or_fill(kit.texts.org_process))
    w.heading(SUB["3.4"])
    w.para(or_fill(kit.texts.staffing_requirements))
    lead = "Руководитель стажировки" if kit.program.type == ProgramType.PK_INTERNSHIP else "Преподаватель"
    w.label_value(lead, f"{BLANK} (ФИО), {BLANK} (должность)")


def _appendices(w: DocWriter, kit: Kit, tables: ProgramTables) -> None:
    w.heading(APPENDIX_1, level=1, page_break=True, align="right")
    w.para("Оценочные материалы текущего контроля", bold=True, align="center", keep_next=True)
    write_current_control(w, kit, "Раздел")
    w.heading(APPENDIX_2, level=1, page_break=True, align="right")
    w.para("Оценочные материалы итоговой аттестации", bold=True, align="center", keep_next=True)
    write_final_attestation(w, kit, tables.fa_title)


__all__ = ["HEADINGS", "ORDER_LINE", "render_pk", "num"]
