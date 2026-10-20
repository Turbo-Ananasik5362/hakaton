"""DocWriter — раскладка документа поверх python-docx: заголовки, абзацы, списки, таблицы в едином оформлении.

Оформление (размеры шрифта, интервалы, заливки) задаёт Look: у ПК — по образцу стажировки (TNR 12, таблицы
10 пт, серые границы), у ДОП — по образцу интенсива (TNR 14, интервал 1,5, красная строка 1,25 см).
Здесь же — оценочные материалы, общие для обоих шаблонов (вопросы тестов, задания, собеседование, ИА).
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass, field

from docx.document import Document as DocumentObject
from docx.table import Table

from app.models import FinalAttestation, Kit, Module, Seminar, TestQuestion, TestScope
from app.render.common import (
    FILL_MARK,
    add_runs,
    cell_text,
    ensure_bullet_numbering,
    font_run,
    format_par,
    hours_phrase,
    items_or_fill,
    num,
    repeat_header,
    set_bullet,
    setup_table,
    upper_first,
)
from app.render.tables import module_numbers, topic_number

LETTERS = "абвгдежзик"


@dataclass(frozen=True)
class Look:
    body_size: float  # pt
    cell_size: float
    header_size: float
    header_fill: str
    total_fill: str
    border_color: str
    border_sz: int
    cell_margins: tuple[int, int, int, int]  # top, left, bottom, right (twips)
    first_line_cm: float
    line_spacing: float
    para_after: float
    list_left_tw: int
    list_hanging_tw: int
    cell_list_left_tw: int
    cell_list_hanging_tw: int
    text_width_tw: int  # ширина области текста книжной страницы: таблицы не шире


@dataclass
class DocWriter:
    doc: DocumentObject
    look: Look
    headings: list[str] = field(default_factory=list)  # для V31: какие разделы записаны
    _bullets: int = 0
    _cell_bullets: int = 0

    def __post_init__(self) -> None:
        self._bullets = ensure_bullet_numbering(self.doc, self.look.list_left_tw, self.look.list_hanging_tw)
        self._cell_bullets = ensure_bullet_numbering(
            self.doc, self.look.cell_list_left_tw, self.look.cell_list_hanging_tw
        )

    @property
    def cell_bullets(self) -> int:
        return self._cell_bullets

    # ---- текст -------------------------------------------------------------------------------------

    def heading(self, text: str, level: int = 2, page_break: bool = False, align: str | None = None) -> None:
        par = self.doc.add_paragraph(style=f"Heading {level}")
        format_par(par, align=align, page_break_before=page_break, keep_next=True, first_line_cm=0)
        font_run(par.add_run(text))
        self.headings.append(text)

    def para(
        self,
        text: str,
        bold: bool = False,
        align: str = "justify",
        after: float | None = None,
        keep_next: bool = False,
        size: float | None = None,
        first_line: bool = True,
    ) -> None:
        par = self.doc.add_paragraph()
        format_par(
            par,
            align=align,
            before=0,
            after=self.look.para_after if after is None else after,
            first_line_cm=self.look.first_line_cm if first_line else 0,
            line_spacing=self.look.line_spacing,
            keep_next=keep_next,
        )
        font_run(par.add_run(text), size_pt=size, bold=bold)

    def label_value(self, label: str, value: str) -> None:
        par = self.doc.add_paragraph()
        format_par(
            par,
            align="justify",
            before=0,
            after=self.look.para_after,
            first_line_cm=self.look.first_line_cm,
            line_spacing=self.look.line_spacing,
        )
        add_runs(par, [(f"{label}: ", True), (value or FILL_MARK, False)])

    def subhead(self, text: str) -> None:
        """Внутренний заголовок без стиля заголовка («Знания:», «Критерии оценивания»): жирный абзац."""
        self.para(text, bold=True, align="left", keep_next=True)

    def bullets(self, items: Sequence[str]) -> None:
        for item in items_or_fill(items):
            par = self.doc.add_paragraph()
            format_par(
                par,
                align="justify",
                before=0,
                after=2,
                first_line_cm=0,
                line_spacing=self.look.line_spacing,
            )
            set_bullet(par, self._bullets)
            font_run(par.add_run(item))

    def numbered(self, items: Sequence[str]) -> None:
        """Нумерованный перечень с номерами в тексте («1. …») — как в образцах (нормативка, литература)."""
        for n, item in enumerate(items_or_fill(items), start=1):
            self.para(f"{n}. {item}", after=3, first_line=False)

    def blank(self) -> None:
        self.doc.add_paragraph()

    # ---- таблицы -----------------------------------------------------------------------------------

    def table(self, rows: int, widths_tw: Sequence[int], header_rows: int = 1) -> Table:
        look = self.look
        table = self.doc.add_table(rows=rows, cols=len(widths_tw))
        setup_table(table, widths_tw, look.border_color, look.border_sz, look.cell_margins)
        repeat_header(table, header_rows)
        return table

    def header_cells(self, table: Table, row: int, texts: Sequence[str]) -> None:
        for idx, text in enumerate(texts):
            cell_text(
                table.rows[row].cells[idx],
                text,
                self.look.cell_size,
                bold=True,
                align="center",
                fill=self.look.header_fill,
                v_center=True,
            )

    def number_cell(self, table: Table, row: int, col: int, value: float | None, bold: bool = False) -> None:
        cell_text(table.rows[row].cells[col], num(value), self.look.cell_size, bold=bold, align="center")

    def text_cell(
        self, table: Table, row: int, col: int, text: str | Sequence[str], bold: bool = False
    ) -> None:
        cell_text(table.rows[row].cells[col], text, self.look.cell_size, bold=bold)

    def after_table(self) -> None:
        par = self.doc.add_paragraph()
        format_par(par, before=0, after=self.look.para_after, first_line_cm=0)


# --------------------------------------------------------------------------------------------------
# Оценочные материалы (приложения) — общие для ПК и ДОП
# --------------------------------------------------------------------------------------------------


def write_current_control(w: DocWriter, kit: Kit, section_word: str) -> None:
    """Оценочные материалы текущего контроля по разделам: тест модуля, вопросы собеседования, практические задания."""
    numbers = module_numbers(kit)
    questions = {q.id: q for q in kit.questions}
    anything = False
    for module in kit.modules_sorted():
        tests = [t for t in kit.tests if t.scope == TestScope.MODULE and t.module_id == module.id]
        interview = [iq for iq in kit.interview_questions if iq.module_id == module.id]
        seminars = _module_seminars(kit, module)
        if not (tests or interview or seminars):
            continue
        anything = True
        number = numbers[module.id]
        title = f"{section_word} {number}. {module.title}" if number != "0" else module.title
        w.subhead(title)
        if module.control_form:
            w.para(f"Форма текущего контроля: {module.control_form}.")
        for test in tests:
            items = [questions[qid] for qid in test.question_ids if qid in questions]
            threshold = (
                f" Тест считается пройденным при {num(test.pass_threshold)} % верных ответов."
                if test.pass_threshold
                else ""
            )
            w.para(f"Тест ({len(items)} вопр.).{threshold}", bold=True, align="left", keep_next=True)
            write_questions(w, items)
        if interview:
            w.para("Вопросы для собеседования", bold=True, align="left", keep_next=True)
            for n, iq in enumerate(interview, start=1):
                w.para(f"{n}. {iq.text}", first_line=False, after=2)
                if iq.answer_points:
                    w.para("Что должно прозвучать в ответе:", first_line=False, after=1)
                    w.bullets(iq.answer_points)
                if iq.criteria:
                    w.para("Критерии оценки:", first_line=False, after=1)
                    w.bullets(iq.criteria)
        for seminar in seminars:
            write_seminar(w, kit, seminar, numbers)
    if not anything:
        w.para(FILL_MARK)


def _module_seminars(kit: Kit, module: Module) -> list[Seminar]:
    topic_ids = [t.id for t in kit.topics_of(module.id)]
    return sorted(
        (s for s in kit.seminars if s.topic_id in topic_ids),
        key=lambda s: (topic_ids.index(s.topic_id), s.id),
    )


def write_questions(w: DocWriter, items: Sequence[TestQuestion]) -> None:
    for n, question in enumerate(items, start=1):
        multiple = " (несколько верных ответов)" if len(question.correct) > 1 else ""
        w.para(f"{n}. {question.stem}{multiple}", first_line=False, after=1, keep_next=True)
        for idx, option in enumerate(question.options):
            letter = LETTERS[idx] if idx < len(LETTERS) else str(idx + 1)
            w.para(f"{letter}) {option}", first_line=False, after=0)
        answer = ", ".join(LETTERS[i] if i < len(LETTERS) else str(i + 1) for i in question.correct)
        w.para(f"Ответ: {answer or FILL_MARK}", first_line=False, after=4, size=w.look.cell_size)


def write_seminar(w: DocWriter, kit: Kit, seminar: Seminar, numbers: dict[str, str]) -> None:
    topic = next((t for t in kit.topics if t.id == seminar.topic_id), None)
    where = f" (тема {topic_number(kit, topic, numbers)})" if topic is not None else ""
    w.para(f"Практическое задание «{seminar.title}»{where}", bold=True, align="left", keep_next=True)
    if seminar.situation:
        w.label_value("Ситуация", seminar.situation)
    w.label_value("Задание", seminar.task)
    if seminar.steps:
        w.para("Порядок выполнения:", first_line=False, after=1)
        w.bullets(seminar.steps)
    w.label_value("Результат", seminar.deliverable)
    if seminar.criteria:
        w.para("Критерии оценивания:", first_line=False, after=1)
        w.bullets([f"{c.text} — до {num(c.max_points)} балл." for c in seminar.criteria])
        total = sum(c.max_points for c in seminar.criteria)
        if seminar.pass_threshold is not None:
            w.para(
                f"Задание зачтено при {num(seminar.pass_threshold)} баллах из {num(total)}.", first_line=False
            )


def write_final_attestation(w: DocWriter, kit: Kit, fa_title: str) -> None:
    """Оценочные материалы итоговой аттестации: требования, пример задания, критерии с баллами, порог, процедура."""
    fa: FinalAttestation | None = kit.final_attestation
    if fa is None:
        w.para(FILL_MARK)
        return
    w.label_value("Форма итоговой аттестации", f"{fa_title.lower()}, {hours_phrase(fa.hours)}")
    if fa.form_text:
        w.para(upper_first(fa.form_text))
    if fa.requirements:
        w.subhead("Требования к итоговой работе")
        w.bullets(fa.requirements)
    for item in fa.portfolio_items:
        w.para(f"Элемент портфолио: {item.title}", bold=True, align="left", keep_next=True)
        w.bullets(item.requirements)
    if fa.example_task is not None:
        task = fa.example_task
        w.subhead(f"Пример задания: {task.title}")
        if task.environment:
            w.label_value("Среда и исходные данные", task.environment)
        w.label_value("Задание", task.description)
        w.label_value("Результат", task.deliverable)
    if fa.data_requirements:
        w.subhead("Требования к данным")
        w.bullets(fa.data_requirements)
    if fa.expected_result:
        w.label_value("Ожидаемый результат", fa.expected_result)
    if fa.report_requirements:
        w.subhead("Требования к отчёту")
        w.bullets(fa.report_requirements)
    w.subhead("Детализированные критерии оценки")
    if fa.criteria:
        table = w.table(len(fa.criteria) + 2, [w.look.text_width_tw - 2300, 2200])
        w.header_cells(table, 0, ["Критерий", "Максимальный балл"])
        for row, crit in enumerate(fa.criteria, start=1):
            w.text_cell(table, row, 0, crit.text)
            w.number_cell(table, row, 1, crit.max_points)
        last = len(fa.criteria) + 1
        cell_text(table.rows[last].cells[0], "Итого", w.look.cell_size, bold=True, fill=w.look.total_fill)
        w.number_cell(table, last, 1, sum(c.max_points for c in fa.criteria), bold=True)
        w.after_table()
    else:
        w.para(FILL_MARK)
    w.label_value("Условие «зачтено»", fa.pass_rule)
    w.subhead("Процедура проведения")
    w.bullets(fa.procedure)
    final_tests = [t for t in kit.tests if t.scope == TestScope.FINAL]
    if final_tests:
        questions = {q.id: q for q in kit.questions}
        for test in final_tests:
            items = [questions[qid] for qid in test.question_ids if qid in questions]
            threshold = f" Порог: {num(test.pass_threshold)} % верных ответов." if test.pass_threshold else ""
            w.subhead(f"Итоговый тест ({len(items)} вопр.).{threshold}")
            write_questions(w, items)
