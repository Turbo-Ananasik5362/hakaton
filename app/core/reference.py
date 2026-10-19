"""Разбор приложенной программы (референса) кодом: .docx (python-docx), .pdf (pypdf), .txt → ParsedReference.

Что извлекается (docs/ARCHITECTURE.md §1.4):
- тип шаблона по титулу («В ФОРМЕ СТАЖИРОВКИ» / «ПОВЫШЕНИЯ КВАЛИФИКАЦИИ» / «ОБЩЕРАЗВИВАЮЩАЯ»);
- оглавление: абзацы со стилем заголовка, жирные нумерованные абзацы («1. ОБЩАЯ …», «2.5.1. …»),
  жирные абзацы с автонумерацией, «Приложение N»;
- таблицы часов (УП, УТП, план ДОП) — числа как в документе, без исправлений и пересчёта:
  ошибки образца должны дойти до проверки V45;
- блоки видов работы рабочей программы: метка «ПР (9 часов)» против числа в колонке «Всего (час.)»;
- обозначения видов работы, реквизиты приказа на титуле, коды ПС и ТФ, упомянутые приказы и письма.

Разбор не падает на любом входе: нечитаемый файл даёт пустой ParsedReference (и запись в лог).
Таблицы часов и рабочая программа извлекаются только из .docx; из .pdf и .txt — текстовая часть.
Текст внутри w:sdt (элементы управления содержимым, экспорт из Google Docs) читается через w:t.
"""

from __future__ import annotations

import io
import logging
import re
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader

from app.models import (
    ParsedReference,
    ProgramType,
    RefHeading,
    RefHoursRow,
    RefHoursTable,
    RefWorkItem,
    WorkKind,
)

logger = logging.getLogger(__name__)

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
W_P, W_TBL, W_TR, W_TC, W_R = f"{_W}p", f"{_W}tbl", f"{_W}tr", f"{_W}tc", f"{_W}r"
W_T, W_TAB, W_BR, W_CR = f"{_W}t", f"{_W}tab", f"{_W}br", f"{_W}cr"
W_NB_HYPHEN = f"{_W}noBreakHyphen"
W_SDT, W_SDT_CONTENT, W_CUSTOM_XML = f"{_W}sdt", f"{_W}sdtContent", f"{_W}customXml"
W_PPR, W_RPR, W_TCPR, W_TRPR = f"{_W}pPr", f"{_W}rPr", f"{_W}tcPr", f"{_W}trPr"
W_PSTYLE, W_B, W_VAL = f"{_W}pStyle", f"{_W}b", f"{_W}val"
W_NUMPR, W_NUMID, W_ILVL = f"{_W}numPr", f"{_W}numId", f"{_W}ilvl"
W_GRID_SPAN, W_VMERGE, W_GRID_BEFORE = f"{_W}gridSpan", f"{_W}vMerge", f"{_W}gridBefore"

HOUR_KEYS = ("total", "contact", "lk", "pr", "lb", "srs")
_KIND_KEYS = {"lk", "pr", "lb", "srs"}
_DASHES = {"-", "—", "–", "−", "‒"}
_FALSE_VALUES = {"0", "false", "off", "none"}
_TITLE_ZONE_MAX_LINES = 60
_MAX_HEADING_LEN = 200
_TEXT_HEADING_MAX_LEN = 120

_RU = "А-Яа-яЁё"
_ABBR_LABEL = "ЛК|ПР|ЛБ|ЛР|СРС|СР"
_WORD_LABEL = (
    r"(?:Теоретическ|Лекционн|Практическ|Лабораторн|Самостоятельн)[а-яё]*"
    r"(?:\s+(?:заняти|работ)[а-яё]*)?|Лекци[а-яё]*"
)
_HOURS = r"\d+(?:[.,]\d+)?"
_HOUR_WORD = r"(?:час[а-яё]*|ч\.?)"
LABEL_RE = re.compile(
    rf"^\s*(?:(?P<abbr>{_ABBR_LABEL})|(?P<word>{_WORD_LABEL}))(?![{_RU}])\s*"
    rf"(?:\(\s*(?P<h1>{_HOURS})\s*{_HOUR_WORD}\s*\)|[–—-]\s*(?P<h2>{_HOURS})\s*{_HOUR_WORD}|(?P<colon>:))",
    re.IGNORECASE,
)
_ABBR_TOKEN_RE = re.compile(rf"(?<![{_RU}])(?:{_ABBR_LABEL})(?![{_RU}])")
_PAREN_HOURS_RE = re.compile(rf"\(\s*({_HOURS})\s*{_HOUR_WORD}", re.IGNORECASE)
TOPIC_RE = re.compile(
    r"^(?:Тема\s+(?P<a>\d+(?:\.\d+)*)|(?P<b>\d+\.\d+(?:\.\d+)*))\.?(?:\s+|$)(?P<title>.*)$",
    re.IGNORECASE | re.DOTALL,
)
_SECTION_RE = re.compile(r"^(?:Раздел|Модуль)\s+\d+", re.IGNORECASE)
_TITLE_NUMBER_RE = re.compile(r"^(?:(?P<word>Раздел|Модуль|Тема)\s+)?(?P<num>\d+(?:\.\d+)*\.?)\s+\S", re.I)
_NUMBERED_RE = re.compile(r"^(?P<num>\d{1,2}(?:\.\d{1,2}){0,3})\.?\s+(?P<title>\S.*)$", re.DOTALL)
_APPENDIX_RE = re.compile(r"^Приложение(?:\s+№)?\s*\d*\s*$", re.IGNORECASE)
_HEADING_STYLE_RE = re.compile(r"^(?:heading|заголовок)\s*(\d)", re.IGNORECASE)
PS_CODE_RE = re.compile(r"(?<![\d.,])\d{2}\.\d{3}(?![\d]|[.,]\d)")
TF_CODE_RE = re.compile(rf"(?<![0-9A-Za-z{_RU}])([A-ZА-ЯЁ])\s?/\s?(\d{{2}}\.\d)(?![\d])")
_ORDER_KEYWORD = (
    r"(?:приказ[а-яё]*|письм(?:ом|о|а|у|е)|постановлени[а-яё]*|распоряжени[а-яё]*"
    r"|федеральн[а-яё]*\s+закон[а-яё]*)"
)
ORDER_REF_RE = re.compile(
    rf"(?<![{_RU}]){_ORDER_KEYWORD}(?![{_RU}])[^№\n]{{0,160}}?№\s*[0-9A-Za-z{_RU}][0-9A-Za-z{_RU}\-/.]*",
    re.IGNORECASE,
)
_REQUISITES_RE = re.compile(r"(?<![А-Яа-яЁё])(?:приказом|приказ|протоколом|протокол)\s*№", re.IGNORECASE)
_YEAR_RE = re.compile(r"(?<!\d)(?:19|20)[\d_]{2}(?!\d)")  # «2026» и заготовка «20__»
_CYR_TO_LAT = str.maketrans("АВСЕНКМОРТХ", "ABCEHKMOPTX")


# --------------------------------------------------------------------------------------------------
# Общие помощники
# --------------------------------------------------------------------------------------------------


def _low(text: str) -> str:
    return " ".join(text.lower().replace("ё", "е").split())


def _unique(items: Iterable[str]) -> list[str]:
    return list(dict.fromkeys(item for item in items if item))


def parse_hours(text: str) -> float | None:
    """Число из ячейки как в документе: «246» → 246.0, «0,5» → 0.5; пусто, «—», «-», текст → None."""
    value = text.replace("\xa0", " ").strip()
    if not value or value in _DASHES:
        return None
    match = re.fullmatch(rf"({_HOURS})\s*(?:{_HOUR_WORD})?", value, re.IGNORECASE)
    return float(match.group(1).replace(",", ".")) if match else None


def number_key(number: str | None) -> str:
    """Номер для сопоставления строк разных таблиц: «1.6.» и «1.6» → «1.6»."""
    return (number or "").strip().rstrip(".")


def _kind_of(label: str) -> WorkKind | None:
    low = label.lower()
    if low == "лк" or low.startswith(("лекци", "теоретическ", "лекционн")):
        return WorkKind.LK
    if low == "пр" or low.startswith("практическ"):
        return WorkKind.PR
    if low in ("лб", "лр") or low.startswith("лабораторн"):
        return WorkKind.LB
    if low in ("срс", "ср") or low.startswith("самостоятельн"):
        return WorkKind.SRS
    return None


@dataclass
class WorkLabel:
    text: str
    kind: WorkKind | None
    hours: float | None
    is_abbr: bool


def match_label(paragraph: str) -> WorkLabel | None:
    """Метка вида работы в начале абзаца: «ПР (9 часов): …», «СРС (12часов)», «Практическое занятие – 1 час».
    Без часов принимаются только сокращения с двоеточием («СРС:»)."""
    match = LABEL_RE.match(paragraph)
    if match is None:
        return None
    abbr = match.group("abbr")
    if match.group("colon") and not abbr:
        return None
    label = " ".join((abbr or match.group("word")).split())
    raw_hours = match.group("h1") or match.group("h2")
    hours = float(raw_hours.replace(",", ".")) if raw_hours else None
    return WorkLabel(text=label, kind=_kind_of(label), hours=hours, is_abbr=abbr is not None)


def guess_program_type(title_text: str, full_text: str = "") -> ProgramType | None:
    """Тип программы по титулу; если на титуле признаков нет — по всему тексту."""
    for text in (title_text, full_text):
        low = _low(text)
        if "в форме стажировки" in low:
            return ProgramType.PK_INTERNSHIP
        if "общеразвивающ" in low:
            return ProgramType.DOP
        if "повышения квалификации" in low:
            return ProgramType.PK
    return None


def find_ps_codes(text: str) -> list[str]:
    return _unique(PS_CODE_RE.findall(text))


def find_tf_codes(text: str) -> list[str]:
    """Коды ТФ вида «C/14.6»; буква, набранная кириллицей, приводится к латинице (как в реестре)."""
    return _unique(f"{letter.translate(_CYR_TO_LAT)}/{code}" for letter, code in TF_CODE_RE.findall(text))


def find_order_refs(lines: Iterable[str]) -> list[str]:
    refs = []
    for line in lines:
        for match in ORDER_REF_RE.finditer(line):
            refs.append(" ".join(match.group(0).split()).rstrip(".,;:"))
    return _unique(refs)


def find_requisites(title_lines: list[str], all_lines: list[str]) -> str | None:
    """Строка «Приказом № … от …» — сначала на титуле, затем в начале документа."""
    for line in [*title_lines, *all_lines[:80]]:
        if _REQUISITES_RE.search(line):
            return " ".join(line.split())
    return None


def requisites_filled(line: str | None) -> bool:
    """Номер или дата приказа заполнены цифрами (год без дня и номера не считается)."""
    if not line:
        return False
    match = re.search(r"№(?P<number>.*?)(?:(?<![А-Яа-яЁё])от(?![А-Яа-яЁё])(?P<date>.*))?$", line, re.I)
    if match is None:
        return False
    number = match.group("number") or ""
    date = _YEAR_RE.sub("", match.group("date") or "")
    return bool(re.search(r"\d", number) or re.search(r"\d", date))


def _split_number(text: str) -> tuple[str | None, str]:
    match = _NUMBERED_RE.match(text)
    if match is None:
        return None, text
    return match.group("num"), match.group("title").strip()


def _looks_like_title(title: str) -> bool:
    return bool(title) and (title[0].isupper() or title[0] in "«\"'")


def _mostly_upper(text: str) -> bool:
    letters = [ch for ch in text if ch.isalpha()]
    return bool(letters) and sum(ch.isupper() for ch in letters) / len(letters) >= 0.7


# --------------------------------------------------------------------------------------------------
# Чтение .docx: блоки, текст, стили, нумерация
# --------------------------------------------------------------------------------------------------


def _iter_blocks(parent: Any) -> Iterator[Any]:
    """Абзацы и таблицы тела в порядке документа, включая вложенные в w:sdt и w:customXml."""
    for child in parent.iterchildren():
        if child.tag in (W_P, W_TBL):
            yield child
        elif child.tag == W_SDT:
            content = child.find(W_SDT_CONTENT)
            if content is not None:
                yield from _iter_blocks(content)
        elif child.tag == W_CUSTOM_XML:
            yield from _iter_blocks(child)


def _p_text(paragraph: Any) -> str:
    """Текст абзаца по всем w:r, в том числе внутри w:sdt и гиперссылок (paragraph.text их не видит)."""
    parts: list[str] = []
    for run in paragraph.iter(W_R):
        for child in run.iterchildren():
            if child.tag == W_T:
                parts.append(child.text or "")
            elif child.tag == W_TAB:
                parts.append("\t")
            elif child.tag in (W_BR, W_CR):
                if child.get(f"{_W}type") in (None, "textWrapping"):
                    parts.append("\n")
            elif child.tag == W_NB_HYPHEN:
                parts.append("-")
    return "".join(parts)


def _run_text(run: Any) -> str:
    return "".join(child.text or "" for child in run.iterchildren(W_T))


def _bool_attr(element: Any) -> bool:
    return (element.get(W_VAL) or "true").lower() not in _FALSE_VALUES


@dataclass
class _StyleInfo:
    heading_level: int | None = None
    bold: bool = False


class _DocxContext:
    """Сведения о стилях и нумерации документа, нужные для распознавания заголовков."""

    def __init__(self, document: Any) -> None:
        self.styles: dict[str, _StyleInfo] = {}
        self.default_style = _StyleInfo()
        self.num_formats: dict[tuple[str, str], str] = {}
        self.num_counters: dict[tuple[str, str], int] = {}
        self._load_styles(document)
        self._load_numbering(document)

    def _load_styles(self, document: Any) -> None:
        for style in document.styles:
            style_id = getattr(style, "style_id", None)
            if style_id is None:
                continue
            info = _StyleInfo(heading_level=_style_heading_level(style), bold=_style_bold(style))
            self.styles[style_id] = info
            element = style.element
            is_paragraph = element.get(f"{_W}type") in (None, "paragraph")
            if is_paragraph and element.get(f"{_W}default") in ("1", "true"):
                self.default_style = info

    def _load_numbering(self, document: Any) -> None:
        try:
            numbering = document.part.numbering_part.element
        except (KeyError, NotImplementedError, AttributeError):
            return
        abstract: dict[tuple[str, str], str] = {}
        for abstract_num in numbering.iterchildren(f"{_W}abstractNum"):
            abstract_id = abstract_num.get(f"{_W}abstractNumId")
            for level in abstract_num.iterchildren(f"{_W}lvl"):
                fmt = level.find(f"{_W}numFmt")
                abstract[(abstract_id, level.get(f"{_W}ilvl"))] = fmt.get(W_VAL) if fmt is not None else ""
        for num in numbering.iterchildren(f"{_W}num"):
            ref = num.find(f"{_W}abstractNumId")
            if ref is None:
                continue
            for (abstract_id, ilvl), fmt in abstract.items():
                if abstract_id == ref.get(W_VAL):
                    self.num_formats[(num.get(f"{_W}numId"), ilvl)] = fmt

    def style_of(self, paragraph: Any) -> _StyleInfo:
        ppr = paragraph.find(W_PPR)
        pstyle = ppr.find(W_PSTYLE) if ppr is not None else None
        if pstyle is None:
            return self.default_style
        return self.styles.get(pstyle.get(W_VAL), self.default_style)

    def is_bold(self, paragraph: Any) -> bool:
        """Все непустые фрагменты абзаца жирные (прямое форматирование или стиль абзаца)."""
        style_bold = self.style_of(paragraph).bold
        runs = [run for run in paragraph.iter(W_R) if _run_text(run).strip()]
        if not runs:
            return False
        for run in runs:
            rpr = run.find(W_RPR)
            bold = rpr.find(W_B) if rpr is not None else None
            if not (_bool_attr(bold) if bold is not None else style_bold):
                return False
        return True

    def auto_number(self, paragraph: Any) -> tuple[str, int] | None:
        """Номер автонумерации десятичного списка («1»), если абзац в таком списке."""
        ppr = paragraph.find(W_PPR)
        numpr = ppr.find(W_NUMPR) if ppr is not None else None
        if numpr is None:
            return None
        num_id_el, ilvl_el = numpr.find(W_NUMID), numpr.find(W_ILVL)
        num_id = num_id_el.get(W_VAL) if num_id_el is not None else None
        ilvl = ilvl_el.get(W_VAL) if ilvl_el is not None else "0"
        if num_id is None or self.num_formats.get((num_id, ilvl)) != "decimal":
            return None
        key = (num_id, ilvl)
        self.num_counters[key] = self.num_counters.get(key, 0) + 1
        return str(self.num_counters[key]), int(ilvl) + 1


def _style_heading_level(style: Any) -> int | None:
    for _ in range(10):
        if style is None:
            return None
        match = _HEADING_STYLE_RE.match(getattr(style, "name", "") or "")
        if match:
            return max(1, int(match.group(1)))
        style = getattr(style, "base_style", None)
    return None


def _style_bold(style: Any) -> bool:
    for _ in range(10):
        if style is None:
            return False
        font = getattr(style, "font", None)
        if font is not None and font.bold is not None:
            return bool(font.bold)
        style = getattr(style, "base_style", None)
    return False


def _docx_heading(ctx: _DocxContext, paragraph: Any, text: str) -> RefHeading | None:
    text = " ".join(text.split())
    if not text or len(text) > _MAX_HEADING_LEN:
        return None
    number, title = _split_number(text)
    style_level = ctx.style_of(paragraph).heading_level
    auto = ctx.auto_number(paragraph)  # счётчик автонумерации ведётся для всех абзацев списка
    if style_level is not None:
        level = number.count(".") + 1 if number else style_level
        return RefHeading(level=level, number=number, title=title or text)
    if not ctx.is_bold(paragraph):
        return None
    if auto is not None and number is None and _looks_like_title(text):
        return RefHeading(level=auto[1], number=auto[0], title=text)
    if number is not None and _looks_like_title(title) and int(number.split(".")[0]) <= 20:
        return RefHeading(level=number.count(".") + 1, number=number, title=title)
    if _APPENDIX_RE.match(text):
        return RefHeading(level=1, number=None, title=text)
    return None


# --------------------------------------------------------------------------------------------------
# Таблицы .docx: сетка с учётом объединений
# --------------------------------------------------------------------------------------------------


@dataclass
class _Slot:
    """Ячейка сетки таблицы. key общий у всех позиций одной w:tc (gridSpan и продолжения vMerge)."""

    key: int
    paragraphs: list[str] = field(default_factory=list)
    span_dup: bool = False  # повтор ячейки, объединённой по горизонтали
    vmerge_cont: bool = False  # продолжение ячейки, объединённой по вертикали

    @property
    def text(self) -> str:
        return " ".join(" ".join(p.split()) for p in self.paragraphs if p.strip())

    @property
    def is_new(self) -> bool:
        return not self.span_dup and not self.vmerge_cont


def _int_attr(element: Any, tag: str, default: int) -> int:
    child = element.find(tag) if element is not None else None
    try:
        return int(float(child.get(W_VAL))) if child is not None else default
    except (TypeError, ValueError):
        return default


def _iter_cells(row: Any) -> Iterator[Any]:
    for child in row.iterchildren():
        if child.tag == W_TC:
            yield child
        elif child.tag in (W_SDT, W_CUSTOM_XML):
            content = child.find(W_SDT_CONTENT) if child.tag == W_SDT else child
            if content is not None:
                yield from (tc for tc in content.iterchildren(W_TC))


def _table_grid(table: Any) -> list[list[_Slot]]:
    """Сетка таблицы: по позиции на каждую колонку grid. Дубли объединённых ячеек помечены."""
    grid: list[list[_Slot]] = []
    counter = 0
    for tr in table.iterchildren(W_TR):
        row: list[_Slot] = []
        for _ in range(_int_attr(tr.find(W_TRPR), W_GRID_BEFORE, 0)):
            counter += 1
            row.append(_Slot(key=counter))
        for tc in _iter_cells(tr):
            tcpr = tc.find(W_TCPR)
            span = max(1, _int_attr(tcpr, W_GRID_SPAN, 1))
            vmerge = tcpr.find(W_VMERGE) if tcpr is not None else None
            col = len(row)
            above = grid[-1][col] if grid and col < len(grid[-1]) else None
            if vmerge is not None and vmerge.get(W_VAL) in (None, "continue") and above is not None:
                row.extend(
                    _Slot(above.key, above.paragraphs, span_dup=k > 0, vmerge_cont=True) for k in range(span)
                )
                continue
            counter += 1
            paragraphs = [_p_text(p).strip() for p in tc.iter(W_P)]
            row.extend(_Slot(counter, paragraphs, span_dup=k > 0) for k in range(span))
        grid.append(row)
    return grid


def _slot(row: list[_Slot], col: int | None) -> _Slot | None:
    if col is None or col >= len(row):
        return None
    return row[col]


def _column_key(text: str) -> str | None:
    """Нормализованный ключ колонки по тексту шапки."""
    t = _low(text)
    if not t:
        return None
    tokens = set(re.findall(r"[a-zа-я]+", t))
    if "аудиторн" in t or "контактн" in t:
        return "contact"
    if (
        "трудоемк" in t
        or t.startswith("всего")
        or "объем" in t
        or re.search(r"(?:кол-во|количество)\s+час", t)
        or t in ("часы", "часов", "час", "час.", "ч.", "ч")
    ):
        return "total"
    if "форма" in t:
        return "form"
    if "наименован" in t or "название" in t or t.startswith(("раздел", "модул", "тема")):
        return "title"
    if t.startswith("№") or "п/п" in t:
        return "number"
    if "лаборатор" in t or tokens & {"лб", "лр"}:
        return "lb"
    if "самостоят" in t or tokens & {"срс", "ср"}:
        return "srs"
    if "лекци" in t or "теори" in t or "лк" in tokens:
        return "lk"
    if "практи" in t or "пр" in tokens:
        return "pr"
    return None


def _header_rows(grid: list[list[_Slot]]) -> list[int]:
    """Строки шапки: первая строка с ≥ 2 распознанными колонками и следующие уровни шапки
    (без чисел, с новыми ячейками видов работы — как «Лекции/теория» под «в т.ч.»)."""
    rows: list[int] = []
    for index, row in enumerate(grid[:3]):
        new = [s for s in row if s.is_new and s.text]
        keys = [_column_key(s.text) for s in new]
        if index == 0:
            if sum(key is not None for key in keys) < 2:
                return []
            rows.append(0)
            continue
        if any(parse_hours(s.text) is not None for s in new) or not any(k in HOUR_KEYS for k in keys):
            break
        rows.append(index)
    return rows


def _is_column_numbering(texts: list[str]) -> bool:
    """Строка нумерации колонок «1 | 2 | 3 | …» под шапкой."""
    values = [t for t in texts if t]
    return len(values) >= 3 and values == [str(i) for i in range(1, len(values) + 1)]


def _row_type(number: str, title: str) -> str:
    low = _low(title)
    if "итоговая аттестация" in low or low.startswith("итоговая"):
        return "fa"
    if re.match(r"^(?:итого|всего)(?![а-я])", low):
        return "total"
    if "текущий контроль" in low or "промежуточная аттестация" in low:
        return "control"
    num = number.strip()
    if re.fullmatch(r"\d+\.?", num):
        return "section"
    if re.fullmatch(r"\d+(?:\.\d+)+\.?", num):
        return "topic"
    if _SECTION_RE.match(title.strip()):
        return "section"
    if re.match(r"^тема\s+\d+", low):
        return "topic"
    if low.startswith("вводн"):
        return "section"
    return "other"


@dataclass
class _HoursLayout:
    col_keys: list[str | None]
    raw_header: list[str]
    header_rows: list[int]

    def first(self, key: str) -> int | None:
        return next((j for j, k in enumerate(self.col_keys) if k == key), None)

    def hour_columns(self) -> list[tuple[int, str]]:
        seen: set[str] = set()
        result = []
        for j, key in enumerate(self.col_keys):
            if key in HOUR_KEYS and key not in seen:
                seen.add(key)
                result.append((j, key))
        return result


def _hours_layout(grid: list[list[_Slot]]) -> _HoursLayout | None:
    header_rows = _header_rows(grid)
    if not header_rows:
        return None
    width = max(len(row) for row in grid)
    col_keys: list[str | None] = [None] * width
    raw: list[str] = [""] * width
    for j in range(width):
        for r in header_rows:
            slot = _slot(grid[r], j)
            if slot is None or not slot.text:
                continue
            raw[j] = slot.text
            key = _column_key(slot.text)
            if key is not None:
                col_keys[j] = key
    layout = _HoursLayout(col_keys, raw, header_rows)
    hour_keys = {key for _, key in layout.hour_columns()}
    is_hours = ("total" in hour_keys and len(hour_keys) >= 2) or len(hour_keys & _KIND_KEYS) >= 3
    return layout if is_hours else None


def _raw_header(grid: list[list[_Slot]], layout: _HoursLayout) -> list[str]:
    """Тексты шапки по колонкам (нижний уровень), без повторов объединённых по горизонтали ячеек."""
    last = grid[layout.header_rows[-1]]
    result: list[str] = []
    for j, text in enumerate(layout.raw_header):
        slot = _slot(last, j)
        if slot is not None and slot.span_dup and j > 0 and layout.raw_header[j - 1] == text:
            continue
        result.append(text)
    return result


def _hours_row(row: list[_Slot], layout: _HoursLayout) -> RefHoursRow | None:
    texts = [slot.text if slot.is_new else "" for slot in row]
    if not any(texts) or _is_column_numbering(texts):
        return None
    j_number, j_title, j_form = layout.first("number"), layout.first("title"), layout.first("form")
    number_slot, title_slot = _slot(row, j_number), _slot(row, j_title)
    number = number_slot.text if number_slot is not None and number_slot.is_new else ""
    title = title_slot.text if title_slot is not None and not title_slot.vmerge_cont else ""
    if number_slot is not None and title_slot is not None and number_slot.key == title_slot.key:
        number, title = "", number_slot.text
    if not title:
        title = next((s.text for s in row if s.is_new and s.text and parse_hours(s.text) is None), "")
    if not number:
        match = _TITLE_NUMBER_RE.match(title)
        if match and (match.group("word") or "." in match.group("num")):
            number = match.group("num")
    values: dict[str, float | None] = {}
    for j, key in layout.hour_columns():
        slot = _slot(row, j)
        values[key] = parse_hours(slot.text) if slot is not None and slot.is_new else None
    form_slot = _slot(row, j_form)
    form_text = form_slot.text if form_slot is not None and form_slot.is_new else ""
    form = form_text if form_text and form_text not in _DASHES else None
    return RefHoursRow(
        number=number or None,
        title=title,
        row_type=_row_type(number, title),  # type: ignore[arg-type]
        values=values,
        form=form,
    )


def _table_kind(
    columns: list[str], rows: list[RefHoursRow], context_heading: str, program_type: ProgramType | None
) -> str:
    """plan / utp / dop_plan. Тип программы с титула главнее раскладки колонок: у ПК тоже бывает
    колонка «аудиторных»; без типа план ДОП узнаётся по «Всего аудиторных» при отсутствии ЛБ."""
    context = _low(context_heading)
    is_utp = "тематическ" in context or (
        "учебный план" not in context and any(r.row_type == "topic" for r in rows)
    )
    dop_layout = "contact" in columns and "lb" not in columns
    if program_type is ProgramType.DOP or (program_type is None and dop_layout):
        return "utp" if "тематическ" in context else "dop_plan"
    return "utp" if is_utp else "plan"


def _hours_table(
    grid: list[list[_Slot]], context_heading: str, program_type: ProgramType | None = None
) -> RefHoursTable | None:
    layout = _hours_layout(grid)
    if layout is None:
        return None
    rows = [
        parsed
        for row in grid[layout.header_rows[-1] + 1 :]
        if (parsed := _hours_row(row, layout)) is not None
    ]
    columns = [key for _, key in layout.hour_columns()]
    kind = _table_kind(columns, rows, context_heading, program_type)
    return RefHoursTable(kind=kind, columns=columns, raw_header=_raw_header(grid, layout), rows=rows)  # type: ignore[arg-type]


# --------------------------------------------------------------------------------------------------
# Рабочая программа (ПК) и УТП ДОП: метки видов работы
# --------------------------------------------------------------------------------------------------


def _special_kind_label(first_text: str) -> str | None:
    low = _low(first_text)
    if "текущий контроль" in low or "промежуточная аттестация" in low or "итоговая аттестация" in low:
        return " ".join(first_text.split())
    return None


def _workprogram_hours_col(grid: list[list[_Slot]]) -> int | None:
    if not grid:
        return None
    return next((j for j, s in enumerate(grid[0]) if s.is_new and _column_key(s.text) == "total"), None)


def _row_labels(row: list[_Slot], hours_col: int | None) -> list[WorkLabel]:
    labels = []
    for j, slot in enumerate(row):
        if j == 0 or j == hours_col or not slot.is_new:
            continue
        labels.extend(label for p in slot.paragraphs if (label := match_label(p)) is not None)
    return labels


def _workprogram_items(grid: list[list[_Slot]]) -> list[RefWorkItem]:
    """Блоки видов работы по темам. Таблица учитывается, только если в первой колонке есть темы
    («Тема 1.2. …») и в строках есть метки видов работы. Так же разбирается УТП ДОП
    («Теоретическое занятие – 0,5 часа»): колонки часов там нет, column_hours = None.
    Строки «Текущий контроль» / «Итоговая аттестация» — отдельные блоки с kind=None и без темы;
    строки раздела и «Итого» пропускаются."""
    if not grid:
        return []
    hours_col = _workprogram_hours_col(grid)
    items: list[RefWorkItem] = []
    topic_number: str | None = None
    topic_title = ""
    has_topic = False
    for index, row in enumerate(grid):
        if not row:
            continue
        first = row[0]
        first_text = first.text if first.is_new else ""
        labels = _row_labels(row, hours_col)
        hours_slot = _slot(row, hours_col)
        column_hours = parse_hours(hours_slot.text) if hours_slot is not None and hours_slot.is_new else None
        topic_match = TOPIC_RE.match(first_text) if first_text else None
        if index == 0 and topic_match is None and not labels:
            continue  # шапка
        if topic_match is not None:
            has_topic = True
            topic_number = topic_match.group("a") or topic_match.group("b")
            topic_title = " ".join(topic_match.group("title").split())
        elif first_text:
            full_width = all(slot.key == first.key for slot in row)
            low = _low(first_text)
            if full_width or _SECTION_RE.match(first_text) or re.match(r"^(?:итого|всего)(?![а-я])", low):
                topic_number, topic_title = None, ""
                continue
            special = _special_kind_label(first_text)
            if special is not None:
                content = " ".join(s.text for j, s in enumerate(row) if j and j != hours_col and s.is_new)
                paren = _PAREN_HOURS_RE.search(content)
                items.append(
                    RefWorkItem(
                        kind_label=special,
                        label_hours=float(paren.group(1).replace(",", ".")) if paren else None,
                        column_hours=column_hours,
                    )
                )
                topic_number, topic_title = None, ""
                continue
            topic_number, topic_title = None, " ".join(first_text.split())
        elif topic_number is None and not topic_title:
            continue
        if not labels:
            if column_hours is not None and (topic_number or topic_title):
                items.append(
                    RefWorkItem(topic_number=topic_number, topic_title=topic_title, column_hours=column_hours)
                )
            continue
        for label in labels:
            items.append(
                RefWorkItem(
                    topic_number=topic_number,
                    topic_title=topic_title,
                    kind_label=label.text,
                    kind=label.kind,
                    label_hours=label.hours,
                    column_hours=column_hours if len(labels) == 1 else None,
                )
            )
    has_labels = any(item.kind is not None for item in items)
    return items if has_topic and has_labels else []


def _workprogram_abbr_labels(grid: list[list[_Slot]], items: list[RefWorkItem]) -> list[str]:
    if not items:
        return []
    hours_col = _workprogram_hours_col(grid)
    return [label.text for row in grid for label in _row_labels(row, hours_col) if label.is_abbr]


# --------------------------------------------------------------------------------------------------
# Сборка результата
# --------------------------------------------------------------------------------------------------


@dataclass
class _Collected:
    lines: list[str] = field(default_factory=list)
    title_lines: list[str] = field(default_factory=list)
    headings: list[RefHeading] = field(default_factory=list)
    hour_tables: list[RefHoursTable] = field(default_factory=list)
    workprogram: list[RefWorkItem] = field(default_factory=list)
    labels: list[str] = field(default_factory=list)


def _build(name: str, data: _Collected) -> ParsedReference:
    title_lines = data.title_lines if data.headings else data.lines[:_TITLE_ZONE_MAX_LINES]
    full_text = "\n".join(data.lines)
    requisites = find_requisites(title_lines, data.lines)
    ref_lines = [line for line in data.lines if line != requisites]
    return ParsedReference(
        source_name=name,
        program_type_guess=guess_program_type(" ".join(title_lines), full_text),
        headings=data.headings,
        hour_tables=data.hour_tables,
        workprogram=data.workprogram,
        work_kind_labels=_unique(data.labels),
        order_requisites=requisites,
        requisites_filled=requisites_filled(requisites),
        ps_codes=find_ps_codes(full_text),
        tf_codes=find_tf_codes(full_text),
        order_refs=find_order_refs(ref_lines),
    )


def _collect_docx(data: bytes) -> _Collected:
    document = Document(io.BytesIO(data))
    ctx = _DocxContext(document)
    out = _Collected()
    last_heading = ""
    for block in _iter_blocks(document.element.body):
        if block.tag == W_P:
            text = " ".join(_p_text(block).split())
            if not text:
                continue
            out.lines.append(text)
            heading = _docx_heading(ctx, block, text)
            if heading is not None:
                out.headings.append(heading)
                last_heading = heading.title
            elif not out.headings and len(out.title_lines) < _TITLE_ZONE_MAX_LINES:
                out.title_lines.append(text)
            continue
        grid = _table_grid(block)
        out.lines.extend(text for row in grid for slot in row if slot.is_new and (text := slot.text))
        table = _hours_table(grid, last_heading, guess_program_type(" ".join(out.title_lines)))
        if table is not None:
            out.hour_tables.append(table)
            out.labels.extend(m for text in table.raw_header for m in _ABBR_TOKEN_RE.findall(text))
            continue
        items = _workprogram_items(grid)
        out.workprogram.extend(items)
        out.labels.extend(_workprogram_abbr_labels(grid, items))
    return out


def _text_heading(line: str) -> RefHeading | None:
    """Заголовок в тексте без разметки (.txt, .pdf): «2.1. Учебный план», «1. ОБЩАЯ …», «Приложение 1»."""
    if len(line) > _TEXT_HEADING_MAX_LEN:
        return None
    if _APPENDIX_RE.match(line):
        return RefHeading(level=1, number=None, title=line)
    number, title = _split_number(line)
    if number is None or not _looks_like_title(title) or int(number.split(".")[0]) > 20:
        return None
    if "№" in title or "http" in title.lower() or " — " in title or title.endswith((".", ";", ",", ":")):
        return None
    if "." not in number and not _mostly_upper(title):
        return None  # «1. Текст» без подномера — обычно пункт списка, если не набран прописными
    return RefHeading(level=number.count(".") + 1, number=number, title=title)


def _collect_text(raw_lines: Iterable[str]) -> _Collected:
    out = _Collected()
    for raw in raw_lines:
        line = " ".join(raw.split())
        if not line:
            continue
        out.lines.append(line)
        heading = _text_heading(line)
        if heading is not None:
            out.headings.append(heading)
        elif not out.headings and len(out.title_lines) < _TITLE_ZONE_MAX_LINES:
            out.title_lines.append(line)
        label = match_label(line)
        if label is not None and label.is_abbr:
            out.labels.append(label.text)
    return out


def _pdf_lines(data: bytes) -> list[str]:

    reader = PdfReader(io.BytesIO(data))
    lines: list[str] = []
    for page in reader.pages:
        try:
            lines.extend((page.extract_text() or "").splitlines())
        except Exception as exc:  # битая страница не должна ронять разбор остальных
            logger.warning("Страница PDF не прочитана: %s", exc)
    return lines


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _detect_format(data: bytes, name: str) -> str:
    if data.startswith(b"%PDF"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        return "docx"
    suffix = Path(name).suffix.lower()
    if suffix in (".docx", ".pdf"):
        return suffix[1:]
    return "txt"


def _read_source(source: Path | str | bytes, filename: str | None) -> tuple[bytes, str]:
    if isinstance(source, (bytes, bytearray, memoryview)):
        return bytes(source), filename or ""
    path = Path(source)
    return path.read_bytes(), filename or path.name


def _collect(data: bytes, name: str) -> _Collected:
    kind = _detect_format(data, name)
    if kind == "docx":
        return _collect_docx(data)
    if kind == "pdf":
        return _collect_text(_pdf_lines(data))
    return _collect_text(_decode_text(data).splitlines())


def _safe_collect(source: Path | str | bytes, filename: str | None) -> tuple[_Collected, str]:
    name = filename or (Path(source).name if isinstance(source, (str, Path)) else "")
    try:
        data, name = _read_source(source, filename)
        return _collect(data, name), name
    except Exception as exc:  # вложение методиста может быть любым — разбор не должен падать
        logger.warning("Референс %r не разобран: %s: %s", name, exc.__class__.__name__, exc)
        return _Collected(), name


# --------------------------------------------------------------------------------------------------
# Публичный API
# --------------------------------------------------------------------------------------------------


def parse_reference(source: Path | str | bytes, filename: str | None = None) -> ParsedReference:
    """Разбирает приложенную программу (.docx, .pdf, .txt; формат — по сигнатуре, затем по расширению).
    Никогда не падает: нечитаемый файл даёт ParsedReference только с source_name."""
    collected, name = _safe_collect(source, filename)
    return _build(name, collected)


def extract_text(source: Path | str | bytes, filename: str | None = None) -> str:
    """Плоский текст вложения (абзацы, ячейки таблиц, текст в w:sdt) — для антиплагиата и intake."""
    collected, _ = _safe_collect(source, filename)
    return "\n".join(collected.lines)


def reference_outline(reference: ParsedReference, limit: int = 80) -> str:
    """Оглавление референса для промпта LLM: «2.1. Учебный план» с отступом по уровню."""
    lines = []
    for heading in reference.headings[:limit]:
        prefix = f"{heading.number}. " if heading.number else ""
        lines.append(f"{'  ' * (heading.level - 1)}{prefix}{heading.title}")
    return "\n".join(lines)
